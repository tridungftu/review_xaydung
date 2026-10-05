# -*- coding: utf-8 -*-
"""
Lõi xử lý của app kiểm tra hồ sơ: đọc file, phân loại theo danh mục,
tính tình trạng đủ/thiếu, tìm đoạn văn liên quan để hỏi đáp.
Chạy hoàn toàn trên máy. Local AI (Ollama) là tuỳ chọn.
"""
import csv
import json
import math
import os
import re
import unicodedata
import urllib.request
from collections import Counter
from pathlib import Path

HERE = Path(__file__).parent
CHECKLIST_CSV = HERE / "danh_muc_tai_lieu.csv"


# ----------------------------------------------------------------- tiện ích
def bo_dau(s: str) -> str:
    """Bỏ dấu tiếng Việt, đưa về chữ thường để so khớp từ khoá."""
    s = str(s or "").replace("đ", "d").replace("Đ", "D")
    s = unicodedata.normalize("NFD", s)
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn")
    return re.sub(r"\s+", " ", s.lower()).strip()


def load_checklist(path=CHECKLIST_CSV):
    with open(path, encoding="utf-8-sig") as f:
        rows = list(csv.DictReader(f))
    for r in rows:
        r["bat_buoc"] = str(r["bat_buoc"]).strip() in ("1", "x", "X", "True", "true")
        r["uu_tien"] = int(r.get("uu_tien") or 3)
        r["kw"] = [k.strip() for k in r["tu_khoa"].split("|") if k.strip()]
    return rows


# ----------------------------------------------------------------- đọc file
def extract_text(path: Path, max_chars=200_000) -> tuple[str, dict]:
    """Trả về (text, meta). Không có thư viện thì trả text rỗng, app vẫn chạy."""
    ext = path.suffix.lower()
    meta = {"loai": ext.lstrip("."), "trang": None, "ghi_chu": ""}
    text = ""
    try:
        if ext == ".pdf":
            import pdfplumber
            with pdfplumber.open(path) as pdf:
                meta["trang"] = len(pdf.pages)
                parts = []
                for p in pdf.pages[:60]:
                    parts.append(p.extract_text() or "")
                text = "\n".join(parts)
            if len(text.strip()) < 30:
                meta["ghi_chu"] = "PDF scan (không có lớp chữ) - cần OCR"
                text = ocr_pdf(path) or ""
        elif ext == ".docx":
            import docx
            d = docx.Document(path)
            text = "\n".join(p.text for p in d.paragraphs)
            for t in d.tables:
                for row in t.rows:
                    text += "\n" + " | ".join(c.text for c in row.cells)
        elif ext in (".xlsx", ".xlsm"):
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            meta["sheets"] = wb.sheetnames
            parts = ["Sheets: " + ", ".join(wb.sheetnames)]
            for ws in wb.worksheets[:12]:
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    if i > 400:
                        break
                    vals = [str(v) for v in row if v not in (None, "")]
                    if vals:
                        parts.append(" | ".join(vals))
            text = "\n".join(parts)
        elif ext in (".txt", ".md", ".csv"):
            text = path.read_text(encoding="utf-8", errors="ignore")
        elif ext in (".png", ".jpg", ".jpeg", ".tif", ".tiff"):
            text = ocr_image(path) or ""
            if not text:
                meta["ghi_chu"] = "Ảnh - cần OCR (cài pytesseract + gói tiếng Việt)"
    except Exception as e:  # file hỏng, có mật khẩu...
        meta["ghi_chu"] = f"Lỗi đọc file: {e}"
    return text[:max_chars], meta


def ocr_image(path):
    try:
        import pytesseract
        from PIL import Image
        return pytesseract.image_to_string(Image.open(path), lang="vie+eng")
    except Exception:
        return ""


def ocr_pdf(path, max_pages=10):
    try:
        import pypdfium2 as pdfium
        import pytesseract
        pdf = pdfium.PdfDocument(str(path))
        out = []
        for i in range(min(len(pdf), max_pages)):
            img = pdf[i].render(scale=2).to_pil()
            out.append(pytesseract.image_to_string(img, lang="vie+eng"))
        return "\n".join(out)
    except Exception:
        return ""


# ----------------------------------------------------------------- phân loại
CODE_RE = re.compile(r"^\s*([A-F]\d{2})[\s_\-.]", re.I)


def classify(filename: str, text: str, checklist, top=3):
    """
    Chấm điểm từng mục trong danh mục.
    - Tên file bắt đầu bằng mã (vd 'D01_...') thì nhận luôn, điểm 100.
    - Còn lại: từ khoá khớp tên file x3, khớp nội dung (2.000 ký tự đầu x2, phần còn lại x1).
    """
    m = CODE_RE.match(filename)
    codes = {r["ma"].upper(): r for r in checklist}
    if m and m.group(1).upper() in codes:
        return [(m.group(1).upper(), 100.0, "mã trong tên file")]
    fn = bo_dau(Path(filename).stem).replace("_", " ")
    head = bo_dau(text[:2000])
    body = bo_dau(text[2000:20000])
    scores = []
    for r in checklist:
        sc, hits = 0.0, []
        for kw in r["kw"]:
            k = bo_dau(kw)
            pat = k if k.startswith("\\b") else re.escape(k)
            if re.search(pat, fn):
                sc += 3
                hits.append(kw)
            if re.search(pat, head):
                sc += 2
                hits.append(kw)
            elif re.search(pat, body):
                sc += 1
        if sc:
            scores.append((r["ma"], sc, "từ khoá: " + ", ".join(sorted(set(hits)))[:80]))
    scores.sort(key=lambda x: -x[1])
    return scores[:top]


# ----------------------------------------------------------------- kho hồ sơ
class Kho:
    """Lưu file upload và chỉ mục phân loại trong một thư mục trên máy."""

    def __init__(self, root: Path):
        self.root = Path(root)
        self.files = self.root / "files"
        self.cache = self.root / "text"
        self.files.mkdir(parents=True, exist_ok=True)
        self.cache.mkdir(parents=True, exist_ok=True)
        self.index_path = self.root / "index.json"
        self.index = json.loads(self.index_path.read_text("utf-8")) if self.index_path.exists() else {}

    def save(self):
        self.index_path.write_text(json.dumps(self.index, ensure_ascii=False, indent=1), "utf-8")

    def add(self, name: str, data: bytes, checklist, llm=None):
        p = self.files / name
        p.write_bytes(data)
        text, meta = extract_text(p)
        (self.cache / (name + ".txt")).write_text(text, "utf-8")
        cand = classify(name, text, checklist)
        code, score, how = (cand[0] if cand else ("", 0, "không khớp"))
        if llm and llm.ok() and score < 4:
            ai_code = llm.classify(name, text, checklist)
            if ai_code:
                code, score, how = ai_code, score, "local AI"
        self.index[name] = dict(ma=code, diem=score, cach=how, goi_y=[c[0] for c in cand],
                                kich_thuoc=len(data), **meta, xac_nhan=False)
        self.save()
        return self.index[name]

    def set_code(self, name, code):
        if name in self.index:
            self.index[name]["ma"] = code
            self.index[name]["cach"] = "người dùng chọn"
            self.index[name]["xac_nhan"] = True
            self.save()

    def remove(self, name):
        self.index.pop(name, None)
        for p in (self.files / name, self.cache / (name + ".txt")):
            if p.exists():
                p.unlink()
        self.save()

    def text(self, name):
        p = self.cache / (name + ".txt")
        return p.read_text("utf-8") if p.exists() else ""


def tinh_trang(checklist, index):
    """Bảng tình trạng theo danh mục: Đã có / Cần xác nhận / Thiếu."""
    by_code = {}
    for fn, info in index.items():
        by_code.setdefault(info.get("ma"), []).append((fn, info))
    out = []
    for r in checklist:
        files = by_code.get(r["ma"], [])
        if not files:
            st = "Thiếu" if r["bat_buoc"] else "Chưa có (không bắt buộc)"
        elif any(i.get("xac_nhan") or i.get("diem", 0) >= 4 for _, i in files):
            st = "Đã có"
        else:
            st = "Cần xác nhận"
        out.append(dict(ma=r["ma"], nhom=r["nhom"], ten=r["ten"], bat_buoc=r["bat_buoc"],
                        uu_tien=r["uu_tien"], tinh_trang=st, so_file=len(files),
                        file=", ".join(f for f, _ in files), muc_dich=r["muc_dich"],
                        mat_xich=r["mat_xich"]))
    return out


# ----------------------------------------------------------------- tìm kiếm
def chunks(text, size=900, overlap=150):
    i = 0
    while i < len(text):
        yield text[i:i + size]
        i += size - overlap


def tokenize(s):
    return re.findall(r"[a-z0-9]+", bo_dau(s))


def search(kho: Kho, query: str, k=5):
    """BM25 rút gọn trên các đoạn văn bản đã trích. Không cần vector DB."""
    docs = []
    for fn in kho.index:
        for j, ch in enumerate(chunks(kho.text(fn))):
            docs.append((fn, j, ch, Counter(tokenize(ch))))
    if not docs:
        return []
    q = tokenize(query)
    N = len(docs)
    avg = sum(sum(d[3].values()) for d in docs) / N
    df = Counter()
    for d in docs:
        for t in set(d[3]):
            df[t] += 1
    res = []
    for fn, j, ch, tf in docs:
        L = sum(tf.values()) or 1
        s = 0.0
        for t in q:
            if t in tf:
                idf = math.log(1 + (N - df[t] + 0.5) / (df[t] + 0.5))
                s += idf * tf[t] * 2.2 / (tf[t] + 1.2 * (0.25 + 0.75 * L / avg))
        if s > 0:
            res.append((s, fn, ch))
    res.sort(key=lambda x: -x[0])
    return res[:k]


# ----------------------------------------------------------------- local AI
class LocalLLM:
    """Gọi Ollama chạy trên máy (mặc định http://localhost:11434). Không gửi dữ liệu ra ngoài."""

    def __init__(self, model="qwen3.6:35b", host="http://localhost:11434", timeout=600):
        self.model, self.host, self.timeout = model, host.rstrip("/"), timeout
        self._ok = None

    def ok(self):
        if self._ok is None:
            try:
                urllib.request.urlopen(self.host + "/api/tags", timeout=2)
                self._ok = True
            except Exception:
                self._ok = False
        return self._ok

    def chat(self, system, user):
        body = json.dumps({"model": self.model, "stream": False, "think": False,
                           "options": {"temperature": 0, "num_ctx": 32768},
                           "messages": [{"role": "system", "content": system},
                                        {"role": "user", "content": user}]}).encode()
        req = urllib.request.Request(self.host + "/api/chat", data=body,
                                     headers={"Content-Type": "application/json"})
        with urllib.request.urlopen(req, timeout=self.timeout) as r:
            return json.loads(r.read())["message"]["content"]

    def classify(self, filename, text, checklist):
        menu = "\n".join(f"{r['ma']}: {r['ten']}" for r in checklist)
        ans = self.chat(
            "Bạn phân loại tài liệu hồ sơ xây dựng. Chỉ trả lời đúng một mã trong danh sách, "
            "hoặc NONE nếu không thuộc mã nào.",
            f"Danh sách mã:\n{menu}\n\nTên file: {filename}\nNội dung đầu file:\n{text[:1500]}\n\nMã:")
        m = re.search(r"\b([A-F]\d{2})\b", ans or "")
        return m.group(1) if m else None

    def answer(self, question, passages, status_summary):
        ctx = "\n\n".join(f"[{fn}]\n{ch}" for _, fn, ch in passages)
        return self.chat(
            "Bạn là trợ lý kiểm toán hồ sơ thanh toán xây dựng. Chỉ trả lời dựa trên dữ liệu được cung cấp. "
            "Ghi rõ tên file làm nguồn trong ngoặc vuông. Nếu dữ liệu không có câu trả lời, nói rõ là không tìm thấy "
            "và gợi ý tài liệu cần yêu cầu bổ sung. Trả lời bằng tiếng Việt, ngắn gọn.",
            f"Tình trạng danh mục hồ sơ:\n{status_summary}\n\nCác đoạn trích:\n{ctx}\n\nCâu hỏi: {question}")


# ----------------------------------------------------------------- hỏi đáp
MISSING_Q = re.compile(r"thieu|chua co|con thieu|can bo sung|request|yeu cau", re.I)


def tra_loi(question, kho, checklist, llm=None):
    """Câu hỏi về thiếu/đủ trả lời từ bảng tình trạng (chắc chắn).
    Câu hỏi khác: tìm đoạn liên quan, nếu có local AI thì nhờ AI tóm tắt."""
    st = tinh_trang(checklist, kho.index)
    q = bo_dau(question)
    if MISSING_Q.search(q):
        miss = [r for r in st if r["tinh_trang"] in ("Thiếu", "Cần xác nhận")]
        grp = None
        m = re.search(r"\b(nhom|muc)\s+([a-f])\b", q)
        if m:
            grp = m.group(2).upper()
            miss = [r for r in miss if r["ma"].startswith(grp)]
        if not miss:
            return "Không còn tài liệu bắt buộc nào bị thiếu" + (f" trong nhóm {grp}." if grp else "."), []
        lines = [f"- **{r['ma']}** {r['ten']} ({r['tinh_trang']}, ưu tiên {r['uu_tien']}) - dùng cho: {r['muc_dich']}"
                 for r in sorted(miss, key=lambda r: (r['uu_tien'], r['ma']))]
        return f"Còn {len(miss)} tài liệu thiếu hoặc cần xác nhận:\n" + "\n".join(lines), []
    code = re.search(r"\b([a-f]\d{2})\b", q)
    if code:
        r = next((x for x in st if x["ma"] == code.group(1).upper()), None)
        if r:
            return (f"**{r['ma']} {r['ten']}**: {r['tinh_trang']}. File: {r['file'] or 'chưa có'}. "
                    f"Dùng để: {r['muc_dich']} ({r['mat_xich']})."), []
    hits = search(kho, question)
    if not hits:
        return "Không tìm thấy nội dung liên quan trong các file đã upload.", []
    if llm and llm.ok():
        summ = "; ".join(f"{r['ma']}={r['tinh_trang']}" for r in st)
        try:
            return llm.answer(question, hits, summ), hits
        except Exception as e:
            return f"Local AI lỗi ({e}). Các đoạn liên quan nhất ở bên dưới.", hits
    return "Chưa bật local AI. Các đoạn liên quan nhất:", hits


def xuat_request(status_rows):
    """Danh sách yêu cầu bổ sung dạng Markdown."""
    miss = [r for r in status_rows if r["tinh_trang"] in ("Thiếu", "Cần xác nhận")]
    miss.sort(key=lambda r: (r["uu_tien"], r["ma"]))
    out = ["# Danh sách tài liệu đề nghị bổ sung", "",
           "| Mã | Tài liệu | Tình trạng | Ưu tiên | Để chứng minh | Liên quan |", "|---|---|---|---|---|---|"]
    for r in miss:
        out.append(f"| {r['ma']} | {r['ten']} | {r['tinh_trang']} | {r['uu_tien']} | {r['muc_dich']} | {r['mat_xich']} |")
    return "\n".join(out) + "\n"
