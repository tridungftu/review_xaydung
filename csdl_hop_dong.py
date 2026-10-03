# -*- coding: utf-8 -*-
"""
csdl_hop_dong.py - Cơ sở dữ liệu hợp đồng (SQLite) cho review hồ sơ thanh toán xây dựng.

NGUYÊN TẮC
- Số liệu gốc của hợp đồng (tham số, BOQ) không bao giờ bị sửa. Phụ lục là bản ghi thay đổi riêng.
- "Hiệu lực tại ngày X" = gốc + các phụ lục ĐÃ KÝ có ngày hiệu lực <= X.
  Mốc X khi kiểm tra một đợt thanh toán = NGÀY NGHIỆM THU của đợt đó.
- BOQ lưu dạng dòng: dòng GOC của hợp đồng + dòng THAY / THEM / BO của từng phụ lục.
  Mã hạng mục (HMxxxx) cố định suốt đời hợp đồng.
- Bảng dữ liệu gốc bị trigger chặn sửa/xoá; sửa sai phải qua lệnh có lý do, ghi nhật ký.
- Excel chỉ là phiếu nhập (phiếu phụ lục) và bản xuất để theo dõi; nguồn duy nhất là file .db.

LỆNH (HĐ có thể gõ mã HD0001, số HĐ đầy đủ, hoặc một đoạn duy nhất như 0404)
  python csdl_hop_dong.py nap-master   master_HD_0404.xlsx
  python csdl_hop_dong.py nap-danh-muc MAU_DANH_MUC_HOP_DONG.xlsx
  python csdl_hop_dong.py tao-phieu    0404                    -> phieu_phu_luc_<HĐ>.xlsx
  python csdl_hop_dong.py nap-phu-luc  phieu_phu_luc.xlsx      (mỗi sheet "PL..." là một phụ lục)
  python csdl_hop_dong.py huy-phu-luc  0404 PL01 --ly-do "nhập sai, nạp lại"
  python csdl_hop_dong.py nap-dot      0404 "HSTT dot 2.xlsx" --dot 2 --ngay-nghiem-thu 15/07/2026
  python csdl_hop_dong.py kiem-dot     0404 "HSTT dot 3.xlsx" --ngay-nghiem-thu 20/08/2026 [--dot 3]
  python csdl_hop_dong.py ngay-nghiem-thu 0404 2 15/07/2026     (bổ sung ngày NT cho đợt nạp từ master)
  python csdl_hop_dong.py boq          0404 [--ngay 30/06/2026] (in BOQ hiệu lực)
  python csdl_hop_dong.py xuat-excel   [0404] [--ngay 30/09/2026]
Tuỳ chọn chung: --db <đường dẫn .db> (mặc định: co_so_du_lieu/hop_dong.db cạnh file này)
"""
import argparse
import datetime as dt
import difflib
import re
import sqlite3
import sys
from contextlib import contextmanager
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from master_hd import (BASE, BOLD, GREEN, NUMF, QF, RED, YELLOW, bo_dau, header, khoa, nfc, num, put,
                       tim_bang_chi_tiet, tim_so_tong_hop)

DB_MAC_DINH = Path(__file__).resolve().parent / "co_so_du_lieu" / "hop_dong.db"
TRANG_THAI_PL = ("Đã ký", "Dự thảo", "Đã hủy")
LOAI_PL = ("Điều chỉnh khối lượng", "Điều chỉnh đơn giá", "Bổ sung hạng mục", "Gia hạn tiến độ",
           "Sửa điều khoản", "Kết hợp")
THAO_TAC = ("THAY", "THEM", "BO")

# 36 tham số chuẩn, đúng thứ tự cột của DANH_MUC_HD (MAU_DANH_MUC_HOP_DONG.xlsx)
DS_THAM_SO = [
    ("THÔNG TIN CHUNG", [("TC-01", "Số hợp đồng"), ("TC-02", "Ngày ký"), ("TC-03", "Gói thầu"), ("TC-04", "Dự án"),
                         ("TC-05", "Địa điểm"), ("TC-06", "Bên A (giao thầu)"), ("TC-07", "Bên B (nhận thầu)"),
                         ("TC-08", "Hiệu lực hợp đồng")]),
    ("GIÁ HỢP ĐỒNG", [("TT-01", "Loại hợp đồng"), ("TT-02", "Giá trị hợp đồng"), ("TT-03", "Giá trị HĐ đã gồm VAT"),
                      ("TT-04", "Thuế suất GTGT"), ("TT-05", "Cơ sở thanh toán khối lượng")]),
    ("TẠM ỨNG", [("TT-10", "Tỷ lệ tạm ứng"), ("TT-11", "Thu hồi tạm ứng"), ("TT-12", "Bảo lãnh tạm ứng")]),
    ("THANH TOÁN ĐỢT", [("TT-20", "Tỷ lệ thanh toán mỗi đợt"), ("TT-21", "Tỷ lệ giữ lại mỗi đợt"),
                        ("TT-22", "Thời hạn thanh toán")]),
    ("QUYẾT TOÁN – BẢO HÀNH", [("TT-30", "Tỷ lệ thanh toán khi quyết toán"), ("TT-31", "Tỷ lệ bảo hành / BL bảo hành"),
                               ("TT-32", "Thời hạn bảo hành"), ("TT-33", "Bảo lãnh thực hiện hợp đồng")]),
    ("ĐIỀU CHỈNH – PHÁT SINH", [("TT-40", "Điều chỉnh đơn giá"), ("TT-41", "Ngưỡng điều chỉnh giá"),
                                ("TT-42", "Xử lý khối lượng phát sinh")]),
    ("TIẾN ĐỘ – PHẠT", [("TT-50", "Tiến độ thực hiện"), ("TT-51", "Phạt vi phạm chất lượng (tối đa)"),
                        ("TT-52", "Tạm giữ chậm tiến độ - mức 1"), ("TT-53", "Tạm giữ chậm tiến độ - mức 2"),
                        ("TT-54", "Phạt chậm tổng tiến độ"), ("TT-55", "Trần phạt tiến độ"),
                        ("TT-56", "Khấu trừ tiền phạt vào thanh toán")]),
    ("HỒ SƠ THEO HĐ", [("TT-60", "Hồ sơ thanh toán mỗi đợt"), ("TT-61", "Hồ sơ quyết toán"),
                       ("TT-62", "Số bộ hồ sơ thanh toán")]),
]
MA_TY_LE = {"TT-04", "TT-10", "TT-20", "TT-21", "TT-30", "TT-31", "TT-41", "TT-51", "TT-52", "TT-53", "TT-54", "TT-55"}
GRAY = Font(name="Arial", size=9, italic=True, color="6B6B6B")
TITLE = Font(name="Arial", size=12, bold=True)

# ============================================================== cấu trúc CSDL
SCHEMA = """
PRAGMA foreign_keys = ON;
CREATE TABLE IF NOT EXISTS che_do(id INTEGER PRIMARY KEY CHECK(id = 1), sua INTEGER NOT NULL DEFAULT 0);
INSERT OR IGNORE INTO che_do VALUES (1, 0);

CREATE TABLE IF NOT EXISTS hop_dong(
    ma_hd TEXT PRIMARY KEY,               -- mã nội bộ HD0001, dùng làm khoá
    so_hd TEXT NOT NULL UNIQUE,           -- số hợp đồng theo văn bản (TC-01)
    phan_loai TEXT,                       -- loại HĐ / hạng mục (cột B danh mục)
    loai_gia TEXT,                        -- TT-01: trọn gói / đơn giá cố định / đơn giá điều chỉnh ...
    ngay_ky TEXT,                         -- ISO yyyy-mm-dd
    goi_thau TEXT,
    gia_tri_hd REAL,                      -- TT-02 theo văn bản HĐ gốc
    gom_vat TEXT,                         -- TT-03
    vat REAL,                             -- TT-04 (0.08)
    file_nguon TEXT,
    ngay_nap TEXT);

CREATE TABLE IF NOT EXISTS tham_so(        -- giá trị GỐC theo hợp đồng, mã theo MAP_THAM_SO (TC-xx, TT-xx, Y-xx)
    ma_hd TEXT NOT NULL REFERENCES hop_dong(ma_hd),
    ma_ts TEXT NOT NULL, ten TEXT, gia_tri TEXT, don_vi TEXT, dieu_khoan TEXT, trich_dan TEXT, trang_thai TEXT,
    PRIMARY KEY (ma_hd, ma_ts));

CREATE TABLE IF NOT EXISTS phu_luc(
    id INTEGER PRIMARY KEY,
    ma_hd TEXT NOT NULL REFERENCES hop_dong(ma_hd),
    so_pl TEXT NOT NULL,
    ngay_ky TEXT NOT NULL,
    ngay_hieu_luc TEXT NOT NULL,          -- mặc định = ngày ký
    loai TEXT,
    doi_boq INTEGER NOT NULL,             -- 1 = phụ lục có đổi BOQ
    gt_truoc_vat REAL,                    -- giá trị tăng (+) / giảm (-) ghi trên phụ lục
    gt_sau_vat REAL,
    gt_hd_sau_pl REAL,                    -- giá trị HĐ sau phụ lục ghi trên phụ lục
    noi_dung TEXT, can_cu TEXT, file_nguon TEXT, phieu TEXT,
    trang_thai TEXT NOT NULL,             -- Đã ký / Dự thảo / Đã hủy
    nguoi_nhap TEXT, ngay_nap TEXT, ly_do_huy TEXT);
CREATE UNIQUE INDEX IF NOT EXISTS ux_phu_luc ON phu_luc(ma_hd, so_pl) WHERE trang_thai <> 'Đã hủy';

CREATE TABLE IF NOT EXISTS thay_doi_tham_so(
    pl_id INTEGER NOT NULL REFERENCES phu_luc(id),
    ma_ts TEXT NOT NULL, gia_tri_cu TEXT, gia_tri_moi TEXT, trich_dan TEXT,
    PRIMARY KEY (pl_id, ma_ts));

CREATE TABLE IF NOT EXISTS boq_dong(
    id INTEGER PRIMARY KEY,
    ma_hd TEXT NOT NULL REFERENCES hop_dong(ma_hd),
    pl_id INTEGER REFERENCES phu_luc(id),  -- NULL = BOQ gốc của hợp đồng
    thao_tac TEXT NOT NULL CHECK (thao_tac IN ('GOC', 'THAY', 'THEM', 'BO')),
    ma_hm TEXT NOT NULL, stt TEXT, nhom TEXT, ten TEXT, dvt TEXT,
    kl REAL, don_gia REAL,                -- đơn giá chưa VAT
    thanh_tien_ghi REAL,                  -- thành tiền ghi trên tài liệu (để bắt lỗi số học)
    khoa TEXT,                            -- khoá so khớp nhóm|tên|ĐVT không dấu
    vi_tri_nguon TEXT, ghi_chu TEXT);
CREATE UNIQUE INDEX IF NOT EXISTS ux_boq ON boq_dong(ma_hd, IFNULL(pl_id, 0), ma_hm);

CREATE TABLE IF NOT EXISTS dot_thanh_toan(
    ma_hd TEXT NOT NULL REFERENCES hop_dong(ma_hd), dot INTEGER NOT NULL,
    ngay_nghiem_thu TEXT, ngay_nhan_hs TEXT,
    gt_ky REAL, gt_luy_ke REAL, giu_lai REAL, thu_hoi_tu REAL, gt_tt_ky REAL,
    file_nguon TEXT, ngay_nap TEXT,
    PRIMARY KEY (ma_hd, dot));

CREATE TABLE IF NOT EXISTS kl_thanh_toan(
    ma_hd TEXT NOT NULL, dot INTEGER NOT NULL, ma_hm TEXT NOT NULL,
    kl_ky REAL, kl_luy_ke REAL, don_gia_tt REAL, vi_tri_nguon TEXT,
    PRIMARY KEY (ma_hd, dot, ma_hm),
    FOREIGN KEY (ma_hd, dot) REFERENCES dot_thanh_toan(ma_hd, dot));

CREATE TABLE IF NOT EXISTS khop_hang_muc(   -- dòng trong hồ sơ TT -> Mã HM, nhớ cho các đợt sau
    ma_hd TEXT NOT NULL, khoa_hstt TEXT NOT NULL, ma_hm TEXT NOT NULL,
    cach_khop TEXT, nguoi_xac_nhan TEXT, ngay TEXT,
    PRIMARY KEY (ma_hd, khoa_hstt));

CREATE TABLE IF NOT EXISTS canh_bao(
    id INTEGER PRIMARY KEY, ma_hd TEXT, pl_id INTEGER, dot INTEGER,
    ma_kt TEXT, muc_do TEXT, noi_dung TEXT, thoi_diem TEXT);

CREATE TABLE IF NOT EXISTS nhat_ky(
    id INTEGER PRIMARY KEY, thoi_diem TEXT, nguoi TEXT, thao_tac TEXT, ma_hd TEXT, chi_tiet TEXT);
"""

# bảng gốc: chặn UPDATE/DELETE trừ khi đang ở chế độ sửa có lý do
KHOA_BANG = ["hop_dong", "tham_so", "thay_doi_tham_so", "boq_dong", "kl_thanh_toan", "dot_thanh_toan", "khop_hang_muc"]
PL_COT_DU_LIEU = "ma_hd, so_pl, ngay_ky, ngay_hieu_luc, loai, doi_boq, gt_truoc_vat, gt_sau_vat, gt_hd_sau_pl"


def _triggers():
    sql = []
    for t in KHOA_BANG:
        for ev in ("UPDATE", "DELETE"):
            sql.append(f"CREATE TRIGGER IF NOT EXISTS khoa_{t}_{ev.lower()} BEFORE {ev} ON {t} "
                       f"WHEN (SELECT sua FROM che_do) = 0 BEGIN "
                       f"SELECT RAISE(ABORT, 'Du lieu goc khong duoc sua/xoa ({t}). Dung lenh co --ly-do.'); END;")
    sql.append(f"CREATE TRIGGER IF NOT EXISTS khoa_phu_luc_update BEFORE UPDATE OF {PL_COT_DU_LIEU} ON phu_luc "
               "WHEN (SELECT sua FROM che_do) = 0 BEGIN "
               "SELECT RAISE(ABORT, 'Noi dung phu luc khong duoc sua. Huy phu luc roi nap lai.'); END;")
    sql.append("CREATE TRIGGER IF NOT EXISTS khoa_phu_luc_delete BEFORE DELETE ON phu_luc "
               "WHEN (SELECT sua FROM che_do) = 0 BEGIN SELECT RAISE(ABORT, 'Khong xoa phu luc. Dung huy-phu-luc.'); END;")
    for ev in ("UPDATE", "DELETE"):
        sql.append(f"CREATE TRIGGER IF NOT EXISTS khoa_nhat_ky_{ev.lower()} BEFORE {ev} ON nhat_ky "
                   "BEGIN SELECT RAISE(ABORT, 'Nhat ky khong duoc sua/xoa.'); END;")
    return "\n".join(sql)


def mo_csdl(path=None):
    path = Path(path or DB_MAC_DINH)
    path.parent.mkdir(parents=True, exist_ok=True)
    con = sqlite3.connect(str(path))
    con.row_factory = sqlite3.Row
    con.executescript(SCHEMA)
    con.executescript(_triggers())
    con.execute("PRAGMA foreign_keys = ON")
    return con


def bay_gio():
    return dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def ghi_nhat_ky(con, thao_tac, ma_hd="", chi_tiet="", nguoi=""):
    con.execute("INSERT INTO nhat_ky(thoi_diem, nguoi, thao_tac, ma_hd, chi_tiet) VALUES (?,?,?,?,?)",
                (bay_gio(), nguoi or _nguoi(), thao_tac, ma_hd, chi_tiet))


def _nguoi():
    import getpass
    try:
        return getpass.getuser()
    except Exception:
        return ""


@contextmanager
def che_do_sua(con, ly_do, ma_hd=""):
    if not (ly_do or "").strip():
        raise SystemExit("Thao tác này sửa dữ liệu gốc, phải ghi --ly-do.")
    ghi_nhat_ky(con, "MỞ CHẾ ĐỘ SỬA", ma_hd, ly_do)
    con.execute("UPDATE che_do SET sua = 1")
    try:
        yield
    finally:
        con.execute("UPDATE che_do SET sua = 0")


# ============================================================== tiện ích
def ngay_iso(v):
    """dd/mm/yyyy, yyyy-mm-dd, datetime -> 'yyyy-mm-dd'; rỗng -> None."""
    if v in (None, ""):
        return None
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    t = str(v).strip()
    m = re.fullmatch(r"(\d{1,2})[/.-](\d{1,2})[/.-](\d{4})", t)
    if m:
        return dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1))).isoformat()
    m = re.fullmatch(r"(\d{4})-(\d{2})-(\d{2})(.*)", t)
    if m:
        return f"{m.group(1)}-{m.group(2)}-{m.group(3)}"
    raise ValueError(f"Không hiểu ngày: {v!r} (dùng dd/mm/yyyy)")


def ngay_vn(iso):
    if not iso:
        return ""
    y, m, d = iso[:10].split("-")
    return f"{d}/{m}/{y}"


def ty_le(v):
    """'0.08' / 0.08 / '8%' -> 0.08"""
    if v in (None, ""):
        return None
    if isinstance(v, (int, float)):
        return float(v) / 100 if v > 1 else float(v)
    t = str(v).strip().replace(",", ".")
    m = re.search(r"([\d.]+)\s*%", t)
    if m:
        return float(m.group(1)) / 100
    try:
        f = float(t)
        return f / 100 if f > 1 else f
    except ValueError:
        return None


def tien(v):
    """Số tiền: 1234567890 / '1.234.567.890' / '-1.234.000' -> float."""
    if v in (None, ""):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    t = str(v).strip().replace(" ", "")
    am = t.startswith("-") or (t.startswith("(") and t.endswith(")"))
    t = t.strip("-()")
    if re.fullmatch(r"\d{1,3}([.,]\d{3})+", t):
        f = float(re.sub(r"[.,]", "", t))
    else:
        f = num(t)
        if f is None:
            return None
    return -f if am else f


def loai_gia_chuan(s):
    b = bo_dau(s)
    if "tron goi" in b:
        return "trọn gói"
    if "dieu chinh" in b:
        return "đơn giá điều chỉnh"
    if "co dinh" in b:
        return "đơn giá cố định"
    return s or ""


def tim_hd(con, ref):
    ref = nfc(ref).strip()
    r = con.execute("SELECT * FROM hop_dong WHERE ma_hd = ? OR so_hd = ?", (ref.upper(), ref)).fetchone()
    if r:
        return r
    rows = con.execute("SELECT * FROM hop_dong WHERE so_hd LIKE ?", (f"%{ref}%",)).fetchall()
    if len(rows) == 1:
        return rows[0]
    if not rows:
        raise SystemExit(f"Không có hợp đồng nào khớp '{ref}'.")
    raise SystemExit(f"'{ref}' khớp nhiều hợp đồng: " + ", ".join(f"{x['ma_hd']} ({x['so_hd']})" for x in rows))


def ma_hd_moi(con):
    n = con.execute("SELECT COUNT(*) FROM hop_dong").fetchone()[0]
    while True:
        n += 1
        ma = f"HD{n:04d}"
        if not con.execute("SELECT 1 FROM hop_dong WHERE ma_hd = ?", (ma,)).fetchone():
            return ma


def ma_hm_moi(con, ma_hd, dang_dung=()):
    cur = [r[0] for r in con.execute("SELECT ma_hm FROM boq_dong WHERE ma_hd = ?", (ma_hd,))] + list(dang_dung)
    n = max([int(m[2:]) for m in cur if re.fullmatch(r"HM\d+", m or "")] or [0])
    return f"HM{n + 1:04d}"


def gt_sau_vat_goc(hd):
    g = hd["gia_tri_hd"] or 0
    if bo_dau(hd["gom_vat"]) in ("khong", "chua"):
        return g * (1 + (hd["vat"] or 0))
    return g


def tol_tien(gt):
    return max(1000.0, abs(gt or 0) * 1e-5)


# ============================================================== truy vấn "hiệu lực"
def phu_luc_ap_dung(con, ma_hd, ngay=None, tru_pl_id=None):
    """Phụ lục đã ký, ngày hiệu lực <= ngay (None = tất cả), theo thứ tự hiệu lực."""
    q = "SELECT * FROM phu_luc WHERE ma_hd = ? AND trang_thai = 'Đã ký'"
    a = [ma_hd]
    if ngay:
        q += " AND ngay_hieu_luc <= ?"
        a.append(ngay)
    if tru_pl_id:
        q += " AND id <> ?"
        a.append(tru_pl_id)
    return con.execute(q + " ORDER BY ngay_hieu_luc, ngay_ky, so_pl", a).fetchall()


def boq_hieu_luc(con, ma_hd, ngay=None, pls=None):
    """
    BOQ hiệu lực tại ngày: BOQ gốc + lần lượt các phụ lục đã ký có hiệu lực <= ngày.
    Trả về (dang_hieu_luc: dict ma_hm -> dòng, da_bo: dict ma_hm -> dòng (kèm 'bo_boi')).
    Mỗi dòng có 'nguon' (HĐ gốc / số PL cuối cùng tác động) và 'lich_su' (chuỗi thay đổi).
    """
    cols = ("ma_hm", "stt", "nhom", "ten", "dvt", "kl", "don_gia", "thanh_tien_ghi", "khoa", "vi_tri_nguon")
    hl, bo = {}, {}
    for r in con.execute("SELECT * FROM boq_dong WHERE ma_hd = ? AND pl_id IS NULL ORDER BY id", (ma_hd,)):
        d = {c: r[c] for c in cols}
        d.update(nguon="HĐ gốc", lich_su=[f"Gốc: KL {_f(r['kl'])} × {_f(r['don_gia'])}"])
        hl[r["ma_hm"]] = d
    for pl in (pls if pls is not None else phu_luc_ap_dung(con, ma_hd, ngay)):
        for r in con.execute("SELECT * FROM boq_dong WHERE pl_id = ? ORDER BY id", (pl["id"],)):
            ma = r["ma_hm"]
            if r["thao_tac"] == "THEM":
                d = {c: r[c] for c in cols}
                d.update(nguon=pl["so_pl"], lich_su=[f"{pl['so_pl']} thêm: KL {_f(r['kl'])} × {_f(r['don_gia'])}"])
                hl[ma] = d
            elif r["thao_tac"] == "THAY" and ma in hl:
                d = dict(hl[ma])
                for c in ("stt", "nhom", "ten", "dvt", "kl", "don_gia", "thanh_tien_ghi"):
                    if r[c] not in (None, ""):
                        d[c] = r[c]
                d["nguon"] = pl["so_pl"]
                d["lich_su"] = hl[ma]["lich_su"] + [f"{pl['so_pl']} thay: KL {_f(d['kl'])} × {_f(d['don_gia'])}"]
                hl[ma] = d
            elif r["thao_tac"] == "BO" and ma in hl:
                d = hl.pop(ma)
                d["bo_boi"] = pl["so_pl"]
                bo[ma] = d
    return hl, bo


def _f(x):
    if x is None:
        return "-"
    return f"{x:,.3f}".rstrip("0").rstrip(".") if abs(x - round(x)) > 1e-9 else f"{x:,.0f}"


def tham_so_hieu_luc(con, ma_hd, ma_ts, ngay=None):
    """(giá trị, nguồn): giá trị gốc hoặc giá trị phụ lục gần nhất có hiệu lực <= ngày."""
    gia_tri, nguon = None, "HĐ gốc"
    r = con.execute("SELECT gia_tri FROM tham_so WHERE ma_hd = ? AND ma_ts = ?", (ma_hd, ma_ts)).fetchone()
    if r:
        gia_tri = r[0]
    for pl in phu_luc_ap_dung(con, ma_hd, ngay):
        t = con.execute("SELECT gia_tri_moi FROM thay_doi_tham_so WHERE pl_id = ? AND ma_ts = ?",
                        (pl["id"], ma_ts)).fetchone()
        if t:
            gia_tri, nguon = t[0], pl["so_pl"]
    return gia_tri, nguon


def lich_su_tham_so(con, ma_hd, ngay=None):
    """ma_ts -> list[(so_pl, ngay_hieu_luc, cu, moi, trich_dan)] của các phụ lục áp dụng."""
    out = {}
    for pl in phu_luc_ap_dung(con, ma_hd, ngay):
        for t in con.execute("SELECT * FROM thay_doi_tham_so WHERE pl_id = ?", (pl["id"],)):
            out.setdefault(t["ma_ts"], []).append((pl["so_pl"], pl["ngay_hieu_luc"], t["gia_tri_cu"], t["gia_tri_moi"], t["trich_dan"]))
    return out


def gt_hd_hieu_luc(con, ma_hd, ngay=None, pls=None):
    """GT HĐ sau VAT = gốc + Σ tăng/giảm sau VAT của các phụ lục đã ký có hiệu lực <= ngày."""
    hd = con.execute("SELECT * FROM hop_dong WHERE ma_hd = ?", (ma_hd,)).fetchone()
    g = gt_sau_vat_goc(hd)
    for pl in (pls if pls is not None else phu_luc_ap_dung(con, ma_hd, ngay)):
        g += pl["gt_sau_vat"] or 0
    return g


def kl_da_thanh_toan(con, ma_hd):
    """ma_hm -> (đợt mới nhất, KL lũy kế, ngày nghiệm thu)."""
    out = {}
    q = ("SELECT k.ma_hm, k.dot, k.kl_luy_ke, d.ngay_nghiem_thu FROM kl_thanh_toan k "
         "JOIN dot_thanh_toan d ON d.ma_hd = k.ma_hd AND d.dot = k.dot WHERE k.ma_hd = ? ORDER BY k.dot")
    for r in con.execute(q, (ma_hd,)):
        out[r["ma_hm"]] = (r["dot"], r["kl_luy_ke"] or 0, r["ngay_nghiem_thu"])
    return out


# ============================================================== nạp hợp đồng
def _them_hop_dong(con, so_hd, ts, phan_loai="", file_nguon=""):
    """ts: dict ma_ts -> dict(ten, gia_tri, don_vi, dieu_khoan, trich_dan, trang_thai)."""
    ma = ma_hd_moi(con)
    g = lambda k: (ts.get(k) or {}).get("gia_tri")
    con.execute("INSERT INTO hop_dong VALUES (?,?,?,?,?,?,?,?,?,?,?)",
                (ma, so_hd, phan_loai, loai_gia_chuan(g("TT-01")), _ngay_an_toan(g("TC-02")), g("TC-03"),
                 tien(g("TT-02")), g("TT-03"), ty_le(g("TT-04")), file_nguon, bay_gio()))
    for k, v in ts.items():
        con.execute("INSERT INTO tham_so VALUES (?,?,?,?,?,?,?,?)",
                    (ma, k, v.get("ten"), None if v.get("gia_tri") is None else str(v.get("gia_tri")),
                     v.get("don_vi"), v.get("dieu_khoan"), v.get("trich_dan"), v.get("trang_thai")))
    return ma


def _ngay_an_toan(v):
    try:
        return ngay_iso(v)
    except ValueError:
        return None


def nap_master(con, path, lam_lai=False, ly_do=""):
    """Nạp master_HD_xxx.xlsx (do master_hd.py tạo): tham số, BOQ gốc, lịch sử TT, KL lũy kế."""
    wb = openpyxl.load_workbook(path, data_only=False)
    ts = {}
    for sh in ("Thong tin HD", "Dieu khoan TT"):
        for row in wb[sh].iter_rows(min_row=2, values_only=True):
            if row[0] and re.fullmatch(r"(TC|TT|Y|BS)-\d+", str(row[0])):
                ts[row[0]] = dict(ten=row[2], gia_tri=row[3], don_vi=row[4], dieu_khoan=row[6],
                                  trich_dan=row[7], trang_thai=row[8])
    so_hd = nfc((ts.get("TC-01") or {}).get("gia_tri") or "").strip()
    if not so_hd:
        raise SystemExit("Master chưa có số hợp đồng (TC-01).")
    cu = con.execute("SELECT ma_hd FROM hop_dong WHERE so_hd = ?", (so_hd,)).fetchone()
    if cu and not lam_lai:
        raise SystemExit(f"Hợp đồng {so_hd} đã có trong CSDL ({cu[0]}). Muốn nạp lại BOQ gốc: thêm --lam-lai --ly-do \"...\"")
    if cu:
        _xoa_hd(con, cu[0], ly_do)
    ma = _them_hop_dong(con, so_hd, ts, file_nguon=Path(path).name)
    # BOQ gốc
    n = 0
    for i, row in enumerate(wb["BOQ"].iter_rows(min_row=2, values_only=True), 2):
        if not (row[0] and str(row[0]).startswith("HM")):
            continue
        it = dict(nhom=row[2] or "", ten=nfc(row[3]), dvt=nfc(row[4]))
        con.execute("INSERT INTO boq_dong(ma_hd, pl_id, thao_tac, ma_hm, stt, nhom, ten, dvt, kl, don_gia, "
                    "thanh_tien_ghi, khoa, vi_tri_nguon, ghi_chu) VALUES (?,NULL,'GOC',?,?,?,?,?,?,?,?,?,?,?)",
                    (ma, row[0], None if row[1] is None else str(row[1]), it["nhom"], it["ten"], it["dvt"],
                     num(row[5]), num(row[6]), num(row[7]) if not str(row[7] or "").startswith("=") else None,
                     row[10] if len(row) > 10 and row[10] else khoa(it),
                     f"{Path(path).name} BOQ!dòng {i} (nguồn gốc: {row[8] or ''})", row[9]))
        n += 1
    # lịch sử thanh toán
    nd = 0
    for row in wb["Lich su TT"].iter_rows(min_row=2, values_only=True):
        if row[0] in (None, ""):
            continue
        con.execute("INSERT INTO dot_thanh_toan VALUES (?,?,?,?,?,?,?,?,?,?,?)",
                    (ma, int(num(row[0])), None, None, num(row[2]), num(row[3]), num(row[4]), num(row[5]),
                     num(row[6]), f"{Path(path).name} Lich su TT: {row[9] or ''}", bay_gio()))
        nd += 1
    # KL lũy kế đến hết đợt gần nhất
    nk = 0
    for row in wb["KL ky truoc"].iter_rows(min_row=2, values_only=True):
        if not (row[0] and str(row[0]).startswith("HM")) or row[6] in (None, ""):
            continue
        dot = int(num(row[6]))
        if not con.execute("SELECT 1 FROM dot_thanh_toan WHERE ma_hd = ? AND dot = ?", (ma, dot)).fetchone():
            con.execute("INSERT INTO dot_thanh_toan(ma_hd, dot, file_nguon, ngay_nap) VALUES (?,?,?,?)",
                        (ma, dot, Path(path).name, bay_gio()))
        con.execute("INSERT INTO kl_thanh_toan VALUES (?,?,?,?,?,?,?)",
                    (ma, dot, row[0], None, num(row[3]), None, f"{Path(path).name} KL ky truoc: {row[7] or ''}"))
        nk += 1
    ghi_nhat_ky(con, "Nạp master", ma, f"{Path(path).name}: {len(ts)} tham số, {n} dòng BOQ gốc, {nd} đợt, {nk} dòng KL lũy kế")
    kt = kiem_tra_boq_goc(con, ma)
    con.commit()
    return ma, len(ts), n, nd, nk, kt


def kiem_tra_boq_goc(con, ma_hd):
    """BOQ gốc × (1+VAT) phải bằng GT hợp đồng; thành tiền ghi phải = KL × ĐG."""
    hd = con.execute("SELECT * FROM hop_dong WHERE ma_hd = ?", (ma_hd,)).fetchone()
    rows = con.execute("SELECT * FROM boq_dong WHERE ma_hd = ? AND pl_id IS NULL", (ma_hd,)).fetchall()
    out = []
    if not rows:
        out.append(("B00", "VÀNG", "Chưa có BOQ gốc: chưa đối chiếu được đơn giá/khối lượng."))
    else:
        tong = sum(round((r["kl"] or 0) * (r["don_gia"] or 0)) for r in rows) * (1 + (hd["vat"] or 0))
        gt = gt_sau_vat_goc(hd)
        if abs(tong - gt) > tol_tien(gt):
            out.append(("B01", "ĐỎ", f"BOQ gốc × (1+VAT) = {tong:,.0f} khác GT HĐ {gt:,.0f} (lệch {tong - gt:,.0f})."))
        for r in rows:
            if r["thanh_tien_ghi"] is not None and abs(r["thanh_tien_ghi"] - round((r["kl"] or 0) * (r["don_gia"] or 0))) > 1:
                out.append(("B02", "VÀNG", f"{r['ma_hm']}: thành tiền ghi {r['thanh_tien_ghi']:,.0f} ≠ KL × ĐG."))
    for ma_kt, md, nd in out:
        con.execute("INSERT INTO canh_bao(ma_hd, ma_kt, muc_do, noi_dung, thoi_diem) VALUES (?,?,?,?,?)",
                    (ma_hd, ma_kt, md, nd, bay_gio()))
    return out


def _xoa_hd(con, ma_hd, ly_do):
    if con.execute("SELECT 1 FROM phu_luc WHERE ma_hd = ? AND trang_thai <> 'Đã hủy'", (ma_hd,)).fetchone():
        raise SystemExit("Hợp đồng đã có phụ lục, không nạp lại BOQ gốc được. Huỷ phụ lục trước.")
    with che_do_sua(con, ly_do, ma_hd):
        for t in ("kl_thanh_toan", "dot_thanh_toan", "khop_hang_muc", "boq_dong", "tham_so", "hop_dong"):
            con.execute(f"DELETE FROM {t} WHERE ma_hd = ?", (ma_hd,))
    ghi_nhat_ky(con, "Xoá để nạp lại", ma_hd, ly_do)


def nap_danh_muc(con, path):
    """Nạp các dòng DANH_MUC_HD (mẫu MAU_DANH_MUC_HOP_DONG.xlsx). HĐ đã có thì chỉ so sánh, không ghi đè."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["DANH_MUC_HD"]
    ma_cot = {}
    for c in range(1, ws.max_column + 1):
        code = str(ws.cell(4, c).value or "").strip()
        if re.fullmatch(r"(TC|TT)-\d+", code) or re.fullmatch(r"(Y|BS)-\d+", code):
            ma_cot[code] = c
    ten_cot = {c: ws.cell(5, c).value for c in range(1, ws.max_column + 1)}
    ket_qua = []
    for r in range(6, ws.max_row + 1):
        so_hd = nfc(ws.cell(r, ma_cot.get("TC-01", 3)).value or "").strip()
        ghi_chu = " ".join(str(ws.cell(r, c).value or "") for c in range(1, ws.max_column + 1)
                           if "ghi chu" in bo_dau(ten_cot.get(c)))
        if not so_hd or "vi du" in bo_dau(ghi_chu) or so_hd.upper().startswith("VD-"):
            continue
        ts = {k: dict(ten=ten_cot.get(c), gia_tri=ws.cell(r, c).value, trang_thai="Từ danh mục")
              for k, c in ma_cot.items() if ws.cell(r, c).value not in (None, "")}
        cu = con.execute("SELECT ma_hd FROM hop_dong WHERE so_hd = ?", (so_hd,)).fetchone()
        if not cu:
            ma = _them_hop_dong(con, so_hd, ts, phan_loai=ws.cell(r, 2).value or "", file_nguon=f"{Path(path).name} dòng {r}")
            ghi_nhat_ky(con, "Nạp từ danh mục", ma, f"{Path(path).name} dòng {r}: {len(ts)} tham số")
            ket_qua.append((so_hd, ma, "Thêm mới", []))
        else:
            lech = []
            for k, v in ts.items():
                db = con.execute("SELECT gia_tri FROM tham_so WHERE ma_hd = ? AND ma_ts = ?", (cu[0], k)).fetchone()
                if db is None:
                    continue
                a, b = str(v["gia_tri"]).strip(), str(db[0] or "").strip()
                if bo_dau(a) != bo_dau(b) and not (num(a) is not None and num(b) is not None and abs(num(a) - num(b)) < 1e-6) \
                        and not (k == "TC-02" and _ngay_an_toan(a) == _ngay_an_toan(b)):
                    lech.append(f"{k}: danh mục '{a[:40]}' ≠ CSDL '{b[:40]}'")
            ket_qua.append((so_hd, cu[0], "Đã có - chỉ so sánh", lech))
    con.commit()
    return ket_qua


# ============================================================== phiếu phụ lục
NHAN = [  # (nhãn hiển thị, khoá, bắt buộc, ghi chú)
    ("Số hợp đồng", "so_hd", True, "Tự điền khi tạo phiếu"),
    ("Số phụ lục", "so_pl", True, "VD: PL01"),
    ("Ngày ký phụ lục", "ngay_ky", True, "dd/mm/yyyy"),
    ("Ngày hiệu lực", "ngay_hieu_luc", False, "Để trống = ngày ký. Ghi khác nếu PL quy định hiệu lực riêng"),
    ("Loại phụ lục", "loai", False, " / ".join(LOAI_PL)),
    ("Có đổi BOQ?", "doi_boq", True, "Có / Không. 'Có' thì điền mục B hoặc chỉ ra sheet BOQ kèm theo"),
    ("Giá trị tăng (+)/giảm (-) trước VAT", "gt_truoc_vat", False, "Số ghi trên phụ lục"),
    ("Giá trị tăng (+)/giảm (-) sau VAT", "gt_sau_vat", False, "Số ghi trên phụ lục; trống thì máy suy ra từ trước VAT"),
    ("Giá trị HĐ sau phụ lục (ghi trên PL)", "gt_hd_sau_pl", False, "Để máy kiểm tra nối chuỗi"),
    ("Nội dung tóm tắt", "noi_dung", False, ""),
    ("Căn cứ (điều khoản HĐ / lý do)", "can_cu", False, "Bắt buộc nếu đổi đơn giá"),
    ("File phụ lục (bản ký)", "file_nguon", False, "Tên file scan/PDF lưu trong hồ sơ"),
    ("Trạng thái", "trang_thai", True, " / ".join(TRANG_THAI_PL[:2])),
    ("Người nhập", "nguoi_nhap", False, ""),
    ("Sheet BOQ kèm theo", "sheet_boq", False, "Tuỳ chọn: tên sheet chứa bảng BOQ của PL (thay cho mục B)"),
    ("Phạm vi sheet BOQ", "pham_vi", False, "Toàn bộ BOQ sau PL / Chỉ hạng mục thay đổi"),
]
MUC_A = "A. THAY ĐỔI ĐIỀU KHOẢN / THAM SỐ (không ghi TT-02, máy tự tính)"
MUC_B = "B. THAY ĐỔI BOQ (chỉ điền khi 'Có đổi BOQ?' = Có)"
COT_A = ["Mã tham số", "Tên tham số", "Giá trị cũ", "Giá trị mới", "Trích dẫn nguyên văn phụ lục"]
COT_B = ["Thao tác", "Mã HM", "Nhóm / hạng mục", "Tên công việc", "ĐVT", "KL mới", "Đơn giá mới (chưa VAT)",
         "Thành tiền ghi trên PL", "Ghi chú"]


def tao_phieu(con, ref, out=None):
    hd = tim_hd(con, ref)
    ma = hd["ma_hd"]
    so = [int(m.group(1)) for (s,) in con.execute("SELECT so_pl FROM phu_luc WHERE ma_hd = ? AND trang_thai <> 'Đã hủy'", (ma,))
          for m in [re.search(r"(\d+)", s or "")] if m]
    so_pl = f"PL{(max(so) if so else 0) + 1:02d}"
    out = Path(out or f"phieu_phu_luc_{ma}_{so_pl}.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = so_pl
    _ve_phieu(ws, hd, so_pl)
    # sheet tham chiếu: BOQ hiện hành + tham số hiện hành
    hl, _ = boq_hieu_luc(con, ma)
    w2 = wb.create_sheet("BOQ hien hanh")
    put(w2, 1, 1, f"BOQ hiện hành của {hd['so_hd']} (gốc + các phụ lục đã ký) - để tra Mã HM, không sửa sheet này", BOLD)
    header(w2, 2, ["Mã HM", "STT", "Nhóm / hạng mục", "Tên công việc", "ĐVT", "KL", "Đơn giá (chưa VAT)", "Nguồn"],
           [9, 6, 34, 56, 8, 12, 15, 10])
    for i, d in enumerate(hl.values(), 3):
        for j, v in enumerate([d["ma_hm"], d["stt"], d["nhom"], d["ten"], d["dvt"], d["kl"], d["don_gia"], d["nguon"]], 1):
            put(w2, i, j, v, fmt=QF if j == 6 else (NUMF if j == 7 else None), text=True)
    w3 = wb.create_sheet("Tham so hien hanh")
    header(w3, 1, ["Mã", "Tham số", "Giá trị hiện hành", "Nguồn"], [8, 40, 50, 10])
    for i, r in enumerate(con.execute("SELECT ma_ts, ten FROM tham_so WHERE ma_hd = ? ORDER BY ma_ts", (ma,)), 2):
        v, ng = tham_so_hieu_luc(con, ma, r["ma_ts"])
        for j, x in enumerate([r["ma_ts"], r["ten"], v, ng], 1):
            put(w3, i, j, x, text=True)
    wh = wb.create_sheet("Huong dan")
    for i, t in enumerate([
        "CÁCH DÙNG PHIẾU PHỤ LỤC",
        "1. Mỗi phụ lục một sheet (tên sheet bắt đầu bằng PL). Nhiều phụ lục: copy sheet PL rồi đổi tên.",
        "2. Điền phần thông tin chung theo đúng văn bản phụ lục đã ký. Ô vàng là bắt buộc.",
        "3. Mục A: chỉ ghi tham số bị phụ lục sửa (VD TT-50 gia hạn tiến độ). Mã tham số tra ở sheet 'Tham so hien hanh'.",
        "   KHÔNG ghi TT-02 (giá trị HĐ): máy tự tính từ ô giá trị tăng/giảm ở trên, để không nhập hai lần.",
        "4. 'Có đổi BOQ?' = Có thì máy cập nhật BOQ theo mục B (hoặc theo sheet BOQ kèm theo):",
        "   THAY = sửa KL/đơn giá hạng mục đã có (ghi Mã HM, chỉ điền cột thay đổi, cột trống giữ nguyên)",
        "   THEM = hạng mục mới (để trống Mã HM, máy tự cấp; phải đủ tên, ĐVT, KL, đơn giá)",
        "   BO   = bỏ hạng mục (ghi Mã HM)",
        "5. Thay vì gõ mục B có thể dán bảng BOQ của phụ lục vào một sheet khác rồi ghi tên sheet đó ở ô 'Sheet BOQ kèm theo'.",
        "   Phạm vi 'Toàn bộ BOQ sau PL': hạng mục không còn trong bảng sẽ bị coi là BO.",
        "   Phạm vi 'Chỉ hạng mục thay đổi': chỉ THAY/THEM, không BO.",
        "6. Kéo thả file vào 7_nap_phu_luc.bat. Máy kiểm tra trước khi ghi; lỗi CHẶN thì không ghi gì.",
        "7. Phụ lục áp dụng cho đợt thanh toán có NGÀY NGHIỆM THU >= ngày hiệu lực phụ lục.",
        "8. Nạp rồi không sửa được. Nhập sai: huy-phu-luc (ghi lý do) rồi nạp lại.",
    ], 1):
        put(wh, i, 1, t, BOLD if i == 1 else BASE)
    wh.column_dimensions["A"].width = 120
    wb.save(out)
    return out, so_pl


def _ve_phieu(ws, hd, so_pl):
    put(ws, 1, 1, "PHIẾU NHẬP PHỤ LỤC HỢP ĐỒNG", TITLE)
    put(ws, 2, 1, f"{hd['ma_hd']} - {hd['so_hd']} - {hd['goi_thau'] or ''}", GRAY)
    for c, w in zip("ABCDEFGHI", [34, 30, 34, 50, 8, 12, 16, 16, 30]):
        ws.column_dimensions[c].width = w
    dv_cokhong = DataValidation(type="list", formula1='"Có,Không"', allow_blank=True)
    dv_tt = DataValidation(type="list", formula1='"Đã ký,Dự thảo"', allow_blank=True)
    dv_loai = DataValidation(type="list", formula1='"' + ",".join(LOAI_PL) + '"', allow_blank=True)
    dv_pv = DataValidation(type="list", formula1='"Toàn bộ BOQ sau PL,Chỉ hạng mục thay đổi"', allow_blank=True)
    dv_op = DataValidation(type="list", formula1='"THAY,THEM,BO"', allow_blank=True)
    for dv in (dv_cokhong, dv_tt, dv_loai, dv_pv, dv_op):
        ws.add_data_validation(dv)
    r = 4
    for nhan, k, bb, gc in NHAN:
        put(ws, r, 1, nhan, BOLD)
        v = {"so_hd": hd["so_hd"], "so_pl": so_pl, "trang_thai": "Đã ký", "doi_boq": "Không"}.get(k)
        put(ws, r, 2, v, fill=YELLOW if bb else None, text=True)
        put(ws, r, 3, gc, GRAY)
        if k in ("gt_truoc_vat", "gt_sau_vat", "gt_hd_sau_pl"):
            ws.cell(r, 2).number_format = NUMF
        dv = {"doi_boq": dv_cokhong, "trang_thai": dv_tt, "loai": dv_loai, "pham_vi": dv_pv}.get(k)
        if dv:
            dv.add(ws.cell(r, 2))
        r += 1
    r += 1
    put(ws, r, 1, MUC_A, BOLD)
    _hdr(ws, r + 1, COT_A)
    r += 2 + 8
    put(ws, r, 1, MUC_B, BOLD)
    _hdr(ws, r + 1, COT_B)
    dv_op.add(f"A{r + 2}:A{r + 300}")
    for rr in range(r + 2, r + 302):
        for c in (6, 7, 8):
            ws.cell(rr, c).number_format = QF if c == 6 else NUMF
    ws.freeze_panes = "A4"


def _hdr(ws, row, names):
    from master_hd import HF, HFONT
    for j, h in enumerate(names, 1):
        cl = ws.cell(row, j, h)
        cl.font, cl.fill = HFONT, HF


def doc_phieu(wb, ws):
    """Đọc một sheet phiếu -> (thong_tin dict, tham_so list, boq list, ghi_chu list)."""
    info, a_rows, b_rows, note = {}, [], [], []
    nhan = sorted(((bo_dau(n), k) for n, k, _, _ in NHAN), key=lambda x: -len(x[0]))
    muc = None
    r = 1
    while r <= ws.max_row:
        v0 = ws.cell(r, 1).value
        t = bo_dau(v0)
        if t.startswith(bo_dau(MUC_A)[:20]):
            muc, r = "A", r + 2
            continue
        if t.startswith(bo_dau(MUC_B)[:15]):
            muc, r = "B", r + 2
            continue
        if muc == "A" and t in ("thay", "them", "bo"):
            muc = "B"   # mất dòng tiêu đề mục B: nhận theo cột Thao tác
            note.append(("VÀNG", f"Sheet {ws.title}: không thấy dòng tiêu đề mục B, nhận mục B từ dòng {r}."))
        if muc is None:
            for n, k in nhan:
                if t and t == n:
                    info[k] = ws.cell(r, 2).value
                    break
        elif muc == "A":
            vals = [ws.cell(r, c).value for c in range(1, 6)]
            if any(x not in (None, "") for x in vals):
                a_rows.append(dict(dong=r, ma_ts=str(vals[0] or "").strip().upper(), ten=vals[1],
                                   cu=vals[2], moi=vals[3], trich=vals[4]))
        elif muc == "B":
            vals = [ws.cell(r, c).value for c in range(1, 10)]
            if any(x not in (None, "") for x in vals):
                b_rows.append(dict(dong=r, thao_tac=str(vals[0] or "").strip().upper(),
                                   ma_hm=str(vals[1] or "").strip().upper(), nhom=nfc(vals[2] or "").strip(),
                                   ten=nfc(vals[3] or "").strip(), dvt=nfc(vals[4] or "").strip(),
                                   kl=num(vals[5]), don_gia=num(vals[6]), thanh_tien_ghi=num(vals[7]), ghi_chu=vals[8],
                                   vi_tri=f"{ws.title}!dòng {r}"))
        r += 1
    # BOQ dán ở sheet khác
    sh = str(info.get("sheet_boq") or "").strip()
    if sh:
        if sh not in wb.sheetnames:
            note.append(("CHẶN", f"Không có sheet '{sh}' trong file."))
        else:
            info["_boq_sheet"] = doc_bang_boq(wb[sh])
    return info, a_rows, b_rows, note


def doc_bang_boq(ws):
    """Đọc bảng BOQ bất kỳ: cần cột tên công việc, ĐVT, khối lượng, đơn giá; tuỳ chọn Mã HM, thành tiền."""
    for r in range(1, min(ws.max_row, 30) + 1):
        hdr = {c: bo_dau(ws.cell(r, c).value) + " " + bo_dau(ws.cell(r + 1, c).value) for c in range(1, min(ws.max_column, 30) + 1)}
        cten = next((c for c, t in hdr.items() if re.search(r"^(ten cong (viec|tac)|noi dung|ten hang muc|hang muc cong viec)", t)), None)
        cdg = next((c for c, t in hdr.items() if re.search(r"^don gia", t)), None)
        cdvt = next((c for c, t in hdr.items() if re.search(r"^(don vi|dvt|d\.? ?vi)", t)), None)
        if not (cten and cdg and cdvt):
            continue
        kls = [c for c, t in hdr.items() if "khoi luong" in t and c < cdg] or [c for c, t in hdr.items() if t.startswith("kl")]
        if not kls:
            continue
        ckl = next((c for c in kls if re.search(r"sau|moi|dieu chinh|phu luc|pl", hdr[c])), kls[-1])
        ctt = next((c for c, t in hdr.items() if c > cdg and re.search(r"thanh tien|gia tri", t)), None)
        cma = next((c for c, t in hdr.items() if re.search(r"^ma (hm|hang muc)", t)), None)
        cstt = next((c for c, t in hdr.items() if c < cten and re.search(r"^(stt|tt)\b", t)), None)
        items, grp = [], ""
        for rr in range(r + 1, ws.max_row + 1):
            ten = ws.cell(rr, cten).value
            if not ten or re.fullmatch(r"[\[\]\-\d\s()]+", str(ten)):
                continue
            dvt = ws.cell(rr, cdvt).value
            low = nfc(ten).strip().lower()
            if dvt in (None, "") and re.match(r"^(cộng|tổng cộng|thuế|tổng giá trị)(\s|\(|:|$)", low):
                break
            if dvt in (None, ""):
                grp = nfc(ten).strip()
                continue
            items.append(dict(ma_hm=str(ws.cell(rr, cma).value or "").strip().upper() if cma else "",
                              stt=str(ws.cell(rr, cstt).value or "") if cstt else "", nhom=grp, ten=nfc(ten).strip(),
                              dvt=nfc(dvt).strip(), kl=num(ws.cell(rr, ckl).value), don_gia=num(ws.cell(rr, cdg).value),
                              thanh_tien_ghi=num(ws.cell(rr, ctt).value) if ctt else None, vi_tri=f"{ws.title}!dòng {rr}"))
        return dict(sheet=ws.title, items=items, cot_kl=ws.cell(r, ckl).value)
    return dict(sheet=ws.title, items=[], loi="Không tìm thấy tiêu đề (tên công việc, ĐVT, khối lượng, đơn giá).")


def _khop_ten(it, hl):
    """Tìm Mã HM cho một dòng BOQ phụ lục (không có mã): khoá đầy đủ -> tên + ĐVT duy nhất."""
    if it.get("ma_hm") and it["ma_hm"] in hl:
        return it["ma_hm"]
    k = khoa(it)
    for ma, d in hl.items():
        if d["khoa"] == k:
            return ma
    t = (bo_dau(it["ten"])[:120], bo_dau(it["dvt"]))
    cand = [ma for ma, d in hl.items() if (bo_dau(d["ten"])[:120], bo_dau(d["dvt"])) == t]
    return cand[0] if len(cand) == 1 else None


def dien_giai_boq_sheet(bs, pham_vi, hl):
    """Bảng BOQ phụ lục -> các dòng THAY/THEM/BO so với BOQ hiện hành."""
    out, seen = [], set()
    for it in bs["items"]:
        ma = _khop_ten(it, hl)
        if ma:
            seen.add(ma)
            d = hl[ma]
            if abs((it["kl"] or 0) - (d["kl"] or 0)) > 1e-9 or abs((it["don_gia"] or 0) - (d["don_gia"] or 0)) > 0.5:
                out.append(dict(it, thao_tac="THAY", ma_hm=ma, ghi_chu="Tự suy từ sheet BOQ"))
        else:
            out.append(dict(it, thao_tac="THEM", ma_hm="", ghi_chu="Tự suy từ sheet BOQ: không khớp hạng mục nào"))
    if bo_dau(pham_vi).startswith("toan bo"):
        for ma, d in hl.items():
            if ma not in seen:
                out.append(dict(thao_tac="BO", ma_hm=ma, nhom=d["nhom"], ten=d["ten"], dvt=d["dvt"], kl=None, don_gia=None,
                                thanh_tien_ghi=None, vi_tri=f"{bs['sheet']} (không có)", ghi_chu="Tự suy: không còn trong BOQ sau PL"))
    return out


def kiem_tra_phu_luc(con, hd, info, a_rows, b_rows, note):
    """Trả về (ket_qua list[(ma, muc_do, noi_dung)], pl_chuan dict, dong_boq list, dong_ts list)."""
    kq = list((("P00", md, nd) for md, nd in note))
    ma = hd["ma_hd"]
    vat = hd["vat"] or 0
    loai_gia = loai_gia_chuan(hd["loai_gia"])
    so_pl = nfc(info.get("so_pl") or "").strip().upper()
    try:
        ngay_ky = ngay_iso(info.get("ngay_ky"))
        ngay_hl = ngay_iso(info.get("ngay_hieu_luc")) or ngay_ky
    except ValueError as e:
        kq.append(("P01", "CHẶN", str(e)))
        ngay_ky = ngay_hl = None
    tt = nfc(info.get("trang_thai") or "").strip()
    doi = bo_dau(info.get("doi_boq"))
    if not so_pl:
        kq.append(("P01", "CHẶN", "Thiếu số phụ lục."))
    if not ngay_ky:
        kq.append(("P01", "CHẶN", "Thiếu ngày ký phụ lục."))
    if tt not in TRANG_THAI_PL[:2]:
        kq.append(("P01", "CHẶN", f"Trạng thái '{tt}' không hợp lệ (Đã ký / Dự thảo)."))
    if doi not in ("co", "khong"):
        kq.append(("P01", "CHẶN", "Ô 'Có đổi BOQ?' phải là Có hoặc Không."))
    if so_pl and con.execute("SELECT 1 FROM phu_luc WHERE ma_hd = ? AND so_pl = ? AND trang_thai <> 'Đã hủy'",
                             (ma, so_pl)).fetchone():
        kq.append(("P02", "CHẶN", f"{so_pl} đã có trong CSDL. Nhập lại thì huỷ bản cũ trước (huy-phu-luc)."))

    # BOQ hiện hành ngay trước phụ lục này
    truoc = [p for p in phu_luc_ap_dung(con, ma) if ngay_hl and (p["ngay_hieu_luc"], p["so_pl"]) <= (ngay_hl, so_pl)]
    hl, bo = boq_hieu_luc(con, ma, pls=truoc)
    gt_truoc_pl = gt_hd_hieu_luc(con, ma, pls=truoc)

    # dòng BOQ: từ mục B hoặc từ sheet kèm theo
    bs = info.get("_boq_sheet")
    if bs and b_rows:
        kq.append(("P03", "CHẶN", "Vừa điền mục B vừa chỉ sheet BOQ kèm theo. Chỉ dùng một cách."))
    if bs:
        if bs.get("loi"):
            kq.append(("P03", "CHẶN", f"Sheet {bs['sheet']}: {bs['loi']}"))
        elif not str(info.get("pham_vi") or "").strip():
            kq.append(("P03", "CHẶN", "Có sheet BOQ kèm theo thì phải chọn 'Phạm vi sheet BOQ'."))
        else:
            b_rows = dien_giai_boq_sheet(bs, info.get("pham_vi"), hl)
            kq.append(("P03", "THÔNG TIN", f"Đọc {len(bs['items'])} dòng từ sheet {bs['sheet']} (cột KL: {bs.get('cot_kl')}) "
                                          f"-> {len(b_rows)} dòng thay đổi."))
    if doi == "co" and not b_rows:
        kq.append(("P03", "CHẶN", "'Có đổi BOQ?' = Có nhưng không có dòng BOQ nào thay đổi."))
    if doi == "khong" and b_rows:
        kq.append(("P03", "CHẶN", "'Có đổi BOQ?' = Không nhưng mục B có dữ liệu."))

    # chuẩn hoá từng dòng BOQ
    dong, moi_cap, delta = [], [], 0.0
    for b in b_rows:
        op = b["thao_tac"]
        if op not in THAO_TAC:
            kq.append(("P04", "CHẶN", f"{b['vi_tri']}: thao tác '{op}' không hợp lệ (THAY/THEM/BO)."))
            continue
        if op in ("THAY", "BO"):
            if b["ma_hm"] not in hl:
                ly = f"đã bị bỏ theo {bo[b['ma_hm']]['bo_boi']}" if b["ma_hm"] in bo else "không có trong BOQ hiện hành"
                kq.append(("P04", "CHẶN", f"{b['vi_tri']}: {op} {b['ma_hm'] or '(trống)'} - {ly}."))
                continue
            cu = hl[b["ma_hm"]]
            kl_m = b["kl"] if b.get("kl") is not None else cu["kl"]
            dg_m = b["don_gia"] if b.get("don_gia") is not None else cu["don_gia"]
            if op == "BO":
                kl_m, dg_m = 0, cu["don_gia"]
            d = round((kl_m or 0) * (dg_m or 0)) - round((cu["kl"] or 0) * (cu["don_gia"] or 0))
            if op == "THAY" and abs((dg_m or 0) - (cu["don_gia"] or 0)) > 0.5 and loai_gia in ("trọn gói", "đơn giá cố định"):
                md = "ĐỎ" if not str(info.get("can_cu") or "").strip() else "VÀNG"
                kq.append(("P16", md, f"{b['ma_hm']}: đổi đơn giá {_f(cu['don_gia'])} -> {_f(dg_m)} trong HĐ {loai_gia}"
                                      + (" mà không ghi căn cứ." if md == "ĐỎ" else f"; căn cứ: {info.get('can_cu')}")))
            if op == "THAY" and b.get("ten") and bo_dau(b["ten"])[:60] != bo_dau(cu["ten"])[:60]:
                kq.append(("P04", "VÀNG", f"{b['vi_tri']}: tên '{b['ten'][:50]}' khác tên hiện hành '{cu['ten'][:50]}' của {b['ma_hm']}."))
        else:  # THEM
            if b["ma_hm"] and (b["ma_hm"] in hl or b["ma_hm"] in bo):
                kq.append(("P04", "CHẶN", f"{b['vi_tri']}: THEM nhưng {b['ma_hm']} đã tồn tại. Dùng THAY hoặc để trống Mã HM."))
                continue
            if not (b["ten"] and b["dvt"] and b["kl"] is not None and b["don_gia"] is not None):
                kq.append(("P04", "CHẶN", f"{b['vi_tri']}: THEM phải đủ tên, ĐVT, KL, đơn giá."))
                continue
            if not b["ma_hm"]:
                b["ma_hm"] = ma_hm_moi(con, ma, moi_cap)
            moi_cap.append(b["ma_hm"])
            dup = _khop_ten(dict(b, ma_hm=""), hl)
            if dup:
                kq.append(("P04", "VÀNG", f"{b['vi_tri']}: hạng mục thêm mới trùng tên/ĐVT với {dup} đang có. Có phải THAY không?"))
            d = round(b["kl"] * b["don_gia"])
        if b.get("thanh_tien_ghi") is not None and op != "BO":
            kl_c = b["kl"] if b.get("kl") is not None else hl.get(b["ma_hm"], {}).get("kl")
            dg_c = b["don_gia"] if b.get("don_gia") is not None else hl.get(b["ma_hm"], {}).get("don_gia")
            if abs(b["thanh_tien_ghi"] - round((kl_c or 0) * (dg_c or 0))) > 1:
                kq.append(("P04", "VÀNG", f"{b['vi_tri']}: thành tiền ghi {b['thanh_tien_ghi']:,.0f} ≠ KL × ĐG {round((kl_c or 0) * (dg_c or 0)):,.0f}."))
        delta += d
        b["_delta"] = d
        dong.append(b)

    # giá trị phụ lục
    gtt, gts = tien(info.get("gt_truoc_vat")), tien(info.get("gt_sau_vat"))
    suy = ""
    if gts is None and gtt is not None:
        gts, suy = round(gtt * (1 + vat)), " (suy ra từ trước VAT)"
    if gtt is None and gts is not None:
        gtt = round(gts / (1 + vat))
    if gts is None:
        gts = gtt = 0.0
        if doi == "co" and abs(delta) > 0:
            kq.append(("P10", "ĐỎ", f"Phụ lục không ghi giá trị tăng/giảm, nhưng các dòng BOQ làm thay đổi {delta * (1 + vat):,.0f} (sau VAT)."))
    if doi == "co" and dong:
        tinh = delta * (1 + vat)
        if abs(tinh - gts) > tol_tien(gts):
            kq.append(("P10", "ĐỎ", f"Σ thay đổi BOQ × (1+VAT) = {tinh:,.0f} khác giá trị ghi trên phụ lục {gts:,.0f}{suy} "
                                    f"(lệch {tinh - gts:,.0f})."))
        else:
            kq.append(("P10", "ĐẠT", f"Σ thay đổi BOQ × (1+VAT) = {tinh:,.0f} khớp giá trị phụ lục {gts:,.0f}."))
    if doi == "khong" and abs(gts) > 0 and loai_gia != "trọn gói":
        kq.append(("P17", "VÀNG", f"Giá trị HĐ đổi {gts:,.0f} nhưng phụ lục không đổi BOQ: không đối chiếu được hạng mục."))
    gt_sau = tien(info.get("gt_hd_sau_pl"))
    if gt_sau is not None:
        if abs(gt_truoc_pl + gts - gt_sau) > tol_tien(gt_sau):
            kq.append(("P11", "ĐỎ", f"GT HĐ trước PL {gt_truoc_pl:,.0f} + tăng/giảm {gts:,.0f} = {gt_truoc_pl + gts:,.0f}, "
                                    f"khác GT sau PL ghi trên phụ lục {gt_sau:,.0f}."))
        else:
            kq.append(("P11", "ĐẠT", f"Nối chuỗi giá trị HĐ khớp: {gt_truoc_pl:,.0f} + {gts:,.0f} = {gt_sau:,.0f}."))

    # thứ tự / ngày
    so_cu = sorted(int(m.group(1)) for (s,) in con.execute(
        "SELECT so_pl FROM phu_luc WHERE ma_hd = ? AND trang_thai <> 'Đã hủy'", (ma,)) for m in [re.search(r"(\d+)", s)] if m)
    m = re.search(r"(\d+)", so_pl)
    if m and int(m.group(1)) != (so_cu[-1] if so_cu else 0) + 1:
        kq.append(("P12", "VÀNG", f"Số phụ lục không liên tục: đã có {so_cu or 'chưa có PL nào'}, nạp {so_pl}."))
    if ngay_ky and hd["ngay_ky"] and ngay_ky < hd["ngay_ky"]:
        kq.append(("P13", "ĐỎ", f"Ngày ký PL {ngay_vn(ngay_ky)} trước ngày ký HĐ {ngay_vn(hd['ngay_ky'])}."))
    if ngay_ky and ngay_hl and ngay_hl < ngay_ky:
        kq.append(("P13", "VÀNG", f"Phụ lục có hiệu lực hồi tố: hiệu lực {ngay_vn(ngay_hl)} trước ngày ký {ngay_vn(ngay_ky)}."))

    # đối chiếu với KL đã thanh toán
    da_tt = kl_da_thanh_toan(con, ma)
    for b in dong:
        if b["ma_hm"] not in da_tt:
            continue
        dot, lk, ngay_nt = da_tt[b["ma_hm"]]
        cu = hl.get(b["ma_hm"])
        if b["thao_tac"] == "BO" and lk > 1e-9:
            kq.append(("P14", "ĐỎ", f"BO {b['ma_hm']} nhưng đã thanh toán lũy kế {_f(lk)} {cu['dvt']} đến đợt {dot}."))
        if b["thao_tac"] == "THAY" and b.get("kl") is not None and b["kl"] + 1e-9 < lk:
            kq.append(("P14", "ĐỎ", f"THAY {b['ma_hm']} giảm KL xuống {_f(b['kl'])} < lũy kế đã thanh toán {_f(lk)} (đợt {dot})."))
        if b["thao_tac"] == "THAY" and cu and lk > (cu["kl"] or 0) + 1e-9 and ngay_nt and ngay_ky and ngay_nt < ngay_ky:
            kq.append(("P15", "ĐỎ", f"{b['ma_hm']}: đợt {dot} (nghiệm thu {ngay_vn(ngay_nt)}) đã thanh toán lũy kế {_f(lk)} "
                                    f"vượt KL {_f(cu['kl'])} trước khi ký {so_pl} ({ngay_vn(ngay_ky)}): hợp thức hoá sau."))
    if tt == "Dự thảo":
        kq.append(("P18", "VÀNG", "Phụ lục dự thảo: lưu để theo dõi, KHÔNG áp dụng khi đối chiếu."))

    # tham số
    dong_ts = []
    ma_ts_co = {r[0] for r in con.execute("SELECT ma_ts FROM tham_so WHERE ma_hd = ?", (ma,))}
    for a in a_rows:
        if not re.fullmatch(r"(TC|TT|Y|BS)-\d+", a["ma_ts"]):
            kq.append(("P05", "CHẶN", f"Mục A dòng {a['dong']}: mã tham số '{a['ma_ts']}' không hợp lệ."))
            continue
        if a["moi"] in (None, ""):
            kq.append(("P05", "CHẶN", f"Mục A dòng {a['dong']}: {a['ma_ts']} thiếu giá trị mới."))
            continue
        if a["ma_ts"] not in ma_ts_co:
            kq.append(("P05", "VÀNG", f"{a['ma_ts']} chưa có trong tham số gốc của HĐ."))
        hien, _ = tham_so_hieu_luc(con, ma, a["ma_ts"], ngay_hl)
        if a["cu"] not in (None, "") and bo_dau(a["cu"]) != bo_dau(hien) and \
                not (num(a["cu"]) is not None and num(hien) is not None and abs(num(a["cu"]) - num(hien)) < 1e-9):
            kq.append(("P05", "VÀNG", f"{a['ma_ts']}: 'giá trị cũ' trên phiếu '{a['cu']}' khác giá trị hiện hành '{hien}'."))
        if not str(a["trich"] or "").strip():
            kq.append(("P05", "VÀNG", f"{a['ma_ts']}: chưa có trích dẫn nguyên văn phụ lục."))
        if a["ma_ts"] == "TT-02":
            kq.append(("P05", "THÔNG TIN", "TT-02 ở mục A bỏ qua: máy tự tính từ giá trị tăng/giảm của phụ lục."))
            continue
        dong_ts.append(a)

    # TT-02 không nhập ở mục A: máy tự ghi thay đổi TT-02 từ giá trị phụ lục, để không nhập hai lần
    if abs(gts or 0) > 0:
        tt02_cu, _ = tham_so_hieu_luc(con, ma, "TT-02", ngay_hl)
        cu = tien(tt02_cu) or 0
        moi = cu + (gtt if bo_dau(hd["gom_vat"]) in ("khong", "chua") else gts)
        dong_ts = [a for a in dong_ts if a["ma_ts"] != "TT-02"]
        dong_ts.append(dict(dong=0, ma_ts="TT-02", ten="Giá trị hợp đồng", cu=f"{cu:.0f}", moi=f"{moi:.0f}",
                            trich=f"Tự tính: TT-02 trước PL + giá trị tăng/giảm ghi trên {so_pl}"
                                  + (f"; GT HĐ sau PL ghi trên phụ lục: {gt_sau:,.0f}" if gt_sau is not None else "")))
    pl = dict(ma_hd=ma, so_pl=so_pl, ngay_ky=ngay_ky, ngay_hieu_luc=ngay_hl, loai=info.get("loai"),
              doi_boq=1 if doi == "co" else 0, gt_truoc_vat=gtt, gt_sau_vat=gts, gt_hd_sau_pl=gt_sau,
              noi_dung=info.get("noi_dung"), can_cu=info.get("can_cu"), file_nguon=info.get("file_nguon"),
              trang_thai=tt, nguoi_nhap=info.get("nguoi_nhap") or _nguoi())
    return kq, pl, dong, dong_ts


def nap_phu_luc(con, path, chi_kiem_tra=False):
    wb = openpyxl.load_workbook(path, data_only=True)
    ket_qua = []
    sheets = [ws for ws in wb.worksheets if bo_dau(ws.cell(1, 1).value).startswith("phieu nhap phu luc")]
    if not sheets:
        raise SystemExit("Không thấy sheet phiếu phụ lục (ô A1 = 'PHIẾU NHẬP PHỤ LỤC HỢP ĐỒNG'). Tạo phiếu bằng lệnh tao-phieu.")
    for ws in sheets:
        info, a_rows, b_rows, note = doc_phieu(wb, ws)
        hd = tim_hd(con, info.get("so_hd") or "")
        kq, pl, dong, dong_ts = kiem_tra_phu_luc(con, hd, info, a_rows, b_rows, note)
        chan = [x for x in kq if x[1] == "CHẶN"]
        if chan or chi_kiem_tra:
            ket_qua.append((ws.title, pl, kq, "KHÔNG GHI" if chan else "Chỉ kiểm tra"))
            continue
        cur = con.execute("INSERT INTO phu_luc(ma_hd, so_pl, ngay_ky, ngay_hieu_luc, loai, doi_boq, gt_truoc_vat, gt_sau_vat, "
                          "gt_hd_sau_pl, noi_dung, can_cu, file_nguon, phieu, trang_thai, nguoi_nhap, ngay_nap) "
                          "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                          (pl["ma_hd"], pl["so_pl"], pl["ngay_ky"], pl["ngay_hieu_luc"], pl["loai"], pl["doi_boq"],
                           pl["gt_truoc_vat"], pl["gt_sau_vat"], pl["gt_hd_sau_pl"], pl["noi_dung"], pl["can_cu"],
                           pl["file_nguon"], f"{Path(path).name}!{ws.title}", pl["trang_thai"], pl["nguoi_nhap"], bay_gio()))
        pl_id = cur.lastrowid
        for b in dong:
            it = dict(nhom=b.get("nhom") or "", ten=b.get("ten") or "", dvt=b.get("dvt") or "")
            con.execute("INSERT INTO boq_dong(ma_hd, pl_id, thao_tac, ma_hm, stt, nhom, ten, dvt, kl, don_gia, thanh_tien_ghi, "
                        "khoa, vi_tri_nguon, ghi_chu) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                        (pl["ma_hd"], pl_id, b["thao_tac"], b["ma_hm"], b.get("stt") or None, it["nhom"] or None,
                         it["ten"] or None, it["dvt"] or None, b.get("kl"), b.get("don_gia"), b.get("thanh_tien_ghi"),
                         khoa(it) if b["thao_tac"] == "THEM" else None, f"{Path(path).name}!{b['vi_tri']}", b.get("ghi_chu")))
        for a in dong_ts:
            con.execute("INSERT INTO thay_doi_tham_so VALUES (?,?,?,?,?)",
                        (pl_id, a["ma_ts"], None if a["cu"] is None else str(a["cu"]), str(a["moi"]), a["trich"]))
        for ma_kt, md, nd in kq:
            if md not in ("ĐẠT", "THÔNG TIN"):
                con.execute("INSERT INTO canh_bao(ma_hd, pl_id, ma_kt, muc_do, noi_dung, thoi_diem) VALUES (?,?,?,?,?,?)",
                            (pl["ma_hd"], pl_id, ma_kt, md, nd, bay_gio()))
        ghi_nhat_ky(con, "Nạp phụ lục", pl["ma_hd"],
                    f"{pl['so_pl']} ({pl['trang_thai']}) từ {Path(path).name}!{ws.title}: {len(dong)} dòng BOQ, {len(dong_ts)} tham số; "
                    f"{sum(1 for x in kq if x[1] == 'ĐỎ')} cờ đỏ")
        ket_qua.append((ws.title, pl, kq, "ĐÃ GHI"))
    con.commit()
    return ket_qua


def huy_phu_luc(con, ref, so_pl, ly_do):
    if not (ly_do or "").strip():
        raise SystemExit("Huỷ phụ lục phải có --ly-do.")
    hd = tim_hd(con, ref)
    pl = con.execute("SELECT * FROM phu_luc WHERE ma_hd = ? AND so_pl = ? AND trang_thai <> 'Đã hủy'",
                     (hd["ma_hd"], so_pl.upper())).fetchone()
    if not pl:
        raise SystemExit(f"Không có {so_pl} đang hiệu lực của {hd['so_hd']}.")
    sau = con.execute("SELECT so_pl FROM phu_luc WHERE ma_hd = ? AND trang_thai <> 'Đã hủy' AND id <> ? AND "
                      "(ngay_hieu_luc > ? OR (ngay_hieu_luc = ? AND so_pl > ?))",
                      (hd["ma_hd"], pl["id"], pl["ngay_hieu_luc"], pl["ngay_hieu_luc"], pl["so_pl"])).fetchall()
    if sau:
        print("Lưu ý: còn phụ lục sau phụ lục này (" + ", ".join(r[0] for r in sau) + "); kiểm tra lại chúng sau khi nạp lại.")
    con.execute("UPDATE phu_luc SET trang_thai = 'Đã hủy', ly_do_huy = ? WHERE id = ?", (ly_do, pl["id"]))
    ghi_nhat_ky(con, "Huỷ phụ lục", hd["ma_hd"], f"{pl['so_pl']}: {ly_do}")
    con.commit()


# ============================================================== đợt thanh toán
def _khop_hstt(con, ma_hd, items, hl):
    """Gán Mã HM cho từng dòng hồ sơ: bảng khop_hang_muc -> khoá đầy đủ -> tên + ĐVT duy nhất -> gợi ý gần đúng."""
    nho = {r["khoa_hstt"]: r["ma_hm"] for r in con.execute("SELECT * FROM khop_hang_muc WHERE ma_hd = ?", (ma_hd,))}
    by_khoa = {}
    for ma, d in hl.items():
        by_khoa.setdefault(d["khoa"], ma)
    for it in items:
        if not it["dvt"]:
            continue
        it["ma_hm"], it["cach_khop"], it["goi_y"] = None, "", ""
        if it["khoa"] in nho:
            it["ma_hm"], it["cach_khop"] = nho[it["khoa"]], "đã xác nhận"
        elif it["khoa"] in by_khoa:
            it["ma_hm"], it["cach_khop"] = by_khoa[it["khoa"]], "khoá"
        else:
            ma = _khop_ten(dict(it, ma_hm=""), hl)
            if ma:
                it["ma_hm"], it["cach_khop"] = ma, "tên + ĐVT"
            else:
                best = max(((difflib.SequenceMatcher(None, bo_dau(it["ten"]), bo_dau(d["ten"])).ratio(), m)
                            for m, d in hl.items() if bo_dau(d["dvt"]) == bo_dau(it["dvt"])), default=(0, None))
                if best[0] >= 0.85:
                    it["goi_y"] = f"{best[1]} ({best[0]:.0%})"


def nap_dot(con, ref, path, dot, ngay_nt, ngay_hs=None, ly_do=""):
    """Nạp đợt ĐÃ DUYỆT: KL lũy kế từng hạng mục + số tổng hợp. Là nguồn KL kỳ trước cho đợt sau."""
    hd = tim_hd(con, ref)
    ma = hd["ma_hd"]
    ngay_nt = ngay_iso(ngay_nt)
    wv = openpyxl.load_workbook(path, data_only=True)
    bang, tong = tim_bang_chi_tiet(wv), tim_so_tong_hop(wv)
    hl, bo = boq_hieu_luc(con, ma, ngay_nt)
    _khop_hstt(con, ma, bang["items"], {**bo, **hl})
    cu = con.execute("SELECT * FROM dot_thanh_toan WHERE ma_hd = ? AND dot = ?", (ma, dot)).fetchone()
    if cu:
        if not ly_do:
            raise SystemExit(f"Đợt {dot} đã có trong CSDL (nguồn: {cu['file_nguon']}). Nạp lại thì thêm --ly-do.")
        with che_do_sua(con, ly_do, ma):
            con.execute("DELETE FROM kl_thanh_toan WHERE ma_hd = ? AND dot = ?", (ma, dot))
            con.execute("DELETE FROM dot_thanh_toan WHERE ma_hd = ? AND dot = ?", (ma, dot))
    g = lambda k: tong.get(k, {}).get("gia_tri")
    con.execute("INSERT INTO dot_thanh_toan VALUES (?,?,?,?,?,?,?,?,?,?,?)",
                (ma, dot, ngay_nt, ngay_iso(ngay_hs), g("gt_kn"), g("gt_lk"), g("giu_lai"), g("thu_hoi_tu"), g("tt_ky_nay"),
                 f"{Path(path).name}: " + "; ".join(f"{k}={v['o']}" for k, v in tong.items()), bay_gio()))
    ok, miss = 0, []
    for it in bang["items"]:
        if not it["dvt"]:
            continue
        if not it["ma_hm"]:
            miss.append(it)
            continue
        con.execute("INSERT OR REPLACE INTO kl_thanh_toan VALUES (?,?,?,?,?,?,?)",
                    (ma, dot, it["ma_hm"], it["kl_kn"], it["kl_lk"], it["don_gia"],
                     f"{Path(path).name} {bang['sheet']}!dòng {it['dong']}"))
        if it["cach_khop"] in ("khoá", "tên + ĐVT"):
            con.execute("INSERT OR IGNORE INTO khop_hang_muc VALUES (?,?,?,?,?,?)",
                        (ma, it["khoa"], it["ma_hm"], f"tự động ({it['cach_khop']})", "", bay_gio()))
        ok += 1
    ghi_nhat_ky(con, f"Nạp đợt {dot}", ma, f"{Path(path).name}; NT {ngay_vn(ngay_nt)}; {ok} hạng mục khớp, {len(miss)} không khớp")
    con.commit()
    return ok, miss


def ngay_nghiem_thu_dot(con, ref, dot, ngay, ly_do=""):
    """Bổ sung / sửa ngày nghiệm thu của một đợt đã nạp (đợt nạp từ master thường chưa có ngày)."""
    hd = tim_hd(con, ref)
    r = con.execute("SELECT ngay_nghiem_thu FROM dot_thanh_toan WHERE ma_hd = ? AND dot = ?", (hd["ma_hd"], dot)).fetchone()
    if not r:
        raise SystemExit(f"Chưa có đợt {dot} của {hd['so_hd']} trong CSDL.")
    if r[0] and not ly_do:
        raise SystemExit(f"Đợt {dot} đã có ngày nghiệm thu {ngay_vn(r[0])}. Sửa thì thêm --ly-do.")
    with che_do_sua(con, ly_do or f"Bổ sung ngày nghiệm thu đợt {dot} (trước đó trống)", hd["ma_hd"]):
        con.execute("UPDATE dot_thanh_toan SET ngay_nghiem_thu = ? WHERE ma_hd = ? AND dot = ?",
                    (ngay_iso(ngay), hd["ma_hd"], dot))
    con.commit()


def _doan_dot(path):
    m = re.search(r"đợt\s*0*(\d+)", nfc(Path(path).name), re.I) or re.search(r"dot\s*0*(\d+)", bo_dau(Path(path).name))
    return int(m.group(1)) if m else None


def kiem_dot(con, ref, path, ngay_nt, out=None, dot=None):
    """Đối chiếu hồ sơ đợt mới với BOQ / GT hợp đồng hiệu lực tại ngày nghiệm thu (Bước 1-3 mục 4 B2)."""
    hd = tim_hd(con, ref)
    ma = hd["ma_hd"]
    ngay = ngay_iso(ngay_nt)
    loai = loai_gia_chuan(hd["loai_gia"])
    vat = hd["vat"] or 0
    pls = phu_luc_ap_dung(con, ma, ngay)
    hl, bo = boq_hieu_luc(con, ma, pls=pls)
    gt_hl = gt_hd_hieu_luc(con, ma, pls=pls)
    tat_ca = {**bo, **hl}  # để nhận ra hạng mục đã bị bỏ
    con_lai = con.execute("SELECT * FROM phu_luc WHERE ma_hd = ? AND trang_thai <> 'Đã hủy' AND "
                          "(trang_thai = 'Dự thảo' OR ngay_hieu_luc > ?) ORDER BY ngay_hieu_luc", (ma, ngay)).fetchall()
    wv = openpyxl.load_workbook(path, data_only=True)
    bang, tong = tim_bang_chi_tiet(wv), tim_so_tong_hop(wv)
    _khop_hstt(con, ma, bang["items"], tat_ca)
    # KL lũy kế đợt trước đã duyệt
    dot = dot or _doan_dot(path)
    if dot:
        dot_truoc = con.execute("SELECT MAX(dot) FROM dot_thanh_toan WHERE ma_hd = ? AND dot < ?", (ma, dot)).fetchone()[0]
    else:
        dot_truoc = con.execute("SELECT MAX(dot) FROM dot_thanh_toan WHERE ma_hd = ? AND (ngay_nghiem_thu IS NULL OR "
                                "ngay_nghiem_thu < ?)", (ma, ngay)).fetchone()[0]
    kt = {r["ma_hm"]: r["kl_luy_ke"] for r in con.execute(
        "SELECT ma_hm, kl_luy_ke FROM kl_thanh_toan WHERE ma_hd = ? AND dot = ?", (ma, dot_truoc))} if dot_truoc else {}
    kt_chi_tiet = bool(kt)
    # PL ngoài phạm vi: hạng mục chỉ xuất hiện ở đó
    ngoai = {}
    for pl in con_lai:
        for r in con.execute("SELECT * FROM boq_dong WHERE pl_id = ?", (pl["id"],)):
            ngoai.setdefault(r["ma_hm"], []).append(pl)

    # BƯỚC 1: giá trị hợp đồng
    b1 = []
    gt_file = tong.get("gt_hd", {}).get("gia_tri")
    if gt_file is None:
        b1.append(("Không đọc được GT HĐ trong hồ sơ", None, gt_hl, "THIẾU SỐ LIỆU", ""))
    elif abs(gt_file - gt_hl) <= tol_tien(gt_hl):
        b1.append(("GT HĐ trong hồ sơ = GT HĐ hiệu lực", gt_file, gt_hl, "ĐẠT", ""))
    else:
        ly = "CẦN PHỤ LỤC: GT HĐ trong hồ sơ khác GT hiệu lực và không phụ lục nào giải thích."
        g = gt_sau_vat_goc(hd)
        for k_, pl in enumerate([None] + list(pls)):
            g += (pl["gt_sau_vat"] or 0) if pl else 0
            if k_ < len(pls) and abs(gt_file - g) <= tol_tien(g):
                ly = ("HỒ SƠ CHƯA CẬP NHẬT PHỤ LỤC: hồ sơ dùng GT " + ("HĐ gốc" if pl is None else f"sau {pl['so_pl']}")
                      + f", chưa tính {', '.join(p['so_pl'] for p in pls[k_:])} đã có hiệu lực trước ngày nghiệm thu.")
                break
        g = gt_hl
        for pl in con_lai:
            g += pl["gt_sau_vat"] or 0
            if abs(gt_file - g) <= tol_tien(g):
                ly = (f"PHỤ LỤC CHƯA HIỆU LỰC: hồ sơ đã dùng GT sau {pl['so_pl']} ({pl['trang_thai']}, hiệu lực "
                      f"{ngay_vn(pl['ngay_hieu_luc'])}) trong khi nghiệm thu ngày {ngay_vn(ngay)}.")
                break
        b1.append(("GT HĐ trong hồ sơ ≠ GT HĐ hiệu lực", gt_file, gt_hl, ly, tong["gt_hd"]["o"]))

    # BƯỚC 3: từng hạng mục
    rows, seen = [], set()
    for it in bang["items"]:
        if not it["dvt"]:
            continue
        m = hl.get(it["ma_hm"]) if it["ma_hm"] else None
        loi = []
        if not m:
            if it["ma_hm"] and it["ma_hm"] in bo:
                loi.append(f"Hạng mục đã bị bỏ theo {bo[it['ma_hm']]['bo_boi']}")
            else:
                for pl_ma, pl_list in ngoai.items():
                    r0 = con.execute("SELECT * FROM boq_dong WHERE ma_hm = ? AND ma_hd = ? AND pl_id = ?",
                                     (pl_ma, ma, pl_list[0]["id"])).fetchone()
                    if r0 and r0["khoa"] == it["khoa"]:
                        p = pl_list[0]
                        loi.append(f"Chỉ có trong {p['so_pl']} ({p['trang_thai']}, hiệu lực {ngay_vn(p['ngay_hieu_luc'])}) - "
                                   f"chưa áp dụng tại ngày nghiệm thu")
                        break
                if not loi:
                    loi.append("Không có trong BOQ hiệu lực: phát sinh chưa có phụ lục"
                               + (f" (gần giống {it['goi_y']})" if it["goi_y"] else ""))
        else:
            seen.add(it["ma_hm"])
            if abs((it["don_gia"] or 0) - (m["don_gia"] or 0)) > 0.5:
                loi.append("Đơn giá khác BOQ hiệu lực")
            if it["kl_hd"] is not None and abs(it["kl_hd"] - (m["kl"] or 0)) > 1e-6:
                loi.append("KL HĐ khác BOQ hiệu lực")
            if (it["kl_lk"] or 0) > (m["kl"] or 0) + 1e-6:
                loi.append("Lũy kế vượt KL HĐ hiệu lực: " + ("HĐ trọn gói không thanh toán vượt" if loai == "trọn gói"
                                                            else "cần phụ lục / biên bản phát sinh theo TT-42"))
            if kt_chi_tiet and it["ma_hm"] in kt and abs((it["kl_kt"] or 0) - (kt[it["ma_hm"]] or 0)) > 1e-6:
                loi.append(f"KL kỳ trước khác lũy kế đợt {dot_truoc} đã duyệt")
            if kt_chi_tiet and it["ma_hm"] not in kt and (it["kl_kt"] or 0) > 0:
                loi.append(f"Có KL kỳ trước nhưng đợt {dot_truoc} không có hạng mục này")
            if not dot_truoc and (it["kl_kt"] or 0) > 0:
                loi.append("Chưa có đợt trước trong CSDL: KL kỳ trước phải nhập tay/đối chiếu hồ sơ gốc")
        # ảnh hưởng = GT lũy kế theo hồ sơ - GT lũy kế được chấp nhận theo BOQ hiệu lực (chưa VAT)
        lk = it["kl_lk"] or 0
        gt_hs = lk * (it["don_gia"] or 0)
        gt_ok = min(lk, (m or {}).get("kl") or 0) * ((m or {}).get("don_gia") or 0) if m else 0
        rows.append(dict(it=it, m=m, loi=loi, anh_huong=round(gt_hs - gt_ok) if loi else 0))
    thieu = [d for ma_hm, d in hl.items() if ma_hm not in seen]

    # Σ kiểm tra giá trị
    vat1 = 1 + vat
    gt_kn = tong.get("gt_kn", {}).get("gia_tri")
    sum_kn = sum((x["it"]["kl_kn"] or 0) * ((x["m"] or {}).get("don_gia") or 0) for x in rows) * vat1

    out = Path(out or f"kiem_dot_{ma}_{Path(path).stem[:40]}.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tong hop"
    put(ws, 1, 1, "KIỂM TRA HỒ SƠ THANH TOÁN VỚI HỢP ĐỒNG + PHỤ LỤC HIỆU LỰC", TITLE)
    put(ws, 2, 1, f"{hd['so_hd']} ({ma}) - loại giá: {loai or '?'} - hồ sơ: {Path(path).name} - bảng: {bang['sheet']}")
    put(ws, 4, 1, f"Đợt kiểm tra: {dot or '?'}; đợt trước trong CSDL: {dot_truoc or 'không có'}"
                  + ("" if kt_chi_tiet or not dot_truoc else f" (đợt {dot_truoc} chỉ có số tổng, chưa có KL từng hạng mục: "
                     "KL kỳ trước chưa đối chiếu được)"), GRAY)
    put(ws, 3, 1, f"Mốc áp dụng phụ lục = ngày nghiệm thu {ngay_vn(ngay)}. Phụ lục áp dụng: "
                  + (", ".join(f"{p['so_pl']} ({ngay_vn(p['ngay_hieu_luc'])})" for p in pls) or "không có")
                  + (". Chưa áp dụng: " + ", ".join(f"{p['so_pl']} ({p['trang_thai']}, {ngay_vn(p['ngay_hieu_luc'])})"
                                                     for p in con_lai) if con_lai else ""), GRAY)
    header(ws, 6, ["Kiểm tra", "Theo hồ sơ", "Theo CSDL (hiệu lực)", "Chênh lệch", "Kết luận", "Ô trong hồ sơ"],
           [48, 18, 20, 16, 70, 26])
    r = 7
    put(ws, r, 1, "BƯỚC 1 - Giá trị hợp đồng", BOLD); r += 1
    for t, a, b, kl, o in b1:
        put(ws, r, 1, t); put(ws, r, 2, a, fmt=NUMF); put(ws, r, 3, b, fmt=NUMF)
        put(ws, r, 4, (a - b) if a is not None else None, fmt=NUMF)
        put(ws, r, 5, kl, fill=GREEN if kl == "ĐẠT" else RED, wrap=True); put(ws, r, 6, o); r += 1
    r += 1
    put(ws, r, 1, "BƯỚC 2 - BOQ hiệu lực", BOLD); r += 1
    put(ws, r, 1, "Số hạng mục BOQ hiệu lực / đã bỏ"); put(ws, r, 2, len(hl)); put(ws, r, 3, len(bo)); r += 1
    put(ws, r, 1, "Hạng mục chịu tác động của phụ lục"); put(ws, r, 2, sum(1 for d in hl.values() if d["nguon"] != "HĐ gốc")); r += 1
    r += 1
    put(ws, r, 1, f"BƯỚC 3 - Từng hạng mục (HĐ {loai})", BOLD); r += 1
    from collections import Counter
    cnt = Counter(l.split(":")[0].split(" (")[0] for x in rows for l in x["loi"])
    if not cnt:
        put(ws, r, 1, "Không có lệch"); put(ws, r, 5, "ĐẠT", fill=GREEN); r += 1
    for k_, n_ in cnt.most_common():
        put(ws, r, 1, k_); put(ws, r, 2, n_); put(ws, r, 5, "Xem sheet Hang muc", fill=RED); r += 1
    put(ws, r, 1, "Hạng mục BOQ hiệu lực không có trong hồ sơ"); put(ws, r, 2, len(thieu)); r += 1
    tong_ah = sum(x["anh_huong"] for x in rows) * vat1
    put(ws, r, 1, "Giá trị lũy kế chưa có căn cứ HĐ/phụ lục (sau VAT) - tạm chưa thanh toán", BOLD)
    put(ws, r, 3, round(tong_ah), fmt=NUMF, fill=RED if abs(tong_ah) > 1 else GREEN)
    put(ws, r, 5, "Chi tiết ở sheet Yeu cau phu luc"); r += 1
    r += 1
    put(ws, r, 1, "Σ(KL kỳ này × đơn giá hiệu lực) × (1+VAT)"); put(ws, r, 2, gt_kn, fmt=NUMF); put(ws, r, 3, round(sum_kn), fmt=NUMF)
    if gt_kn is not None:
        put(ws, r, 4, gt_kn - round(sum_kn), fmt=NUMF)
        put(ws, r, 5, "ĐẠT" if abs(gt_kn - sum_kn) <= tol_tien(gt_kn) else "Lệch - kiểm tra",
            fill=GREEN if abs(gt_kn - sum_kn) <= tol_tien(gt_kn) else RED)
        put(ws, r, 6, tong.get("gt_kn", {}).get("o", ""))
    r += 1
    if dot_truoc:
        lk_db = con.execute("SELECT gt_luy_ke FROM dot_thanh_toan WHERE ma_hd = ? AND dot = ?", (ma, dot_truoc)).fetchone()[0]
        a = tong.get("gt_lk_kt", {}).get("gia_tri")
        put(ws, r, 1, f"GT lũy kế hết kỳ trước = lũy kế đợt {dot_truoc} đã duyệt"); put(ws, r, 2, a, fmt=NUMF); put(ws, r, 3, lk_db, fmt=NUMF)
        if a is not None and lk_db is not None:
            put(ws, r, 4, a - lk_db, fmt=NUMF)
            put(ws, r, 5, "ĐẠT" if abs(a - lk_db) <= 1 else "Lệch - kiểm tra", fill=GREEN if abs(a - lk_db) <= 1 else RED)
        put(ws, r, 6, tong.get("gt_lk_kt", {}).get("o", ""))
        r += 1

    w2 = wb.create_sheet("Hang muc")
    header(w2, 1, ["Dòng HS", "Mã HM", "Khớp bằng", "Tên công việc", "ĐVT", "KL HĐ hồ sơ", "KL hiệu lực", "Nguồn KL/ĐG",
                   "Đơn giá hồ sơ", "Đơn giá hiệu lực", "KL kỳ trước HS", f"Lũy kế đợt {dot_truoc or '-'}", "KL kỳ này",
                   "KL lũy kế", "Ảnh hưởng đơn giá (× lũy kế)", "Lịch sử BOQ", "Lệch"],
           [7, 8, 12, 50, 8, 12, 12, 9, 13, 13, 12, 12, 12, 12, 16, 40, 60])
    w2.freeze_panes = "E2"
    for i, x in enumerate(rows, 2):
        it, m = x["it"], x["m"] or {}
        vals = [it["dong"], it["ma_hm"], it["cach_khop"], it["ten"], it["dvt"], it["kl_hd"], m.get("kl"), m.get("nguon"),
                it["don_gia"], m.get("don_gia"), it["kl_kt"], kt.get(it["ma_hm"]) if it["ma_hm"] else None, it["kl_kn"], it["kl_lk"]]
        for j, v in enumerate(vals, 1):
            put(w2, i, j, v, fmt=QF if j in (6, 7, 11, 12, 13, 14) else (NUMF if j in (9, 10) else None), text=True)
        put(w2, i, 15, f'=IF(J{i}="","",(I{i}-J{i})*N{i})', fmt=NUMF)
        put(w2, i, 16, " → ".join(m.get("lich_su", [])), text=True)
        put(w2, i, 17, "; ".join(x["loi"]), text=True)
        if x["loi"]:
            w2.cell(i, 17).fill = RED
        if m.get("nguon") and m["nguon"] != "HĐ gốc":
            w2.cell(i, 8).fill = YELLOW
    wy = wb.create_sheet("Yeu cau phu luc")
    header(wy, 1, ["Mã HM", "Tên công việc (theo hồ sơ)", "ĐVT", "KL lũy kế hồ sơ", "KL hiệu lực", "Đơn giá hồ sơ",
                   "Đơn giá hiệu lực", "Ảnh hưởng chưa VAT", "Ảnh hưởng sau VAT", "Nội dung cần phụ lục / giải trình"],
           [8, 50, 8, 13, 13, 14, 14, 16, 16, 70])
    i = 2
    for x in rows:
        if not x["loi"]:
            continue
        it, m = x["it"], x["m"] or {}
        for j, v in enumerate([it["ma_hm"], it["ten"], it["dvt"], it["kl_lk"], m.get("kl"), it["don_gia"], m.get("don_gia"),
                               x["anh_huong"], round(x["anh_huong"] * vat1), "; ".join(x["loi"])], 1):
            put(wy, i, j, v, fmt=QF if j in (4, 5) else (NUMF if j in (6, 7, 8, 9) else None), text=True)
        i += 1
    put(wy, i, 2, "Cộng", BOLD)
    put(wy, i, 8, f"=SUM(H2:H{i - 1})" if i > 2 else 0, BOLD, NUMF)
    put(wy, i, 9, f"=SUM(I2:I{i - 1})" if i > 2 else 0, BOLD, NUMF)
    w3 = wb.create_sheet("Thieu trong ho so")
    header(w3, 1, ["Mã HM", "Tên công việc", "ĐVT", "KL hiệu lực", "Nguồn"], [9, 60, 9, 12, 10])
    for i, d in enumerate(thieu, 2):
        for j, v in enumerate([d["ma_hm"], d["ten"], d["dvt"], d["kl"], d["nguon"]], 1):
            put(w3, i, j, v, fmt=QF if j == 4 else None, text=True)
    w4 = wb.create_sheet("Phu luc")
    header(w4, 1, ["Số PL", "Ngày ký", "Ngày hiệu lực", "Trạng thái", "Áp dụng cho đợt này?", "GT tăng/giảm sau VAT", "Nội dung"],
           [8, 11, 12, 10, 18, 18, 60])
    for i, p in enumerate(con.execute("SELECT * FROM phu_luc WHERE ma_hd = ? ORDER BY ngay_hieu_luc, so_pl", (ma,)), 2):
        ad = "Có" if p["id"] in {q["id"] for q in pls} else "Không"
        for j, v in enumerate([p["so_pl"], ngay_vn(p["ngay_ky"]), ngay_vn(p["ngay_hieu_luc"]), p["trang_thai"], ad,
                               p["gt_sau_vat"], p["noi_dung"]], 1):
            put(w4, i, j, v, fmt=NUMF if j == 6 else None, text=True)
    wb.save(out)
    ghi_nhat_ky(con, "Kiểm tra đợt", ma, f"{Path(path).name}; NT {ngay_vn(ngay)}; kết quả {out.name}")
    con.commit()
    return out, b1, cnt, len(thieu), pls, con_lai, dot, dot_truoc


# ============================================================== xuất Excel
def _gon(ma_ts, v):
    """Hiển thị gọn giá trị tham số: số tiền có dấu phân cách, tỷ lệ dạng %."""
    if v in (None, ""):
        return ""
    if ma_ts in MA_TY_LE and ty_le(v) is not None and not re.search(r"[a-zA-ZÀ-ỹ]", str(v)):
        return f"{ty_le(v) * 100:g}%"
    if ma_ts == "TT-02" and tien(v) is not None:
        return f"{tien(v):,.0f}"
    return str(v)


def _gia_tri_o(ma_ts, v):
    """Giá trị ghi vào ô Excel + định dạng."""
    if v in (None, ""):
        return None, None
    if ma_ts in MA_TY_LE and not re.search(r"[a-zA-ZÀ-ỹ]", str(v)) and ty_le(v) is not None:
        return ty_le(v), "0.00%"
    if ma_ts == "TT-02" and tien(v) is not None:
        return tien(v), NUMF
    if ma_ts in ("TT-22", "TT-32", "TT-62") and num(v) is not None:
        return num(v), None
    return str(v), None


def _sheet_hop_dong(con, ws, hds, ngay):
    """Mỗi hợp đồng một dòng: 36 tham số hiệu lực (thứ tự như DANH_MUC_HD) + các cột kiểm soát.
    Ô bị phụ lục sửa tô vàng; mọi ô có ghi chú điều khoản + trích dẫn nguyên văn."""
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment
    from master_hd import HF, HFONT
    put(ws, 1, 1, f"BẢN XUẤT TỪ CƠ SỞ DỮ LIỆU lúc {bay_gio()} - tham số và giá trị hiệu lực tính đến ngày {ngay_vn(ngay)}. "
                  "Ô vàng = đã bị phụ lục sửa. Rê chuột vào ô để xem điều khoản, trích dẫn nguyên văn và lịch sử phụ lục. "
                  "Sửa trên file này KHÔNG làm thay đổi CSDL.", Font(name="Arial", size=10, bold=True, color="9A5B00"))
    ma_list = [m for _, g in DS_THAM_SO for m, _ in g]
    ten_ma = {m: t for _, g in DS_THAM_SO for m, t in g}
    # mã bổ sung (Y-xx, BS-xx) nếu có
    ph = ",".join("?" * len(hds)) or "''"
    them = [r[0] for r in con.execute(f"SELECT DISTINCT ma_ts FROM tham_so WHERE ma_hd IN ({ph}) ORDER BY ma_ts",
                                      [h["ma_hd"] for h in hds]) if r[0] not in ten_ma]
    nhom = [(n, [m for m, _ in g]) for n, g in DS_THAM_SO] + ([("THAM SỐ BỔ SUNG", them)] if them else [])
    kiem_soat = ["Số PL áp dụng", "GT HĐ gốc (sau VAT)", "GT HĐ hiệu lực (sau VAT)", "Tăng/giảm do PL",
                 "BOQ hiệu lực × (1+VAT)", "Chênh BOQ - GT", "Số hạng mục BOQ", "Đợt TT gần nhất", "GT lũy kế đã TT",
                 "Cờ đỏ", "Tham số chưa xác nhận"]
    # 3 dòng tiêu đề: nhóm / mã / tên
    c = 1
    for r, v in ((3, "MÃ HĐ"), (4, ""), (5, "Mã HĐ")):
        put(ws, r, c, v)
    cols = {}
    c = 2
    for ten_nhom, mas in nhom:
        put(ws, 3, c, ten_nhom)
        if len(mas) > 1:
            ws.merge_cells(start_row=3, start_column=c, end_row=3, end_column=c + len(mas) - 1)
        for m in mas:
            cols[m] = c
            put(ws, 4, c, m)
            put(ws, 5, c, ten_ma.get(m) or m)
            c += 1
    c_ks = c
    put(ws, 3, c_ks, "KIỂM SOÁT (tính từ CSDL)")
    ws.merge_cells(start_row=3, start_column=c_ks, end_row=3, end_column=c_ks + len(kiem_soat) - 1)
    for k, t in enumerate(kiem_soat):
        put(ws, 5, c_ks + k, t)
    for r in (3, 4, 5):
        for cc in range(1, c_ks + len(kiem_soat)):
            cl = ws.cell(r, cc)
            cl.font, cl.fill = HFONT, HF
            cl.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[5].height = 45
    from openpyxl.utils import get_column_letter
    ws.column_dimensions["A"].width = 8
    for m, cc in cols.items():
        ws.column_dimensions[get_column_letter(cc)].width = {"TC-01": 24, "TC-03": 34, "TC-04": 30, "TC-06": 30, "TC-07": 30,
                                                             "TT-02": 17, "TT-05": 30, "TT-40": 30, "TT-42": 30, "TT-56": 30,
                                                             "TT-60": 34, "TT-61": 34}.get(m, 13)
    for k in range(len(kiem_soat)):
        ws.column_dimensions[get_column_letter(c_ks + k)].width = 15
    ws.freeze_panes = ws.cell(6, 3)

    for i, h in enumerate(hds, 6):
        ma = h["ma_hd"]
        put(ws, i, 1, ma, BOLD)
        goc = {r["ma_ts"]: r for r in con.execute("SELECT * FROM tham_so WHERE ma_hd = ?", (ma,))}
        ls = lich_su_tham_so(con, ma, ngay)
        chua_xn = []
        for m, cc in cols.items():
            r0 = goc.get(m)
            v = ls[m][-1][3] if m in ls else (r0["gia_tri"] if r0 else None)
            val, fmt = _gia_tri_o(m, v)
            cl = put(ws, i, cc, val, fmt=fmt, text=True)
            if r0 is None and m in ten_ma:
                cl.fill = RED
                cl.value = "THIẾU"
                continue
            note = []
            if r0 is not None:
                if r0["dieu_khoan"]:
                    note.append(str(r0["dieu_khoan"])[:120])
                if r0["trich_dan"]:
                    note.append("HĐ: \"" + str(r0["trich_dan"])[:400] + "\"")
                if r0["trang_thai"] and r0["trang_thai"] != "Đã xác nhận":
                    chua_xn.append(m)
            if m in ls:
                cl.fill = YELLOW
                note.append(f"Gốc: {_gon(m, r0['gia_tri'] if r0 else '')}")
                for so_pl, nhl, cu, moi, td in ls[m]:
                    note.append(f"{so_pl} (hiệu lực {ngay_vn(nhl)}): {_gon(m, cu)} → {_gon(m, moi)}"
                                + (f" | \"{str(td)[:200]}\"" if td else ""))
            if note:
                cm = Comment("\n".join(note), "CSDL")
                cm.width, cm.height = 420, 160
                cl.comment = cm
        # cột kiểm soát
        pls = phu_luc_ap_dung(con, ma, ngay)
        hl, _ = boq_hieu_luc(con, ma, pls=pls)
        g0 = gt_sau_vat_goc(h)
        gth = gt_hd_hieu_luc(con, ma, pls=pls)
        boqv = sum(round((d["kl"] or 0) * (d["don_gia"] or 0)) for d in hl.values()) * (1 + (h["vat"] or 0)) if hl else None
        d_tt = con.execute("SELECT dot, gt_luy_ke FROM dot_thanh_toan WHERE ma_hd = ? ORDER BY dot DESC LIMIT 1", (ma,)).fetchone()
        do = con.execute("SELECT COUNT(*) FROM canh_bao c LEFT JOIN phu_luc p ON p.id = c.pl_id WHERE c.ma_hd = ? "
                         "AND c.muc_do = 'ĐỎ' AND IFNULL(p.trang_thai, '') <> 'Đã hủy'", (ma,)).fetchone()[0]
        vals = [", ".join(p["so_pl"] for p in pls) or "-", g0, gth, gth - g0, boqv,
                (boqv - gth) if boqv is not None else "Chưa có BOQ", len(hl), d_tt["dot"] if d_tt else None,
                d_tt["gt_luy_ke"] if d_tt else None, do, f"{len(chua_xn)}/{len(cols)}"]
        for k, v in enumerate(vals):
            cl = put(ws, i, c_ks + k, v, fmt=NUMF if k in (1, 2, 3, 4, 5, 8) else None, text=True)
        if boqv is None or abs(boqv - gth) > tol_tien(gth):
            ws.cell(i, c_ks + 5).fill = RED
        if do:
            ws.cell(i, c_ks + 9).fill = RED
        if chua_xn:
            ws.cell(i, c_ks + 10).fill = YELLOW


def xuat_excel(con, ref=None, ngay=None, out=None):
    ngay = ngay_iso(ngay) if ngay else dt.date.today().isoformat()
    hds = [tim_hd(con, ref)] if ref else con.execute("SELECT * FROM hop_dong ORDER BY ma_hd").fetchall()
    ids = [h["ma_hd"] for h in hds]
    ph = ",".join("?" * len(ids)) or "''"
    out = Path(out or f"CSDL_hop_dong_{ref or 'tat_ca'}_{ngay}.xlsx".replace("/", "-"))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "HOP_DONG"
    _sheet_hop_dong(con, ws, hds, ngay)

    def bang(ten, sql, args, cot, rong, fmt=None, map_row=None):
        w = wb.create_sheet(ten)
        header(w, 1, cot, rong)
        w.freeze_panes = "A2"
        for i, r in enumerate(con.execute(sql, args), 2):
            vals = map_row(r) if map_row else list(r)
            for j, v in enumerate(vals, 1):
                put(w, i, j, v, fmt=(fmt or {}).get(j), text=True)
        return w

    def _pl_row(r):
        ts = "; ".join(f"{t['ma_ts']}: {_gon(t['ma_ts'], t['gia_tri_cu'])} → {_gon(t['ma_ts'], t['gia_tri_moi'])}"
                       for t in con.execute("SELECT * FROM thay_doi_tham_so WHERE pl_id = ? ORDER BY ma_ts", (r["id"],)))
        bq = {k: n for k, n in con.execute("SELECT thao_tac, COUNT(*) FROM boq_dong WHERE pl_id = ? GROUP BY thao_tac", (r["id"],))}
        do = con.execute("SELECT COUNT(*) FROM canh_bao WHERE pl_id = ? AND muc_do = 'ĐỎ'", (r["id"],)).fetchone()[0]
        return [r["ma_hd"], r["so_pl"], r["trang_thai"], ngay_vn(r["ngay_ky"]), ngay_vn(r["ngay_hieu_luc"]), r["loai"],
                r["gt_truoc_vat"], r["gt_sau_vat"], r["gt_hd_sau_pl"],
                ("Có: " + ", ".join(f"{n} {k}" for k, n in sorted(bq.items()))) if r["doi_boq"] else "Không",
                ts, do or "", r["noi_dung"], r["can_cu"], r["file_nguon"], r["phieu"], r["nguoi_nhap"], r["ngay_nap"], r["ly_do_huy"]]

    wp = bang("PHU_LUC", f"SELECT * FROM phu_luc WHERE ma_hd IN ({ph}) ORDER BY ma_hd, ngay_hieu_luc, so_pl", ids,
              ["Mã HĐ", "Số PL", "Trạng thái", "Ngày ký", "Ngày hiệu lực", "Loại", "Tăng/giảm trước VAT", "Tăng/giảm sau VAT",
               "GT HĐ sau PL (ghi trên PL)", "Đổi BOQ (số dòng)", "Tham số thay đổi (cũ → mới)", "Cờ đỏ", "Nội dung", "Căn cứ",
               "File PL", "Phiếu nhập", "Người nhập", "Ngày nạp", "Lý do huỷ"],
              [8, 7, 9, 11, 11, 16, 16, 16, 18, 16, 50, 6, 40, 30, 18, 24, 10, 17, 22], {7: NUMF, 8: NUMF, 9: NUMF}, _pl_row)
    for r in range(2, wp.max_row + 1):
        if wp.cell(r, 12).value:
            wp.cell(r, 12).fill = RED
        if wp.cell(r, 3).value == "Đã hủy":
            for c in range(1, 20):
                wp.cell(r, c).font = GRAY
    # BOQ hiệu lực
    w = wb.create_sheet("BOQ_HIEU_LUC")
    header(w, 1, ["Mã HĐ", "Mã HM", "STT", "Nhóm / hạng mục", "Tên công việc", "ĐVT", "KL", "Đơn giá (chưa VAT)", "Thành tiền",
                  "Nguồn", "Lịch sử"], [8, 8, 6, 30, 56, 8, 12, 15, 16, 9, 50])
    w.freeze_panes = "C2"
    i = 2
    for h in hds:
        hl, bo = boq_hieu_luc(con, h["ma_hd"], ngay)
        for d in list(hl.values()) + [dict(x, nguon=f"Bỏ ({x['bo_boi']})") for x in bo.values()]:
            vals = [h["ma_hd"], d["ma_hm"], d["stt"], d["nhom"], d["ten"], d["dvt"], d["kl"] if "bo_boi" not in d else 0,
                    d["don_gia"], round((d["kl"] or 0) * (d["don_gia"] or 0)) if "bo_boi" not in d else 0, d["nguon"],
                    " → ".join(d["lich_su"])]
            for j, v in enumerate(vals, 1):
                put(w, i, j, v, fmt=QF if j == 7 else (NUMF if j in (8, 9) else None), text=True)
            if d["nguon"] != "HĐ gốc":
                w.cell(i, 10).fill = YELLOW if "bo_boi" not in d else RED
            i += 1
    bang("BOQ_DONG", f"SELECT b.*, p.so_pl, p.trang_thai FROM boq_dong b LEFT JOIN phu_luc p ON p.id = b.pl_id "
                     f"WHERE b.ma_hd IN ({ph}) ORDER BY b.ma_hd, IFNULL(p.ngay_hieu_luc, ''), b.id", ids,
         ["Mã HĐ", "Nguồn", "Trạng thái PL", "Thao tác", "Mã HM", "STT", "Nhóm", "Tên công việc", "ĐVT", "KL", "Đơn giá",
          "Thành tiền ghi", "Vị trí nguồn", "Ghi chú"], [8, 8, 10, 8, 8, 6, 26, 50, 8, 12, 14, 16, 40, 24],
         {10: QF, 11: NUMF, 12: NUMF},
         lambda r: [r["ma_hd"], r["so_pl"] or "HĐ gốc", r["trang_thai"] or "", r["thao_tac"], r["ma_hm"], r["stt"], r["nhom"],
                    r["ten"], r["dvt"], r["kl"], r["don_gia"], r["thanh_tien_ghi"], r["vi_tri_nguon"], r["ghi_chu"]])
    bang("DOT_THANH_TOAN", f"SELECT * FROM dot_thanh_toan WHERE ma_hd IN ({ph}) ORDER BY ma_hd, dot", ids,
         ["Mã HĐ", "Đợt", "Ngày nghiệm thu", "Ngày nhận HS", "GT kỳ", "GT lũy kế", "Giữ lại", "Thu hồi TƯ", "GT TT kỳ",
          "Nguồn", "Ngày nạp"], [8, 5, 12, 12, 16, 16, 14, 14, 16, 50, 17], {5: NUMF, 6: NUMF, 7: NUMF, 8: NUMF, 9: NUMF},
         lambda r: [r["ma_hd"], r["dot"], ngay_vn(r["ngay_nghiem_thu"]) or "(chưa nhập)", ngay_vn(r["ngay_nhan_hs"]),
                    r["gt_ky"], r["gt_luy_ke"], r["giu_lai"], r["thu_hoi_tu"], r["gt_tt_ky"], r["file_nguon"], r["ngay_nap"]])
    bang("KL_THANH_TOAN", f"SELECT * FROM kl_thanh_toan WHERE ma_hd IN ({ph}) ORDER BY ma_hd, dot, ma_hm", ids,
         ["Mã HĐ", "Đợt", "Mã HM", "KL kỳ", "KL lũy kế", "Đơn giá TT", "Vị trí nguồn"], [8, 5, 8, 12, 12, 14, 50],
         {4: QF, 5: QF, 6: NUMF})
    bang("CANH_BAO", f"SELECT c.*, p.so_pl, p.trang_thai FROM canh_bao c LEFT JOIN phu_luc p ON p.id = c.pl_id "
                     f"WHERE c.ma_hd IN ({ph}) ORDER BY c.ma_hd, c.id", ids,
         ["Mã HĐ", "Phụ lục", "Trạng thái PL", "Mã KT", "Mức độ", "Nội dung", "Thời điểm"], [8, 8, 10, 7, 8, 100, 17], None,
         lambda r: [r["ma_hd"], r["so_pl"] or "", r["trang_thai"] or "", r["ma_kt"], r["muc_do"], r["noi_dung"], r["thoi_diem"]])
    bang("NHAT_KY", f"SELECT * FROM nhat_ky WHERE ma_hd IN ({ph}) OR ma_hd = '' ORDER BY id", ids,
         ["#", "Thời điểm", "Người", "Thao tác", "Mã HĐ", "Chi tiết"], [5, 17, 12, 18, 8, 100])
    wb.save(out)
    return out


# ============================================================== CLI
def _in_kq(kq):
    for ma_kt, md, nd in kq:
        print(f"   [{md:^9}] {ma_kt} {nd}")


def main(argv=None):
    ap = argparse.ArgumentParser(description="Cơ sở dữ liệu hợp đồng - phụ lục - BOQ")
    ap.add_argument("--db", default=None)
    sub = ap.add_subparsers(dest="lenh", required=True)
    a = sub.add_parser("nap-master"); a.add_argument("files", nargs="+"); a.add_argument("--lam-lai", action="store_true"); a.add_argument("--ly-do", default="")
    a = sub.add_parser("nap-danh-muc"); a.add_argument("file")
    a = sub.add_parser("tao-phieu"); a.add_argument("hd"); a.add_argument("--out")
    a = sub.add_parser("nap-phu-luc"); a.add_argument("files", nargs="+"); a.add_argument("--chi-kiem-tra", action="store_true")
    a = sub.add_parser("huy-phu-luc"); a.add_argument("hd"); a.add_argument("so_pl"); a.add_argument("--ly-do", required=True)
    a = sub.add_parser("nap-dot"); a.add_argument("hd"); a.add_argument("file"); a.add_argument("--dot", type=int, required=True)
    a.add_argument("--ngay-nghiem-thu", required=True); a.add_argument("--ngay-nhan-hs"); a.add_argument("--ly-do", default="")
    a = sub.add_parser("kiem-dot"); a.add_argument("hd"); a.add_argument("file"); a.add_argument("--ngay-nghiem-thu", required=True)
    a.add_argument("--dot", type=int); a.add_argument("--out")
    a = sub.add_parser("ngay-nghiem-thu"); a.add_argument("hd"); a.add_argument("dot", type=int); a.add_argument("ngay"); a.add_argument("--ly-do", default="")
    a = sub.add_parser("boq"); a.add_argument("hd"); a.add_argument("--ngay")
    a = sub.add_parser("xuat-excel"); a.add_argument("hd", nargs="?"); a.add_argument("--ngay"); a.add_argument("--out")
    ar = ap.parse_args(argv)
    con = mo_csdl(ar.db)
    if ar.lenh == "nap-master":
        for f in ar.files:
            ma, nts, nb, nd, nk, kt = nap_master(con, f, ar.lam_lai, ar.ly_do)
            print(f"{Path(f).name} -> {ma}: {nts} tham số, {nb} dòng BOQ gốc, {nd} đợt, {nk} dòng KL lũy kế")
            _in_kq(kt)
    elif ar.lenh == "nap-danh-muc":
        for so_hd, ma, tt, lech in nap_danh_muc(con, ar.file):
            print(f"{so_hd} -> {ma}: {tt}")
            for l in lech:
                print("   lệch:", l)
    elif ar.lenh == "tao-phieu":
        out, so_pl = tao_phieu(con, ar.hd, ar.out)
        print(f"Đã tạo phiếu {so_pl}: {out}")
    elif ar.lenh == "nap-phu-luc":
        loi = False
        for f in ar.files:
            for sh, pl, kq, tt in nap_phu_luc(con, f, ar.chi_kiem_tra):
                print(f"\n{Path(f).name} / sheet {sh}: {pl['so_pl']} - {tt}")
                _in_kq(kq)
                loi |= tt == "KHÔNG GHI"
        return 1 if loi else 0
    elif ar.lenh == "huy-phu-luc":
        huy_phu_luc(con, ar.hd, ar.so_pl, ar.ly_do)
        print(f"Đã huỷ {ar.so_pl}.")
    elif ar.lenh == "nap-dot":
        ok, miss = nap_dot(con, ar.hd, ar.file, ar.dot, ar.ngay_nghiem_thu, ar.ngay_nhan_hs, ar.ly_do)
        print(f"Đợt {ar.dot}: {ok} hạng mục ghi vào CSDL; {len(miss)} không khớp BOQ hiệu lực")
        for it in miss[:20]:
            print(f"   dòng {it['dong']}: {it['ten'][:70]} ({it['dvt']})" + (f" - gợi ý {it['goi_y']}" if it["goi_y"] else ""))
    elif ar.lenh == "kiem-dot":
        out, b1, cnt, nthieu, pls, con_lai, dot, dot_truoc = kiem_dot(con, ar.hd, ar.file, ar.ngay_nghiem_thu, ar.out, ar.dot)
        print(f"Đợt {dot}, đợt trước trong CSDL: {dot_truoc}")
        print("Phụ lục áp dụng:", ", ".join(p["so_pl"] for p in pls) or "không",
              "| chưa áp dụng:", ", ".join(p["so_pl"] for p in con_lai) or "không")
        for t, a_, b_, kl, _ in b1:
            print(f"Bước 1: {kl}  (hồ sơ {a_ if a_ is None else f'{a_:,.0f}'} / hiệu lực {b_:,.0f})")
        print("Bước 3:", "; ".join(f"{k}: {n}" for k, n in cnt.most_common()) or "không lệch", f"| thiếu trong hồ sơ: {nthieu}")
        print("Kết quả:", out)
    elif ar.lenh == "ngay-nghiem-thu":
        ngay_nghiem_thu_dot(con, ar.hd, ar.dot, ar.ngay, ar.ly_do)
        print(f"Đã ghi ngày nghiệm thu đợt {ar.dot}: {ar.ngay}")
    elif ar.lenh == "boq":
        hd = tim_hd(con, ar.hd)
        hl, bo = boq_hieu_luc(con, hd["ma_hd"], ngay_iso(ar.ngay) if ar.ngay else None)
        for d in hl.values():
            print(f"{d['ma_hm']}  {d['ten'][:50]:50} {d['dvt']:>8} {_f(d['kl']):>12} {_f(d['don_gia']):>14}  {d['nguon']}")
        for d in bo.values():
            print(f"{d['ma_hm']}  {d['ten'][:50]:50}  ĐÃ BỎ theo {d['bo_boi']}")
    elif ar.lenh == "xuat-excel":
        print("Đã xuất:", xuat_excel(con, ar.hd, ar.ngay, ar.out))
    return 0


if __name__ == "__main__":
    sys.exit(main())
