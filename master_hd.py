# -*- coding: utf-8 -*-
"""
master_hd.py - Master file hợp đồng cho review hồ sơ thanh toán xây dựng.

Ý tưởng: mọi con số "gốc" của một đợt thanh toán đều phải lấy từ một nguồn cố định:
    1-2. KL hợp đồng, đơn giá      <- Hợp đồng + BOQ (+ phụ lục)
    3-4. KL kỳ trước, lũy kế trước  <- File thanh toán đợt trước (lần đầu: nhập tay)
    5.   KL kỳ này                  <- Biên bản nghiệm thu (+ bản vẽ)
    6.   Tham số thanh toán         <- Điều khoản hợp đồng
    7.   Số đã thanh toán           <- Đợt thanh toán trước
Master file gom 1, 2, 3-4, 6, 7 vào một file Excel. Hồ sơ đợt mới được đối chiếu với master.

Lệnh:
    python master_hd.py tao      <hop_dong.doc/.docx/.pdf> [--boq <file.xlsx>] [--out master.xlsx]
    python master_hd.py nap-dot  <master.xlsx> <ho_so_thanh_toan_dot_truoc.xlsx>
    python master_hd.py doi-chieu <master.xlsx> <ho_so_thanh_toan_dot_moi.xlsx> [--out ket_qua.xlsx]
    python master_hd.py ai-trich <hop_dong> [--an-danh] [--model qwen3.6:35b] [--host http://<IP máy AI>:11434]
Tuỳ chọn: --ai (nhờ model AI trích các điều khoản quy tắc không bắt được)
Địa chỉ máy AI và model mặc định lấy từ cau_hinh_may_ai.json (sửa ở tab "Máy AI" của cửa sổ 5_trich_hop_dong_AI.bat).

Chỉ đọc file hồ sơ, không sửa. Mọi giá trị trích tự động có trạng thái "Chưa xác nhận"
kèm trích dẫn nguyên văn, người review phải xác nhận trước khi dùng.
"""
import argparse
import datetime as dt
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata
import urllib.request
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# ============================================================== tiện ích
def nfc(s):
    return unicodedata.normalize("NFC", str(s or ""))


def bo_dau(s):
    s = nfc(s).replace("đ", "d").replace("Đ", "D")
    s = unicodedata.normalize("NFD", s)
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn")
    return re.sub(r"\s+", " ", s.lower()).strip()


def so(txt):
    """'1.234.567.890' -> 1234567890 ; '0,15' -> 0.15"""
    t = str(txt).strip()
    if re.fullmatch(r"\d{1,3}(\.\d{3})+", t):
        return float(t.replace(".", ""))
    return float(t.replace(".", "").replace(",", ".")) if "," in t else float(t)


def num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", ""))
    except ValueError:
        return None


# ============================================================== đọc hợp đồng
def doc_to_docx(path: Path) -> Path:
    """Chuyển .doc sang .docx: LibreOffice nếu có, nếu không thì MS Word (Windows)."""
    out = Path(tempfile.mkdtemp())
    cands = [shutil.which("soffice"), shutil.which("libreoffice"),
             r"C:\Program Files\LibreOffice\program\soffice.exe",
             r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"]
    for exe in cands:
        if exe and Path(exe).exists():
            subprocess.run([exe, "--headless", "--convert-to", "docx", "--outdir", str(out), str(path)],
                           stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
            res = out / (path.stem + ".docx")
            if res.exists():
                return res
    try:  # Windows + MS Word
        import win32com.client  # type: ignore
        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        d = word.Documents.Open(str(path.resolve()))
        res = out / (path.stem + ".docx")
        d.SaveAs2(str(res), FileFormat=16)
        d.Close()
        word.Quit()
        return res
    except Exception:
        pass
    raise RuntimeError("Không đọc được file .doc. Hãy mở bằng Word và lưu thành .docx, hoặc cài LibreOffice.")


def doc_hop_dong(path):
    """Trả về list đoạn văn [(điều, đoạn)] đã chuẩn hoá NFC."""
    path = Path(path)
    ext = path.suffix.lower()
    lines = []
    if ext == ".doc":
        path, ext = doc_to_docx(path), ".docx"
    if ext == ".docx":
        import docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        d = docx.Document(path)
        # Đọc theo đúng thứ tự trong văn bản: đoạn văn và bảng xen kẽ.
        # Bảng 1 cột (hay dùng để trình bày) -> tách thành từng dòng; bảng nhiều cột -> 'ô | ô'.
        for el in d.element.body.iterchildren():
            tag = el.tag.split("}")[-1]
            if tag == "p":
                lines.append(Paragraph(el, d).text)
            elif tag == "tbl":
                for row in Table(el, d).rows:
                    cells = []
                    for c in row.cells:
                        t = c.text.strip()
                        if t and t not in cells:
                            cells.append(t)
                    if len(cells) == 1:
                        lines += cells[0].splitlines()
                    elif cells:
                        lines.append(" | ".join(x.replace("\n", " / ") for x in cells))
    elif ext == ".pdf":
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            for pg in pdf.pages:
                lines += (pg.extract_text() or "").splitlines()
    else:
        lines = Path(path).read_text(encoding="utf-8", errors="ignore").splitlines()
    out, dieu = [], ""
    for ln in lines:
        ln = nfc(ln).strip()
        if not ln or re.search(r"\t\d+$", ln):  # dòng mục lục
            continue
        m = re.match(r"^Điều\s+(\d+|[IVXLC]+)\b\s*[\.:]?\s*(.*)", ln, re.I)  # cả "ĐIỀU 1", "Điều II"
        if m:
            dieu = f"Điều {m.group(1)}. {m.group(2)[:60]}".strip()
        out.append((dieu, re.sub(r"\s+", " ", ln)))
    return out


# ============================================================== quy tắc trích xuất
# Mỗi quy tắc: (mã, nhóm, tên tham số, đơn vị, regex, hàm lấy giá trị)
def _pct(m):
    return so(m.group("v")) / 100


def _txt(m):
    return m.group("v").strip(" .;:")


def _money(m):
    return so(m.group("v"))


RULES = [
    # ---- thông tin chung
    ("TC-01", "Thông tin chung", "Số hợp đồng", "", r"^Số\s*:\s*(?P<v>\S+)", _txt),
    ("TC-02", "Thông tin chung", "Ngày ký", "", r"Hôm nay,?\s*ngày\s*(?P<d>\d{1,2})\s*tháng\s*(?P<m>\d{1,2})\s*năm\s*(?P<y>\d{4})",
     lambda m: f"{int(m.group('d')):02d}/{int(m.group('m')):02d}/{m.group('y')}"),
    ("TC-03", "Thông tin chung", "Gói thầu", "", r"^Gói thầu\s*:\s*(?P<v>.+)", _txt),
    ("TC-04", "Thông tin chung", "Dự án", "", r"^Dự án\s*:?\s*(?P<v>.+)", _txt),
    ("TC-05", "Thông tin chung", "Địa điểm", "", r"^Địa điểm\s*:\s*(?P<v>.+)", _txt),
    ("TC-06", "Thông tin chung", "Bên A (giao thầu)", "", r"^Bên giao thầu\s*\((?P<lab>[^)]*)\)\s*:\s*(?P<v>.+)", _txt),
    ("TC-07", "Thông tin chung", "Bên B (nhận thầu)", "", r"^Bên nhận thầu\s*\((?P<lab>[^)]*)\)\s*:\s*(?P<v>.+)", _txt),
    ("TC-08", "Thông tin chung", "Hiệu lực hợp đồng", "", r"Hợp đồng này có hiệu lực\s*(?P<v>.+)", _txt),
    # ---- giá và loại hợp đồng
    ("TT-01", "Giá hợp đồng", "Loại hợp đồng", "",
     r"(Hợp đồng|HĐ)[^.]{0,40}?(?P<v>trọn gói|đơn giá cố định|đơn giá điều chỉnh|theo thời gian|theo chi phí cộng phí|kết hợp)", _txt),
    ("TT-02", "Giá hợp đồng", "Giá trị hợp đồng", "đồng",
     r"Giá trị (của )?hợp đồng\s*(là|:)?\s*(?P<v>\d{1,3}(\.\d{3}){2,})\s*(VNĐ|VND|đồng)", _money),
    ("TT-03", "Giá hợp đồng", "Giá trị HĐ đã gồm VAT", "",
     r"(?P<v>đã bao gồm|chưa bao gồm)\s*(các loại )?thuế\s*(VAT|GTGT)", lambda m: "Có" if "đã" in m.group("v") else "Không"),
    ("TT-04", "Giá hợp đồng", "Thuế suất GTGT", "%", r"(thuế\s*)?(VAT|GTGT)\s*(?P<v>\d+(,\d+)?)\s*%", _pct),
    ("TT-05", "Giá hợp đồng", "Cơ sở thanh toán khối lượng", "",
     r"(?P<v>(Thanh toán, quyết toán|Khối lượng thực tế thanh toán)[^.]{0,120})", _txt),
    # ---- tạm ứng
    ("TT-10", "Tạm ứng", "Tỷ lệ tạm ứng", "%",
     r"Tạm ứng( hợp đồng)?\s*:\s*(?P<v>Không áp dụng|\d+(,\d+)?\s*%)",
     lambda m: "Không áp dụng" if "không" in m.group("v").lower() else so(m.group("v").rstrip("% ")) / 100),
    ("TT-11", "Tạm ứng", "Thu hồi tạm ứng", "", r"(?P<v>[^.]*thu hồi[^.]*tạm ứng[^.]*)", _txt),
    ("TT-12", "Tạm ứng", "Bảo lãnh tạm ứng", "", r"Bảo lãnh (hoàn trả )?(tiền )?tạm ứng\s*:\s*(?P<v>[^.;]+)", _txt),
    # ---- thanh toán từng đợt
    ("TT-20", "Thanh toán đợt", "Tỷ lệ thanh toán mỗi đợt", "%",
     r"thanh toán cho (bên B|Nhà thầu)\s*(tối đa\s*)?(?P<v>\d+(,\d+)?)\s*%\s*giá trị", _pct),
    ("TT-21", "Thanh toán đợt", "Tỷ lệ giữ lại mỗi đợt", "%",
     r"giữ lại\s*(?P<v>\d+(,\d+)?)\s*%", _pct),
    ("TT-22", "Thanh toán đợt", "Thời hạn thanh toán", "ngày",
     r"Thời gian thanh toán\s*:\s*(Tối đa\s*)?(?P<v>\d+)\s*ngày", lambda m: int(m.group("v"))),
    # ---- quyết toán, bảo hành
    ("TT-30", "Quyết toán - bảo hành", "Tỷ lệ thanh toán khi quyết toán", "%",
     r"(?P<v>\d+(,\d+)?)\s*%\s*giá trị (theo )?hồ sơ quyết toán", _pct),
    ("TT-31", "Quyết toán - bảo hành", "Tỷ lệ bảo hành / bảo lãnh bảo hành", "%",
     r"bảo (lãnh )?(về )?bảo hành[^.%]{0,60}?(?P<v>\d+(,\d+)?)\s*%", _pct),
    ("TT-32", "Quyết toán - bảo hành", "Thời hạn bảo hành", "tháng",
     r"Thời hạn bảo hành (là\s*)?(?P<v>\d+)\s*tháng", lambda m: int(m.group("v"))),
    ("TT-33", "Quyết toán - bảo hành", "Bảo lãnh thực hiện hợp đồng", "",
     r"Bảo lãnh thực hiện Hợp đồng\s*:\s*(?P<v>[^.;]+)", _txt),
    # ---- điều chỉnh giá, phát sinh
    ("TT-40", "Điều chỉnh - phát sinh", "Điều chỉnh đơn giá", "",
     r"(?P<v>[^.]*điều chỉnh[^.]*(đơn giá|giá)[^.]*(biến động|>|trượt giá)[^.]*)", _txt),
    ("TT-41", "Điều chỉnh - phát sinh", "Ngưỡng điều chỉnh giá", "%",
     r"biến động[^.]{0,40}?>\s*(?P<v>\d+(,\d+)?)\s*%", _pct),
    ("TT-42", "Điều chỉnh - phát sinh", "Xử lý khối lượng phát sinh", "",
     r"(?P<v>[^.]*(hạng mục|khối lượng)[^.]*phát sinh[^.]*(thống nhất|phụ lục|biên bản)[^.]*)", _txt),
    # ---- tiến độ, phạt
    ("TT-50", "Tiến độ - phạt", "Tiến độ thực hiện", "",
     r"(Tiến độ thực hiện|Thời gian thực hiện hợp đồng)\s*:\s*(?P<v>.+)", _txt),
    ("TT-51", "Tiến độ - phạt", "Phạt vi phạm chất lượng (tối đa)", "%",
     r"phạt vi phạm tối đa\s*(?P<v>\d+(,\d+)?)\s*%\s*giá trị", _pct),
    ("TT-52", "Tiến độ - phạt", "Tạm giữ chậm tiến độ - mức 1", "%/ngày",
     r"Chậm\s*0?7 ngày đầu,?\s*tạm giữ[^%]*?(?P<v>\d+(,\d+)?)\s*%", _pct),
    ("TT-53", "Tiến độ - phạt", "Tạm giữ chậm tiến độ - mức 2", "%/ngày",
     r"Chậm từ ngày thứ\s*0?8[^%]*?(?P<v>\d+(,\d+)?)\s*%", _pct),
    ("TT-54", "Tiến độ - phạt", "Phạt chậm tổng tiến độ", "%/ngày",
     r"phạt\s*(?P<v>\d+(,\d+)?)\s*%\s*giá trị quyết toán[^.]*ngày chậm", _pct),
    ("TT-55", "Tiến độ - phạt", "Trần phạt tiến độ", "%",
     r"phạt tiến độ[^.]*không quá\s*(?P<v>\d+(,\d+)?)\s*%", _pct),
    ("TT-56", "Tiến độ - phạt", "Cách khấu trừ tiền phạt vào thanh toán", "",
     r"(?P<v>(?:[^.]|\.(?=\d))*phạt(?:[^.]|\.(?=\d))*(được trừ vào|khấu trừ vào|trừ vào)\s*(đợt|lần|kỳ)(?:[^.]|\.(?=\d))*)", _txt),
    ("TT-62", "Hồ sơ theo HĐ", "Số bộ hồ sơ thanh toán phải nộp", "bộ",
     r"Số lượng hồ sơ\s*:\s*(?P<v>\d+)\s*bộ", lambda m: int(m.group("v"))),
]


def _collect_list(paras, start_re, max_items=15):
    """Lấy các dòng '+ ...' ngay sau câu mở đầu (vd 'Hồ sơ thanh toán gồm')."""
    for i, (dieu, t) in enumerate(paras):
        if re.search(start_re, t, re.I):
            items = []
            for d2, t2 in paras[i + 1:i + 1 + max_items]:
                if t2.startswith("+"):
                    it = t2.lstrip("+ ").strip().rstrip(" ;.")
                    if not re.match(r"Số lượng hồ sơ", it, re.I):
                        items.append(it)
                elif items:
                    break
            if items:
                return dieu, t, items
    return None


def trich_xuat(paras, llm=None):
    """Chạy quy tắc trên từng đoạn; lấy kết quả đầu tiên trong phần thân hợp đồng."""
    rows = []
    # bỏ phần trang bìa: tìm đoạn 'Căn cứ' đầu tiên
    start = next((i for i, (_, t) in enumerate(paras) if t.startswith("Căn cứ")), 0)
    body = paras[max(0, start - 12):]
    for ma, nhom, ten, dv, pat, fn in RULES:
        hit = None
        for dieu, t in body:
            m = re.search(pat, t, re.I)
            if m:
                try:
                    hit = (fn(m), dieu, t, m)
                except (ValueError, IndexError):
                    continue
                break
        note = ""
        if hit and ma in ("TC-06", "TC-07"):
            lab = hit[3].group("lab").upper()
            want = "BÊN A" if ma == "TC-06" else "BÊN B"
            if want not in lab:
                note = f"Hợp đồng ghi '({lab})' - sai ký hiệu bên, cần lưu ý"
        rows.append(dict(ma=ma, nhom=nhom, ten=ten, gia_tri=hit[0] if hit else None, don_vi=dv,
                         cach="Tự động" if hit else "Không tìm thấy", dieu=hit[1] if hit else "",
                         trich=hit[2][:400] if hit else "", ghi_chu=note))
    by = {r["ma"]: r for r in rows}
    # suy ra giữ lại từ tỷ lệ thanh toán mỗi đợt
    if by["TT-21"]["gia_tri"] is None and isinstance(by["TT-20"]["gia_tri"], float):
        by["TT-21"].update(gia_tri=round(1 - by["TT-20"]["gia_tri"], 4), cach="Suy ra = 100% - TT-20",
                           dieu=by["TT-20"]["dieu"], trich=by["TT-20"]["trich"])
    # tạm ứng không áp dụng -> thu hồi không áp dụng
    if by["TT-10"]["gia_tri"] == "Không áp dụng" and by["TT-11"]["gia_tri"] is None:
        by["TT-11"].update(gia_tri="Không áp dụng", cach="Suy ra từ TT-10")
    # kiểm tra nhất quán trong văn bản
    for ma, msg in kiem_tra_nhat_quan(paras, by):
        tgt = by.get(ma)
        if tgt is not None:
            tgt["ghi_chu"] = (tgt["ghi_chu"] + "; " if tgt["ghi_chu"] else "") + msg
    # danh sách hồ sơ
    for ma, ten, pat in (("TT-60", "Hồ sơ thanh toán mỗi đợt", r"Hồ sơ thanh toán gồm"),
                         ("TT-61", "Hồ sơ quyết toán", r"Hồ sơ (thanh )?quyết toán bao gồm")):
        got = _collect_list(body, pat)
        rows.append(dict(ma=ma, nhom="Hồ sơ theo HĐ", ten=ten, gia_tri="; ".join(got[2]) if got else None,
                         don_vi="", cach="Tự động" if got else "Không tìm thấy", dieu=got[0] if got else "",
                         trich=got[1][:400] if got else "", ghi_chu="", items=got[2] if got else []))
    # local AI cho các mục chưa bắt được
    if llm and llm.ok():
        miss = [r for r in rows if r["gia_tri"] is None]
        if miss:
            try:
                ai = llm.extract(paras, miss)
                for r in miss:
                    v = ai.get(r["ma"])
                    if v and v.get("gia_tri") not in (None, "", "null"):
                        r.update(gia_tri=v.get("gia_tri"), cach="Local AI", trich=str(v.get("trich", ""))[:400],
                                 dieu=v.get("dieu", ""))
            except Exception as e:
                print("Local AI lỗi:", e)
    return rows


AI_SYSTEM = """Bạn là chuyên viên kiểm soát hợp đồng xây dựng tại Việt Nam. Bạn đọc hợp đồng và trích tham số chính xác tuyệt đối để phục vụ kiểm tra hồ sơ thanh toán.

QUY TẮC BẮT BUỘC
1. Chỉ dùng thông tin có trong văn bản hợp đồng được cung cấp. Không suy đoán, không dùng kiến thức bên ngoài, không dùng quy định pháp luật để điền thay.
2. Mỗi tham số phải có "trich_dan": chép NGUYÊN VĂN câu trong hợp đồng làm căn cứ, không sửa chữ, không tóm tắt.
3. Không tìm thấy: gia_tri = "KHÔNG TÌM THẤY", trich_dan = "".
4. Hợp đồng ghi rõ "Không áp dụng": gia_tri = "Không áp dụng", vẫn ghi trich_dan.
5. Giá trị phải suy ra (không ghi trực tiếp): ghi vào ghi_chu "Suy ra: <cách suy ra>".
6. Định dạng: tỷ lệ ghi 10%; số tiền ghi số nguyên không dấu phân cách (1234567890); ngày ghi dd/mm/yyyy; thời hạn ghi số, đơn vị ở don_vi.
7. dieu_khoan: số Điều và tên Điều chứa câu trích dẫn.
8. Ghi vào ghi_chu mọi điểm bất thường: ghi sai tên/ký hiệu bên, cùng một nội dung ghi hai con số khác nhau, tỷ lệ mâu thuẫn.
9. Bỏ qua mục lục. Nếu thông tin có ở cả trang bìa và phần thân, lấy phần thân.
10. Chuỗi dạng [BEN_A], [BEN_B], [CA_NHAN], [SO] là thông tin đã được ẩn danh, giữ nguyên khi trích."""

AI_FORMAT = ('Chỉ in một đối tượng JSON hợp lệ, không có ``` và không viết gì thêm, dạng: '
             '{"TC-01": {"gia_tri": "...", "don_vi": "", "dieu_khoan": "...", "trich_dan": "...", "ghi_chu": ""}, ...}. '
             "Đủ tất cả các mã trong danh sách.")

AI_KEEP = r"giá|thanh toán|tạm ứng|bảo lãnh|bảo hành|điều chỉnh|tiến độ|phạt|hiệu lực|quyết toán|nghiệm thu"


def rut_gon(paras, keep=AI_KEEP):
    """Giữ phần mở đầu (trước Điều 1) và các Điều liên quan thanh toán. Bỏ mục lục, phần chữ ký."""
    out = []
    for d, t in paras:
        if not d or re.search(keep, d, re.I):
            if " | " in t and re.search(r"Mã số thuế|Tài khoản|Điện thoại|Đại diện", t):
                continue  # bảng thông tin chữ ký
            out.append((d, t))
    return out


def an_danh(text, rows):
    """Che tên hai bên, tên người, mã số thuế, số tài khoản, số điện thoại trước khi gửi lên máy chủ AI."""
    by = {r["ma"]: r for r in rows}
    for ma, tag in (("TC-06", "[BEN_A]"), ("TC-07", "[BEN_B]")):
        v = by.get(ma, {}).get("gia_tri")
        if isinstance(v, str) and len(v) > 5:
            text = re.sub(re.escape(v), tag, text, flags=re.I)
    text = re.sub(r"\b(Ông|Bà|Ông/Bà)\s*:?\s*(([A-ZĐ][\wÀ-ỹ]*)(\s+[A-ZĐ][\wÀ-ỹ]*){1,4})", r"\1 [CA_NHAN]", text)
    text = re.sub(r"\b\d{9,}\b", "[SO]", text)                       # MST, số tài khoản
    text = re.sub(r"\b0\d{2,3}[ .]?\d{3,4}[ .]?\d{3,4}\b", "[SO]", text)  # điện thoại
    return text


SO_CHU = {"một": 1, "hai": 2, "ba": 3, "bốn": 4, "tư": 4, "năm": 5, "sáu": 6, "bảy": 7, "tám": 8, "chín": 9,
          "mười": 10, "mười một": 11, "mười hai": 12, "mười lăm": 15, "hai mươi": 20, "ba mươi": 30}


def kiem_tra_nhat_quan(paras, by):
    """
    Các lỗi soạn thảo hay gặp mà quy tắc có thể phát hiện:
    - Điều 'định nghĩa' gán tên bên khác với phần các bên tham gia.
    - Số và chữ trong ngoặc không khớp, ví dụ '03 (bốn)'.
    - Cùng văn bản vừa yêu cầu vừa không yêu cầu xác nhận của Bên B cho biên bản vi phạm.
    """
    out = []
    ten_a, ten_b = by.get("TC-06", {}).get("gia_tri"), by.get("TC-07", {}).get("gia_tri")
    for dieu, t in paras:
        m = re.search(r"“Bên giao thầu”\s*\(viết tắt là Bên A\)\s*là\s*(.+?)\s+(như|được nêu)", t)
        if m and ten_a and bo_dau(m.group(1)) != bo_dau(ten_a):
            out.append(("TC-06", f"MÂU THUẪN: {dieu.split('.')[0]} định nghĩa Bên giao thầu là '{m.group(1)[:60]}', "
                                 f"khác phần các bên ('{ten_a[:60]}')"))
        for m in re.finditer(r"\b(\d{1,2})\s*\(\s*([a-zà-ỹ ]{2,12}?)\s*\)", t, re.I):
            chu = m.group(2).strip().lower()
            if chu in SO_CHU and SO_CHU[chu] != int(m.group(1)):
                out.append(("TC-08", f"Số và chữ không khớp tại {dieu.split('.')[0]}: '{m.group(0)}'"))
    full = "\n".join(t for _, t in paras)
    can = re.search(r"có xác nhận của (đại diện )?Bên B", full, re.I)
    khong = re.search(r"không cần (việc )?xác nhận (từ|của) Bên B", full, re.I)
    if can and khong:
        out.append(("TT-56", "MÂU THUẪN: một chỗ yêu cầu biên bản vi phạm có xác nhận của Bên B, "
                             "chỗ khác ghi không cần xác nhận của Bên B"))
    return out


# ============================================================== cấu hình máy AI (dùng chung mọi công cụ)
_CH = Path(__file__).resolve().parent
# app_kiem_tra_ho_so/ dùng chung file cấu hình ở thư mục dự án (thư mục cha)
CAU_HINH_AI = next((d / "cau_hinh_may_ai.json" for d in (_CH, _CH.parent) if (d / "cau_hinh_may_ai.json").exists()),
                   _CH / "cau_hinh_may_ai.json")
MAC_DINH_AI = {"may_ai": "http://localhost:11434", "model": "qwen3:8b", "api": "ollama",
               "an_danh": True, "gui_ca_hop_dong": False, "timeout_giay": 1800}


def cau_hinh_ai():
    """Đọc cau_hinh_may_ai.json (không đẩy lên Git). Thiếu file thì dùng mặc định localhost."""
    cfg = dict(MAC_DINH_AI)
    if CAU_HINH_AI.exists():
        try:
            cfg.update(json.loads(CAU_HINH_AI.read_text(encoding="utf-8-sig")))
        except Exception as e:
            print(f"Không đọc được {CAU_HINH_AI.name} ({e}); dùng cấu hình mặc định.")
    cfg["may_ai"] = str(cfg["may_ai"]).rstrip("/")
    if not cfg["may_ai"].startswith("http"):
        cfg["may_ai"] = "http://" + cfg["may_ai"]
    return cfg


class LocalLLM:
    """
    Gọi model AI qua API. Hai kiểu:
      api="ollama": POST {host}/api/chat           (Ollama gốc; tắt 'think' để model không suy luận dài)
      api="openai": POST {host}/v1/chat/completions (vLLM, LM Studio, Ollama chế độ OpenAI)
    Máy AI trong mạng LAN: host = http://<IP máy AI>:11434 (đặt trong cau_hinh_may_ai.json).
    """

    def __init__(self, model="qwen3.6:35b", host="http://localhost:11434", api="ollama", num_ctx=32768, timeout=900):
        self.model, self.host, self.api, self.num_ctx, self.timeout = model, host.rstrip("/"), api, num_ctx, timeout
        self.stats = {}

    def ok(self):
        try:
            urllib.request.urlopen(self.host + ("/api/tags" if self.api == "ollama" else "/v1/models"), timeout=5)
            return True
        except Exception:
            return False

    def chat(self, system, user, json_mode=True):
        import time
        t0 = time.time()
        if self.api == "ollama":
            body = {"model": self.model, "stream": False, "think": False,
                    "options": {"temperature": 0, "num_ctx": self.num_ctx},
                    "messages": [{"role": "system", "content": system}, {"role": "user", "content": user}]}
            if json_mode:
                body["format"] = "json"
            url = self.host + "/api/chat"
        else:
            body = {"model": self.model, "temperature": 0,
                    "messages": [{"role": "system", "content": system}, {"role": "user", "content": user}]}
            url = self.host + "/v1/chat/completions"
        def goi(b):
            req = urllib.request.Request(url, data=json.dumps(b).encode(), headers={"Content-Type": "application/json"})
            with urllib.request.urlopen(req, timeout=self.timeout) as r:
                return json.loads(r.read())
        def loi(e):
            try:
                chi_tiet = e.read().decode("utf-8", "replace")[:300]
            except Exception:
                chi_tiet = ""
            return RuntimeError(f"Máy AI báo lỗi {e.code}: {chi_tiet or e.reason}")
        try:
            res = goi(body)
        except urllib.request.HTTPError as e:
            # model không có chế độ suy luận (gemma3, llama...) báo lỗi khi gửi "think" -> gửi lại không có "think"
            if self.api == "ollama" and "think" in body and e.code in (400, 500):
                body.pop("think")
                try:
                    res = goi(body)
                except urllib.request.HTTPError as e2:
                    raise loi(e2) from None
            else:
                raise loi(e) from None
        if self.api == "ollama":
            content = res["message"]["content"]
            self.stats = dict(token_vao=res.get("prompt_eval_count"), token_ra=res.get("eval_count"))
        else:
            content = res["choices"][0]["message"]["content"]
            u = res.get("usage", {})
            self.stats = dict(token_vao=u.get("prompt_tokens"), token_ra=u.get("completion_tokens"))
        self.stats["giay"] = round(time.time() - t0, 1)
        return content

    def extract_all(self, text, params):
        ds = "\n".join(f"{ma} | {nhom} | {ten}" + (f" ({dv})" if dv else "") for ma, nhom, ten, dv in params)
        user = (f"Đọc hợp đồng dưới đây và trích đủ các tham số trong DANH SÁCH, giữ đúng mã và thứ tự.\n\n"
                f"DANH SÁCH THAM SỐ (mã | nhóm | tham số)\n{ds}\n\nĐỊNH DẠNG ĐẦU RA\n{AI_FORMAT}\n\nHỢP ĐỒNG\n{text}")
        raw = self.chat(AI_SYSTEM, user)
        m = re.search(r"\{.*\}", raw, re.S)
        return json.loads(m.group(0) if m else raw), raw

    def extract(self, paras, miss):
        """Dùng trong 'tao --ai': chỉ hỏi các mục quy tắc không bắt được."""
        text = "\n".join(t for _, t in rut_gon(paras))
        got, _ = self.extract_all(text, [(r["ma"], r["nhom"], r["ten"], r["don_vi"]) for r in miss])
        return {k: dict(gia_tri=v.get("gia_tri"), dieu=v.get("dieu_khoan", ""), trich=v.get("trich_dan", ""))
                for k, v in got.items() if isinstance(v, dict) and v.get("gia_tri") not in (None, "", "KHÔNG TÌM THẤY")}


def _chuan(v):
    """Chuẩn hoá giá trị để so quy tắc với AI: số, %, chữ bỏ dấu."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return round(float(v), 6)
    t = nfc(v).strip()
    if re.fullmatch(r"-?\d+([.,]\d+)?\s*%", t):
        return round(so(t.rstrip("% ").strip()) / 100, 6)
    if re.fullmatch(r"\d{1,3}([.,]\d{3})+|\d+", t):
        return round(float(re.sub(r"[.,]", "", t)), 6)
    m = re.fullmatch(r"(\d+)\s*(ngày|tháng|bộ)", t)
    if m:
        return float(m.group(1))
    return bo_dau(t)


def so_sanh_ai(rows, ai, full_text, anon_text=""):
    """So kết quả quy tắc với AI; kiểm tra trích dẫn AI có thật trong hợp đồng (bản gốc hoặc bản đã ẩn danh)."""
    full = bo_dau(full_text) + "\n" + bo_dau(anon_text)
    out = []
    for r in rows:
        a = ai.get(r["ma"], {}) if isinstance(ai.get(r["ma"]), dict) else {}
        av, qt = a.get("gia_tri"), a.get("trich_dan") or ""
        if isinstance(av, str) and av.strip().upper() == "KHÔNG TÌM THẤY":
            av = None
        c1, c2 = _chuan(r["gia_tri"]), _chuan(av)
        if c1 is None and c2 is None:
            khop = "Cả hai không thấy"
        elif c1 is None or c2 is None:
            khop = "Chỉ quy tắc thấy" if c2 is None else "Chỉ AI thấy"
        elif isinstance(c1, float) and isinstance(c2, float):
            khop = "Khớp" if abs(c1 - c2) <= max(1e-6, abs(c1) * 1e-9) else "Lệch"
        else:
            s1, s2 = str(c1), str(c2)
            khop = "Khớp" if (s1 == s2 or s1 in s2 or s2 in s1) else "Khác cách ghi"
        q = bo_dau(qt)
        if not q:
            that = ""
        elif q in full:
            that = "Có"
        else:
            # cho phép sai khác nhỏ: xét 60 ký tự đầu
            that = "Gần đúng" if q[:60] in full else "KHÔNG - AI viết lại hoặc bịa"
        out.append(dict(ma=r["ma"], ten=r["ten"], quy_tac=r["gia_tri"], ai=a.get("gia_tri"), khop=khop,
                        trich_ai=qt, that=that, dieu_ai=a.get("dieu_khoan", ""), ghi_chu_ai=a.get("ghi_chu", ""),
                        trich_qt=r["trich"]))
    return out


def ghi_so_sanh(out, hop_dong, cmp, stats, model, n_chars, an):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "So sanh"
    put(ws, 1, 1, "SO SÁNH TRÍCH XUẤT HỢP ĐỒNG: QUY TẮC vs AI", Font(name="Arial", bold=True, size=12))
    put(ws, 2, 1, f"Hợp đồng: {Path(hop_dong).name} | Model: {model} | Văn bản gửi AI: {n_chars:,} ký tự"
                  f"{' (đã ẩn danh)' if an else ''} | Token vào: {stats.get('token_vao')} | Token ra: {stats.get('token_ra')}"
                  f" | Thời gian: {stats.get('giay')} giây", text=True)
    header(ws, 4, ["Mã", "Tham số", "Theo quy tắc", "Theo AI", "So sánh", "Trích dẫn AI có trong HĐ?",
                   "Trích dẫn của AI", "Điều khoản (AI)", "Ghi chú của AI", "Trích dẫn (quy tắc)"],
           [8, 34, 24, 24, 16, 20, 60, 30, 36, 60])
    ws.freeze_panes = "C5"
    fills = {"Khớp": GREEN, "Cả hai không thấy": None, "Lệch": RED, "Chỉ AI thấy": YELLOW,
             "Chỉ quy tắc thấy": YELLOW, "Khác cách ghi": YELLOW}
    for i, x in enumerate(cmp, 5):
        vals = [x["ma"], x["ten"], x["quy_tac"], x["ai"], x["khop"], x["that"], x["trich_ai"], x["dieu_ai"],
                x["ghi_chu_ai"], x["trich_qt"]]
        for j, v in enumerate(vals, 1):
            put(ws, i, j, v, wrap=j in (3, 4, 7, 9, 10), text=True)
        f = fills.get(x["khop"])
        if f:
            ws.cell(i, 5).fill = f
        ws.cell(i, 6).fill = RED if x["that"].startswith("KHÔNG") else (GREEN if x["that"] == "Có" else PatternFill())
    from collections import Counter
    r = len(cmp) + 7
    put(ws, r, 1, "Tổng hợp", BOLD)
    for k, n in Counter(x["khop"] for x in cmp).most_common():
        r += 1
        put(ws, r, 2, k); put(ws, r, 3, n)
    for k, n in Counter(x["that"] or "(không có trích dẫn)" for x in cmp).most_common():
        r += 1
        put(ws, r, 2, "Trích dẫn AI: " + k, text=True); put(ws, r, 3, n)
    wb.save(out)


def doc_dap_an(path):
    gold = {}
    for line in open(path, encoding="utf-8-sig"):
        c = line.rstrip("\r\n").split("\t")
        if c and re.fullmatch(r"(TC|TT)-\d+", c[0].strip()):
            c += [""] * (8 - len(c))
            gold[c[0].strip()] = dict(gia_tri=c[3], don_vi=c[4], dieu=c[5], trich=c[6], ghi_chu=c[7])
    return gold


def _so_trong(txt, pct):
    """Các số trong chuỗi đáp án; nếu đơn vị là % thì chia 100."""
    out = []
    for m in re.finditer(r"\d+(?:[.,]\d+)*", str(txt)):
        t = m.group(0)
        v = float(t.replace(".", "")) if re.fullmatch(r"\d{1,3}(\.\d{3})+", t) else float(t.replace(",", "."))
        out.append(v / 100 if pct else v)
    return out


def cham_diem(ket_qua, gold):
    """
    Chấm tự động (cần người soát lại): ket_qua = {mã: {gia_tri, trich, ghi_chu}}.
    Số: so số đầu tiên; đáp án có thêm số khác mà kết quả không nêu -> Một phần.
    Chữ: tỷ lệ từ của đáp án có trong (giá trị + trích dẫn) >= 0,8 Đúng; >= 0,4 Một phần.
    """
    res = []
    for ma, g in gold.items():
        k = ket_qua.get(ma, {})
        v = k.get("gia_tri")
        if v in (None, "") or (isinstance(v, str) and v.strip().upper() in ("KHÔNG TÌM THẤY", "(THIẾU)")):
            res.append((ma, "Sai", "Không có kết quả")); continue
        pct = g["don_vi"].strip().startswith("%")
        gnums = _so_trong(g["gia_tri"], pct)
        vv = _chuan(v)
        if isinstance(vv, float) and gnums:
            if abs(vv - gnums[0]) < 1e-9 or abs(vv * 100 - gnums[0] * (100 if pct else 1)) < 1e-9:
                extra = [x for x in gnums[1:] if not any(abs(x - y) < 1e-9 for y in _so_trong(f"{k.get('trich','')} {v}", pct))]
                res.append((ma, "Một phần" if extra else "Đúng", "Số khớp" + ("; thiếu thông tin phụ trong đáp án" if extra else "")))
            else:
                res.append((ma, "Sai", f"Số khác: {v} so với {g['gia_tri']}"))
            continue
        gw = set(re.findall(r"[a-z0-9]+", bo_dau(g["gia_tri"])))
        kw = set(re.findall(r"[a-z0-9]+", bo_dau(f"{v} {k.get('trich', '')}")))
        stop = {"va", "cua", "cac", "the", "theo", "la", "tu", "ngay", "ben"}
        gw -= stop
        r = len(gw & kw) / len(gw) if gw else 0
        res.append((ma, "Đúng" if r >= 0.8 else ("Một phần" if r >= 0.4 else "Sai"), f"Tỷ lệ từ khớp {r:.0%}"))
    return res


# ============================================================== đọc bảng chi tiết (mọi mẫu)
def _hdr_text(ws, rows, col):
    return " ".join(bo_dau(ws.cell(r, col).value) for r in rows if ws.cell(r, col).value not in (None, ""))


def tim_bang_chi_tiet(wb_values):
    """
    Tìm bảng xác định giá trị (KL x đơn giá) trong workbook bất kỳ, dựa vào tiêu đề cột:
    cột 'Tên công việc/tác', 'Đơn vị', 'Đơn giá'; cột bên trái đơn giá là khối lượng,
    bên phải là giá trị; phân biệt HĐ / kỳ trước / kỳ này / lũy kế theo chữ trong tiêu đề.
    """
    best = None
    for ws in wb_values.worksheets:
        if ws.max_row < 5:
            continue
        for r in range(1, min(ws.max_row, 40) + 1):
            cols = {}
            for c in range(1, min(ws.max_column, 40) + 1):
                t = bo_dau(ws.cell(r, c).value)
                if re.search(r"^ten cong (viec|tac)|^noi dung cong viec|^ten (hang muc|cong viec)", t):
                    cols["ten"] = c
                elif re.search(r"^don gia", t):
                    cols["dg"] = c
            if "ten" not in cols or "dg" not in cols:
                continue
            hdr_rows = range(r, r + 4)
            m = {"ten": cols["ten"], "dg": cols["dg"]}
            for c in range(1, min(ws.max_column, 40) + 1):
                t = _hdr_text(ws, hdr_rows, c)
                # tiêu đề nhóm (vd 'Khối lượng' gộp ô) chỉ nằm ở ô đầu; ghép với ô dưới
                if not t:
                    continue
                if c < cols["ten"] and re.search(r"\bstt\b|\btt\b", t):
                    m.setdefault("stt", c)
                if re.search(r"^don vi|^d\.? ?vi|^dvt", t) and c != cols["ten"]:
                    m.setdefault("dvt", c)
                side = "kl" if c < cols["dg"] else ("gt" if c > cols["dg"] else None)
                if not side or c <= cols["ten"]:
                    continue
                if re.search(r"hop dong", t) and not re.search(r"ky truoc|ky nay", t):
                    m.setdefault(side + "_hd", c)
                elif re.search(r"ky truoc", t):
                    m.setdefault(side + "_kt", c)
                elif re.search(r"luy ke.*ky nay|het ky nay", t):
                    m.setdefault(side + "_lk", c)
                elif re.search(r"ky nay", t):
                    m.setdefault(side + "_kn", c)
            if "dvt" not in m or not any(k in m for k in ("kl_hd", "kl_lk")):
                continue
            # dòng dữ liệu: bắt đầu sau dòng đánh số cột, dừng ở dòng 'cộng'
            r0 = r + 1
            while r0 < r + 6 and not isinstance(num(ws.cell(r0, m["dg"]).value), float) \
                    and not str(ws.cell(r0, m["ten"]).value or "").strip()[:1].isalpha():
                r0 += 1
            items = []
            for rr in range(r0, ws.max_row + 1):
                name = ws.cell(rr, m["ten"]).value
                nb = bo_dau(name)
                dvt0 = ws.cell(rr, m["dvt"]).value
                low = nfc(name).strip().lower()
                if dvt0 in (None, "", 0) and re.match(r"^(cộng|tổng cộng|thuế( gtgt| vat)?|tổng giá trị)(\s|\(|:|$)", low):
                    break  # dòng cộng / thuế cuối bảng (không nhầm với hạng mục 'Cống ...' vì có đơn vị)
                if not name or re.fullmatch(r"[\[\]\-\d\s]+", str(name)):
                    continue
                dvt = ws.cell(rr, m["dvt"]).value
                it = dict(dong=rr, stt=str(ws.cell(rr, m["stt"]).value or "").strip() if "stt" in m else "",
                          ten=nfc(name).strip(), dvt=nfc(dvt).strip() if dvt not in (None, 0) else "",
                          don_gia=num(ws.cell(rr, m["dg"]).value))
                for k in ("kl_hd", "kl_kt", "kl_kn", "kl_lk", "gt_hd", "gt_kt", "gt_kn", "gt_lk"):
                    it[k] = num(ws.cell(rr, m[k]).value) if k in m else None
                items.append(it)
            n_item = sum(1 for i in items if i["dvt"])
            score = n_item * 10 + len([k for k in m if k.startswith(("kl_", "gt_"))])
            if not best or score > best["score"]:
                best = dict(sheet=ws.title, hdr=r, cols=m, items=items, score=score)
            break
    if not best:
        raise RuntimeError("Không tìm thấy bảng chi tiết (cần cột 'Tên công việc', 'Đơn vị', 'Đơn giá').")
    # gán nhóm (dòng không có đơn vị là tiêu đề nhóm)
    grp = []
    for it in best["items"]:
        if not it["dvt"]:
            if re.match(r"^[A-Z]{1,3}\.?$|^[IVX]+\.?$", it["stt"] or "") or not grp:
                grp = [it["ten"]] if re.match(r"^[A-Z]\.?$", it["stt"] or "") or not grp else grp[:1] + [it["ten"]]
            else:
                grp = grp[:1] + [it["ten"]]
        it["nhom"] = " / ".join(grp)
    seen = {}
    for it in best["items"]:
        if it["dvt"]:
            base = khoa(it)
            seen[base] = seen.get(base, 0) + 1
            it["khoa"] = base if seen[base] == 1 else f"{base}#{seen[base]}"
    return best


def khoa(it):
    """Khoá so khớp hạng mục giữa các file: nhóm + tên + đơn vị (bỏ dấu, bỏ khoảng trắng thừa)."""
    return f"{bo_dau(it.get('nhom', ''))}|{bo_dau(it['ten'])[:120]}|{bo_dau(it['dvt'])}"


# ============================================================== tìm số tổng hợp (mọi mẫu)
LABELS = {
    "gt_hd": r"gia tri hop dong( goc| \+ bo sung| \+ plhd)?",
    "gt_lk_kt": r"luy ke thuc hien (den het|het) ky truoc",
    "gt_kn": r"gia tri (thuc hien|hoan thanh) (ky nay|dot nay)",
    "gt_lk": r"luy ke thuc hien (den het|het) ky nay",
    "giu_lai": r"giu lai",
    "thu_hoi_tu": r"hoan tra tam ung|thu hoi (gia tri )?(da )?tam ung",
    "tt_ky_nay": r"gia tri (thanh toan|de nghi thanh toan) ky nay",
}


def tim_so_tong_hop(wb_values):
    """Quét các sheet tổng hợp tìm số theo nhãn; lấy số đầu tiên bên phải nhãn."""
    found = {}
    for ws in wb_values.worksheets:
        if ws.sheet_state != "visible":
            continue
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 80), max_col=min(ws.max_column, 12)):
            for cell in row:
                t = bo_dau(cell.value) if isinstance(cell.value, str) else ""
                if not t or len(t) > 140:
                    continue
                for k, pat in LABELS.items():
                    if k in found or not re.search(pat, t):
                        continue
                    if k == "gt_hd" and "luy ke" in t:
                        continue
                    for c in range(cell.column + 1, min(ws.max_column, cell.column + 6) + 1):
                        v = ws.cell(cell.row, c).value
                        if isinstance(v, (int, float)) and abs(v) > 1:
                            found[k] = dict(gia_tri=float(v), o=f"{ws.title}!{ws.cell(cell.row, c).coordinate}",
                                            nhan=nfc(cell.value)[:80])
                            break
    return found


# ============================================================== ghi master
THIN = Side(style="thin", color="BFC5C2")
HF = PatternFill("solid", fgColor="1F4A7A")
HFONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
BASE = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
BLUE = Font(name="Arial", size=10, color="0000FF")
YELLOW = PatternFill("solid", fgColor="FFF2CC")
RED = PatternFill("solid", fgColor="F8D7D3")
GREEN = PatternFill("solid", fgColor="DDEFE3")
NUMF = '#,##0;(#,##0);"-"'
QF = '#,##0.000;(#,##0.000);"-"'


def put(ws, r, c, v, font=BASE, fmt=None, fill=None, wrap=False, text=False):
    """Ghi ô. Chuỗi bắt đầu bằng '=' là công thức, trừ khi text=True (nội dung trích từ tài liệu)."""
    cl = ws.cell(r, c, v)
    if text and isinstance(v, str) and v.startswith("="):
        cl.data_type = "s"
    cl.font = font
    if fmt:
        cl.number_format = fmt
    if fill:
        cl.fill = fill
    if wrap:
        cl.alignment = Alignment(wrap_text=True, vertical="top")
    return cl


def header(ws, row, names, widths):
    for j, (h, w) in enumerate(zip(names, widths), 1):
        cl = ws.cell(row, j, h)
        cl.font, cl.fill = HFONT, HF
        cl.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[row].height = 32
    ws.freeze_panes = ws.cell(row + 1, 1)


def ghi_master(out, contract_file, rows, boq=None, boq_src=""):
    wb = openpyxl.Workbook()
    # ---- Hướng dẫn
    ws = wb.active
    ws.title = "Huong dan"
    guide = [
        "MASTER FILE HỢP ĐỒNG",
        f"Hợp đồng nguồn: {Path(contract_file).name}    Lập lúc: {dt.datetime.now():%d/%m/%Y %H:%M}",
        "",
        "Master file là nguồn gốc duy nhất cho các dữ liệu gốc khi review hồ sơ thanh toán:",
        "  1-2. KL hợp đồng, đơn giá      -> sheet BOQ (từ Bảng chi tiết giá trị HĐ, cập nhật khi có phụ lục)",
        "  3-4. KL kỳ trước, lũy kế trước -> sheet KL ky truoc (tự nạp từ file thanh toán đợt trước; đợt đầu: nhập tay)",
        "  5.   KL kỳ này                 -> không lưu ở đây; đối chiếu biên bản nghiệm thu, bản vẽ",
        "  6.   Tham số thanh toán        -> sheet Dieu khoan TT (trích từ hợp đồng)",
        "  7.   Số đã thanh toán          -> sheet Lich su TT (tự nạp từ đợt trước; đối chiếu UNC)",
        "",
        "Quy ước màu: nền vàng = ô người review nhập / xác nhận; chữ xanh = số nhập tay.",
        "Mọi giá trị trích tự động có trạng thái 'Chưa xác nhận'. Đọc cột 'Trích dẫn', sửa nếu sai, rồi chọn 'Đã xác nhận'.",
        "Chỉ dùng master để đối chiếu khi các tham số bắt buộc đã được xác nhận.",
        "",
        "Cập nhật master: (1) khi ký phụ lục -> thêm vào sheet Phu luc và BOQ; (2) sau mỗi đợt được duyệt -> chạy",
        "   python master_hd.py nap-dot master.xlsx <file đợt vừa duyệt.xlsx>",
    ]
    for i, t in enumerate(guide, 1):
        put(ws, i, 1, t, BOLD if i == 1 else BASE)
    ws.column_dimensions["A"].width = 120
    dv_status = DataValidation(type="list", formula1='"Chưa xác nhận,Đã xác nhận,Đã sửa tay,Không áp dụng"', allow_blank=True)

    # ---- Thông tin HĐ + Điều khoản TT
    for title, pref in (("Thong tin HD", "TC"), ("Dieu khoan TT", "TT")):
        ws = wb.create_sheet(title)
        header(ws, 1, ["Mã", "Nhóm", "Tham số", "Giá trị", "Đơn vị", "Cách lấy", "Điều khoản", "Trích dẫn nguyên văn",
                       "Trạng thái", "Người xác nhận", "Ghi chú"], [8, 18, 34, 22, 8, 18, 30, 80, 15, 16, 36])
        ws.add_data_validation(dv_status)
        r = 2
        for row in rows:
            if not row["ma"].startswith(pref):
                continue
            v = row["gia_tri"]
            fmt = "0.00%" if row["don_vi"] in ("%", "%/ngày") and isinstance(v, float) else (NUMF if row["don_vi"] == "đồng" else None)
            vals = [row["ma"], row["nhom"], row["ten"], v, row["don_vi"], row["cach"], row["dieu"], row["trich"],
                    "Chưa xác nhận", "", row["ghi_chu"]]
            for j, x in enumerate(vals, 1):
                put(ws, r, j, x, fmt=fmt if j == 4 else None, wrap=j in (4, 8, 11), text=True)
            ws.cell(r, 4).fill = RED if v is None else YELLOW
            ws.cell(r, 9).fill = YELLOW
            dv_status.add(ws.cell(r, 9))
            r += 1
        if pref == "TT":
            put(ws, r + 1, 1, "Tham số bắt buộc trước khi đối chiếu: TT-01, TT-02, TT-03, TT-04, TT-10, TT-20 hoặc TT-21. "
                              "Ô đỏ = chưa tìm thấy trong hợp đồng, cần nhập tay và ghi điều khoản.", Font(name="Arial", size=9, italic=True))

    # ---- Hồ sơ TT theo HĐ
    ws = wb.create_sheet("Ho so TT theo HD")
    header(ws, 1, ["Loại", "Tài liệu hợp đồng yêu cầu", "Điều khoản", "Đã nộp đợt này?", "Ghi chú"], [18, 70, 30, 16, 40])
    r = 2
    for row in rows:
        if row["ma"] in ("TT-60", "TT-61"):
            for it in row.get("items", []):
                put(ws, r, 1, "Mỗi đợt" if row["ma"] == "TT-60" else "Quyết toán")
                put(ws, r, 2, it, wrap=True)
                put(ws, r, 3, row["dieu"])
                ws.cell(r, 4).fill = YELLOW
                r += 1

    # ---- BOQ
    ws = wb.create_sheet("BOQ")
    header(ws, 1, ["Mã HM", "STT", "Nhóm / hạng mục", "Tên công việc", "ĐVT", "KL hợp đồng", "Đơn giá (chưa VAT)",
                   "Thành tiền", "Nguồn", "Ghi chú", "Khoá so khớp (không sửa)"], [9, 7, 36, 60, 9, 13, 15, 17, 26, 30, 30])
    r = 2
    if boq:
        n = 0
        for it in boq["items"]:
            if not it["dvt"]:
                continue
            n += 1
            put(ws, r, 1, f"HM{n:04d}")
            put(ws, r, 2, it["stt"])
            put(ws, r, 3, it["nhom"], wrap=False)
            put(ws, r, 4, it["ten"])
            put(ws, r, 5, it["dvt"])
            put(ws, r, 6, it["kl_hd"], BLUE, QF)
            put(ws, r, 7, it["don_gia"], BLUE, NUMF)
            put(ws, r, 8, f"=ROUND(F{r}*G{r},0)", fmt=NUMF)
            put(ws, r, 9, f"{boq_src} ({boq['sheet']}!dòng {it['dong']})")
            put(ws, r, 11, it["khoa"], Font(name="Arial", size=8, color="808080"), text=True)
            r += 1
    last = max(r - 1, 2)
    put(ws, r + 1, 4, "Cộng (chưa VAT)", BOLD)
    put(ws, r + 1, 8, f"=SUM(H2:H{last})", BOLD, NUMF)
    put(ws, r + 2, 4, "Thuế GTGT (theo Dieu khoan TT!TT-04)", BOLD)
    put(ws, r + 2, 8, f"=ROUND(H{r + 1}*INDEX('Dieu khoan TT'!D:D,MATCH(\"TT-04\",'Dieu khoan TT'!A:A,0)),0)", BOLD, NUMF)
    put(ws, r + 3, 4, "Cộng sau thuế", BOLD)
    put(ws, r + 3, 8, f"=H{r + 1}+H{r + 2}", BOLD, NUMF)
    put(ws, r + 4, 4, "Giá trị HĐ theo văn bản hợp đồng (TT-02)", BOLD)
    put(ws, r + 4, 8, "=INDEX('Dieu khoan TT'!D:D,MATCH(\"TT-02\",'Dieu khoan TT'!A:A,0))", BOLD, NUMF)
    put(ws, r + 5, 4, "Chênh lệch BOQ so với hợp đồng (phải = 0 hoặc giải thích được do làm tròn)", BOLD)
    put(ws, r + 5, 8, f"=H{r + 3}-H{r + 4}", BOLD, NUMF, fill=YELLOW)

    # ---- Phụ lục
    ws = wb.create_sheet("Phu luc")
    header(ws, 1, ["Số phụ lục", "Ngày ký", "Nội dung thay đổi", "Giá trị tăng/giảm (sau VAT)", "Mã HM bị ảnh hưởng",
                   "File", "Đã cập nhật vào BOQ?"], [14, 12, 60, 20, 22, 30, 16])
    for c in range(1, 8):
        ws.cell(2, c).fill = YELLOW

    # ---- KL kỳ trước
    ws = wb.create_sheet("KL ky truoc")
    header(ws, 1, ["Mã HM", "Tên công việc", "ĐVT", "KL lũy kế đến hết kỳ trước", "Đơn giá (BOQ)", "GT lũy kế kỳ trước",
                   "Đến hết đợt", "Nguồn"], [9, 60, 9, 16, 15, 18, 10, 50])
    r = 2
    if boq:
        n = 0
        for it in boq["items"]:
            if not it["dvt"]:
                continue
            n += 1
            put(ws, r, 1, f"HM{n:04d}")
            put(ws, r, 2, it["ten"])
            put(ws, r, 3, it["dvt"])
            put(ws, r, 4, None, BLUE, QF, fill=YELLOW)
            put(ws, r, 5, f"=INDEX(BOQ!G:G,MATCH(A{r},BOQ!A:A,0))", fmt=NUMF)
            put(ws, r, 6, f"=ROUND(D{r}*E{r},0)", fmt=NUMF)
            put(ws, r, 8, "Đợt đầu: nhập tay (hoặc 0 nếu chưa thi công)")
            r += 1
    put(ws, r + 1, 2, "Cộng GT lũy kế kỳ trước (chưa VAT)", BOLD)
    put(ws, r + 1, 6, f"=SUM(F2:F{max(r - 1, 2)})", BOLD, NUMF)

    # ---- Lịch sử TT
    ws = wb.create_sheet("Lich su TT")
    header(ws, 1, ["Đợt", "Ngày duyệt", "GT thực hiện kỳ (sau VAT)", "GT lũy kế thực hiện (sau VAT)", "Giữ lại kỳ",
                   "Thu hồi tạm ứng kỳ", "GT thanh toán kỳ", "Đã chi thực tế", "Số UNC / chứng từ", "Nguồn"],
           [6, 12, 20, 20, 16, 16, 18, 18, 20, 50])
    for c in range(1, 11):
        ws.cell(2, c).fill = YELLOW

    # ---- Nhật ký
    ws = wb.create_sheet("Nhat ky")
    header(ws, 1, ["Thời điểm", "Thao tác", "File nguồn", "Ghi chú"], [18, 30, 70, 50])
    put(ws, 2, 1, f"{dt.datetime.now():%d/%m/%Y %H:%M}")
    put(ws, 2, 2, "Tạo master")
    put(ws, 2, 3, str(contract_file))
    put(ws, 2, 4, f"BOQ: {boq_src or 'chưa nạp'}")
    wb.save(out)


# ============================================================== đọc master
def doc_master(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    p = {}
    for sh in ("Thong tin HD", "Dieu khoan TT"):
        for row in wb[sh].iter_rows(min_row=2, values_only=True):
            if row[0] and re.fullmatch(r"(TC|TT)-\d+", str(row[0])):
                p[row[0]] = dict(ten=row[2], gia_tri=row[3], trang_thai=row[8])
    boq = {}
    for row in wb["BOQ"].iter_rows(min_row=2, values_only=True):
        if row[0] and str(row[0]).startswith("HM"):
            it = dict(ma=row[0], stt=row[1], nhom=row[2] or "", ten=row[3] or "", dvt=row[4] or "",
                      kl_hd=num(row[5]), don_gia=num(row[6]))
            boq[row[10] if len(row) > 10 and row[10] else khoa(it)] = it
    kt = {}
    for row in wb["KL ky truoc"].iter_rows(min_row=2, values_only=True):
        if row[0] and str(row[0]).startswith("HM"):
            kt[row[0]] = dict(kl=num(row[3]), dot=row[6], nguon=row[7])
    ls = [r for r in wb["Lich su TT"].iter_rows(min_row=2, values_only=True) if r[0] not in (None, "")]
    return p, boq, kt, ls


# ============================================================== nạp đợt đã duyệt vào master
def nap_dot(master_path, file_dot, so_dot=None):
    wv = openpyxl.load_workbook(file_dot, data_only=True)
    bang = tim_bang_chi_tiet(wv)
    tong = tim_so_tong_hop(wv)
    wb = openpyxl.load_workbook(master_path)
    boq_ws, kt_ws, ls_ws, log = wb["BOQ"], wb["KL ky truoc"], wb["Lich su TT"], wb["Nhat ky"]
    ma_by_key = {}
    for r in range(2, boq_ws.max_row + 1):
        ma = boq_ws.cell(r, 1).value
        if ma and str(ma).startswith("HM"):
            it = dict(nhom=boq_ws.cell(r, 3).value or "", ten=boq_ws.cell(r, 4).value or "", dvt=boq_ws.cell(r, 5).value or "")
            ma_by_key[boq_ws.cell(r, 11).value or khoa(it)] = ma
    row_by_ma = {kt_ws.cell(r, 1).value: r for r in range(2, kt_ws.max_row + 1) if kt_ws.cell(r, 1).value}
    so_dot = so_dot or _doan_so_dot(file_dot, wv)
    ok, miss = 0, []
    for it in bang["items"]:
        if not it["dvt"]:
            continue
        ma = ma_by_key.get(it["khoa"])
        if not ma or ma not in row_by_ma:
            miss.append(it["ten"])
            continue
        r = row_by_ma[ma]
        put(kt_ws, r, 4, it["kl_lk"] or 0, BLUE, QF)
        put(kt_ws, r, 7, so_dot)
        put(kt_ws, r, 8, f"{Path(file_dot).name} ({bang['sheet']}!dòng {it['dong']}, cột lũy kế đến hết kỳ này)")
        ok += 1
    # lịch sử thanh toán
    r = 2
    while ls_ws.cell(r, 1).value not in (None, ""):
        r += 1
    vals = [so_dot, None, tong.get("gt_kn", {}).get("gia_tri"), tong.get("gt_lk", {}).get("gia_tri"),
            tong.get("giu_lai", {}).get("gia_tri"), tong.get("thu_hoi_tu", {}).get("gia_tri"),
            tong.get("tt_ky_nay", {}).get("gia_tri"), None, None,
            f"{Path(file_dot).name}: " + "; ".join(f"{k}={v['o']}" for k, v in tong.items())]
    for j, v in enumerate(vals, 1):
        put(ls_ws, r, j, v, BLUE if isinstance(v, float) else BASE, NUMF if isinstance(v, float) else None,
            fill=YELLOW if j in (2, 8, 9) else None)
    lr = 2
    while log.cell(lr, 1).value:
        lr += 1
    put(log, lr, 1, f"{dt.datetime.now():%d/%m/%Y %H:%M}")
    put(log, lr, 2, f"Nạp đợt {so_dot}")
    put(log, lr, 3, str(file_dot))
    put(log, lr, 4, f"{ok} hạng mục khớp; {len(miss)} không khớp BOQ" + (f": {'; '.join(miss[:5])}…" if miss else ""))
    wb.save(master_path)
    return ok, miss, tong


def _doan_so_dot(path, wv):
    m = re.search(r"đợt\s*0*(\d+)", nfc(Path(path).name), re.I)
    return int(m.group(1)) if m else None


# ============================================================== đối chiếu hồ sơ mới với master
def doi_chieu(master_path, file_moi, out):
    p, boq, kt, ls = doc_master(master_path)
    wv = openpyxl.load_workbook(file_moi, data_only=True)
    bang = tim_bang_chi_tiet(wv)
    tong = tim_so_tong_hop(wv)
    vat = p.get("TT-04", {}).get("gia_tri") or 0
    vat = vat if isinstance(vat, float) else 0
    giu = p.get("TT-21", {}).get("gia_tri")
    rows, seen = [], set()
    for it in bang["items"]:
        if not it["dvt"]:
            continue
        m = boq.get(it["khoa"])
        seen.add(it["khoa"])
        k = kt.get(m["ma"]) if m else None
        rec = dict(dong=it["dong"], ten=it["ten"], dvt=it["dvt"], ma=m["ma"] if m else "",
                   kl_hd_file=it["kl_hd"], kl_hd_master=m["kl_hd"] if m else None,
                   dg_file=it["don_gia"], dg_master=m["don_gia"] if m else None,
                   kt_file=it["kl_kt"], kt_master=k["kl"] if k else None,
                   kn=it["kl_kn"], lk=it["kl_lk"], loi=[])
        if not m:
            rec["loi"].append("Không có trong BOQ master (phát sinh chưa có phụ lục?)")
        else:
            if (rec["kl_hd_file"] or 0) != (rec["kl_hd_master"] or 0) and abs((rec["kl_hd_file"] or 0) - (rec["kl_hd_master"] or 0)) > 1e-6:
                rec["loi"].append("KL HĐ khác master")
            if abs((rec["dg_file"] or 0) - (rec["dg_master"] or 0)) > 0.5:
                rec["loi"].append("Đơn giá khác master")
            if rec["kt_master"] is not None and abs((rec["kt_file"] or 0) - rec["kt_master"]) > 1e-6:
                rec["loi"].append("KL kỳ trước khác master (đợt trước đã duyệt)")
            if rec["kt_master"] is None:
                rec["loi"].append("Master chưa có KL kỳ trước (đợt đầu: nhập tay)")
            if (rec["lk"] or 0) > (rec["kl_hd_master"] or 0) + 1e-6:
                rec["loi"].append("Lũy kế vượt KL HĐ")
        rows.append(rec)
    thieu = [v for kx, v in boq.items() if kx not in seen]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tong hop"
    put(ws, 1, 1, "ĐỐI CHIẾU HỒ SƠ THANH TOÁN VỚI MASTER HỢP ĐỒNG", Font(name="Arial", bold=True, size=12))
    put(ws, 2, 1, f"Hồ sơ: {Path(file_moi).name}   |   Master: {Path(master_path).name}   |   Bảng chi tiết đọc được: {bang['sheet']}")
    header(ws, 4, ["Kiểm tra", "Theo hồ sơ", "Theo master", "Chênh lệch", "Kết luận", "Ô trong hồ sơ"], [52, 20, 20, 18, 36, 34])
    ws.freeze_panes = "A5"

    def line(r, name, v_file, v_master, cell, pct=False, tol=1):
        put(ws, r, 1, name)
        put(ws, r, 2, v_file, fmt="0.00%" if pct else NUMF)
        put(ws, r, 3, v_master, fmt="0.00%" if pct else NUMF)
        if isinstance(v_file, (int, float)) and isinstance(v_master, (int, float)):
            put(ws, r, 4, f"=B{r}-C{r}", fmt="0.00%" if pct else NUMF)
            put(ws, r, 5, f'=IF(ABS(D{r})<={tol},"Khớp","Lệch - kiểm tra")')
        else:
            put(ws, r, 5, "Thiếu số liệu để so", fill=YELLOW)
        put(ws, r, 6, cell)

    r = 5
    line(r, "Giá trị hợp đồng (sau VAT)", tong.get("gt_hd", {}).get("gia_tri"), p.get("TT-02", {}).get("gia_tri"),
         tong.get("gt_hd", {}).get("o", "")); r += 1
    gt_kn, gl = tong.get("gt_kn", {}).get("gia_tri"), tong.get("giu_lai", {}).get("gia_tri")
    line(r, "Tỷ lệ giữ lại = Giữ lại / GT thực hiện kỳ này", (gl / gt_kn) if gl and gt_kn else None,
         giu if isinstance(giu, float) else None, tong.get("giu_lai", {}).get("o", ""), pct=True, tol=0.0005); r += 1
    last_lk = next((num(x[3]) for x in reversed(ls) if num(x[3])), None)
    line(r, "GT lũy kế đến hết kỳ trước = GT lũy kế đợt trước đã duyệt", tong.get("gt_lk_kt", {}).get("gia_tri"),
         last_lk, tong.get("gt_lk_kt", {}).get("o", "")); r += 1
    kt_sum = sum((x["kt_file"] or 0) * (x["dg_master"] or 0) for x in rows) * (1 + vat)
    line(r, "Σ(KL kỳ trước trong hồ sơ × đơn giá master) × (1+VAT)", round(kt_sum), last_lk, bang["sheet"]); r += 1
    ln = sum((x["kn"] or 0) * (x["dg_master"] or 0) for x in rows) * (1 + vat)
    line(r, "Σ(KL kỳ này × đơn giá master) × (1+VAT) so với GT kỳ này trong hồ sơ", round(ln), gt_kn,
         tong.get("gt_kn", {}).get("o", ""), tol=max(1000, (gt_kn or 0) * 1e-5)); r += 1
    r += 1
    from collections import Counter
    cnt = Counter(l for x in rows for l in x["loi"])
    put(ws, r, 1, "Số hạng mục theo loại lệch", BOLD); r += 1
    for k_, n_ in cnt.most_common():
        put(ws, r, 1, k_); put(ws, r, 2, n_); r += 1
    put(ws, r, 1, "Hạng mục có trong master nhưng không có trong hồ sơ"); put(ws, r, 2, len(thieu)); r += 1
    put(ws, r + 1, 1, "Tham số master chưa xác nhận: " + ", ".join(k for k, v in p.items() if v["trang_thai"] != "Đã xác nhận"),
        Font(name="Arial", size=9, italic=True, color="9A5B00"))

    ws2 = wb.create_sheet("Hang muc")
    header(ws2, 1, ["Dòng hồ sơ", "Mã HM", "Tên công việc", "ĐVT", "KL HĐ hồ sơ", "KL HĐ master", "Đơn giá hồ sơ",
                    "Đơn giá master", "KL kỳ trước hồ sơ", "KL kỳ trước master", "KL kỳ này", "KL lũy kế",
                    "Ảnh hưởng đơn giá (× KL lũy kế)", "Ảnh hưởng kỳ trước (× đơn giá)", "Lệch"],
           [8, 9, 50, 8, 12, 12, 13, 13, 12, 12, 12, 12, 16, 16, 50])
    for i, x in enumerate(rows, 2):
        vals = [x["dong"], x["ma"], x["ten"], x["dvt"], x["kl_hd_file"], x["kl_hd_master"], x["dg_file"], x["dg_master"],
                x["kt_file"], x["kt_master"], x["kn"], x["lk"]]
        for j, v in enumerate(vals, 1):
            put(ws2, i, j, v, fmt=QF if j in (5, 6, 9, 10, 11, 12) else (NUMF if j in (7, 8) else None))
        put(ws2, i, 13, f'=IF(H{i}="","",(G{i}-H{i})*L{i})', fmt=NUMF)
        put(ws2, i, 14, f'=IF(J{i}="","",(I{i}-J{i})*H{i})', fmt=NUMF)
        put(ws2, i, 15, "; ".join(x["loi"]))
        if x["loi"]:
            ws2.cell(i, 15).fill = RED if any(s.startswith(("Đơn giá", "KL HĐ", "Lũy kế", "KL kỳ trước khác", "Không có"))
                                              for s in x["loi"]) else YELLOW
    ws3 = wb.create_sheet("Thieu trong ho so")
    header(ws3, 1, ["Mã HM", "Tên công việc", "ĐVT", "KL HĐ"], [9, 60, 9, 12])
    for i, v in enumerate(thieu, 2):
        for j, x in enumerate([v["ma"], v["ten"], v["dvt"], v["kl_hd"]], 1):
            put(ws3, i, j, x)
    wb.save(out)
    return rows, tong, cnt


# ============================================================== CLI
def main():
    ap = argparse.ArgumentParser(description="Master file hợp đồng")
    sub = ap.add_subparsers(dest="lenh", required=True)
    a = sub.add_parser("tao"); a.add_argument("hop_dong"); a.add_argument("--boq"); a.add_argument("--out")
    a.add_argument("--ai", action="store_true"); a.add_argument("--model")
    a.add_argument("--host"); a.add_argument("--api", choices=["ollama", "openai"])
    d = sub.add_parser("ai-trich", help="Trích toàn bộ tham số bằng AI và so với quy tắc")
    d.add_argument("hop_dong"); d.add_argument("--out"); d.add_argument("--model")
    d.add_argument("--host"); d.add_argument("--api", choices=["ollama", "openai"])
    d.add_argument("--an-danh", action="store_true", help="Che tên các bên, tên người, MST, số tài khoản trước khi gửi")
    d.add_argument("--ca-hop-dong", action="store_true", help="Gửi cả hợp đồng thay vì chỉ các Điều liên quan thanh toán")
    d.add_argument("--num-ctx", type=int, default=None, help="Ngữ cảnh model; mặc định 32768, hoặc 65536 khi --ca-hop-dong")
    b = sub.add_parser("nap-dot"); b.add_argument("master"); b.add_argument("file_dot"); b.add_argument("--dot", type=int)
    c = sub.add_parser("doi-chieu"); c.add_argument("master"); c.add_argument("file_moi"); c.add_argument("--out")
    e = sub.add_parser("cham", help="Chấm kết quả trích xuất với đáp án (.tsv)")
    e.add_argument("dap_an"); e.add_argument("ket_qua", help="Hợp đồng (.doc/.docx/.pdf -> chấm quy tắc) hoặc file .tsv kết quả AI")
    e.add_argument("--out")
    args = ap.parse_args()
    if args.lenh in ("tao", "ai-trich"):
        cfg = cau_hinh_ai()
        args.model = args.model or cfg["model"]
        args.host = args.host or cfg["may_ai"]
        args.api = args.api or cfg["api"]
    if args.lenh == "tao":
        paras = doc_hop_dong(args.hop_dong)
        rows = trich_xuat(paras, LocalLLM(args.model, args.host, args.api) if args.ai else None)
        boq = None
        if args.boq:
            boq = tim_bang_chi_tiet(openpyxl.load_workbook(args.boq, data_only=True))
        out = args.out or f"master_{re.sub(r'[^0-9A-Za-z]+', '_', str(rows[0]['gia_tri'] or 'hd'))[:30]}.xlsx"
        ghi_master(out, args.hop_dong, rows, boq, Path(args.boq).name if args.boq else "")
        found = sum(1 for r in rows if r["gia_tri"] is not None)
        print(f"Đã tạo {out}: trích được {found}/{len(rows)} tham số; BOQ {sum(1 for i in (boq or {'items': []})['items'] if i['dvt'])} hạng mục")
        for r in rows:
            print(f"  {r['ma']:6} {r['ten'][:38]:38} = {str(r['gia_tri'])[:60]}")
    elif args.lenh == "nap-dot":
        ok, miss, tong = nap_dot(args.master, args.file_dot, args.dot)
        print(f"Nạp {ok} hạng mục; {len(miss)} không khớp BOQ. Số tổng hợp: " +
              ", ".join(f"{k}={v['gia_tri']:,.0f}" for k, v in tong.items()))
    elif args.lenh == "cham":
        gold = doc_dap_an(args.dap_an)
        src = Path(args.ket_qua)
        if src.suffix.lower() in (".doc", ".docx", ".pdf"):
            rows = trich_xuat(doc_hop_dong(src))
            kq = {r["ma"]: dict(gia_tri=r["gia_tri"], trich=r["trich"], ghi_chu=r["ghi_chu"]) for r in rows}
            nguon = "Quy tắc (master_hd.py)"
        else:
            kq = {m: dict(gia_tri=g["gia_tri"], trich=g["trich"], ghi_chu=g["ghi_chu"]) for m, g in doc_dap_an(src).items()}
            nguon = f"AI ({src.name})"
        res = cham_diem(kq, gold)
        from collections import Counter
        c = Counter(x[1] for x in res)
        diem = (c["Đúng"] + 0.5 * c["Một phần"]) / len(res) if res else 0
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Cham diem"
        put(ws, 1, 1, f"CHẤM ĐIỂM: {nguon} so với đáp án {Path(args.dap_an).name}", Font(name="Arial", bold=True, size=12), text=True)
        put(ws, 2, 1, f"Đúng {c['Đúng']} | Một phần {c['Một phần']} | Sai {c['Sai']} | Điểm {diem:.1%} "
                      "(chấm tự động - cần người soát các dòng Một phần/Sai)", text=True)
        header(ws, 4, ["Mã", "Đáp án", "Kết quả", "Chấm", "Lý do", "Ghi chú của kết quả", "Ghi chú / bẫy trong đáp án"],
               [8, 45, 45, 11, 32, 50, 50])
        for i, (ma, ch, ly) in enumerate(res, 5):
            k = kq.get(ma, {})
            vals = [ma, gold[ma]["gia_tri"], k.get("gia_tri"), ch, ly, k.get("ghi_chu", ""), gold[ma]["ghi_chu"]]
            for j, v in enumerate(vals, 1):
                put(ws, i, j, v, wrap=j in (2, 3, 6, 7), text=True)
            ws.cell(i, 4).fill = GREEN if ch == "Đúng" else (YELLOW if ch == "Một phần" else RED)
        out = args.out or f"cham_{src.stem[:40]}.xlsx"
        wb.save(out)
        print(f"{nguon}: Đúng {c['Đúng']}, Một phần {c['Một phần']}, Sai {c['Sai']} -> điểm {diem:.1%}. Đã ghi {out}")
    elif args.lenh == "ai-trich":
        paras = doc_hop_dong(args.hop_dong)
        rows = trich_xuat(paras)
        text = "\n".join(t for _, t in (paras if args.ca_hop_dong else rut_gon(paras)))
        if args.an_danh:
            text = an_danh(text, rows)
        llm = LocalLLM(args.model, args.host, args.api,
                       num_ctx=args.num_ctx or (65536 if args.ca_hop_dong else 32768))
        if not llm.ok():
            sys.exit(f"Không kết nối được máy AI tại {args.host}. Kiểm tra máy AI đang bật và Ollama đang chạy; "
                     "mở 5_trich_hop_dong_AI.bat, tab Máy AI để kiểm tra kết nối.")
        print(f"Gửi {len(text):,} ký tự tới {args.model} …")
        params = [(r["ma"], r["nhom"], r["ten"], r["don_vi"]) for r in rows]
        try:
            ai, raw = llm.extract_all(text, params)
        except json.JSONDecodeError:
            sys.exit("AI trả về không phải JSON hợp lệ. Thử lại hoặc đổi model.")
        goc = "\n".join(t for _, t in paras)
        cmp = so_sanh_ai(rows, ai, goc, an_danh(goc, rows) if args.an_danh else "")
        out = args.out or f"so_sanh_AI_{Path(args.hop_dong).stem[:40]}.xlsx"
        ghi_so_sanh(out, args.hop_dong, cmp, llm.stats, args.model, len(text), args.an_danh)
        Path(out).with_suffix(".json").write_text(raw, encoding="utf-8")
        from collections import Counter
        print(f"Đã ghi {out} ({llm.stats}). So sánh: {dict(Counter(x['khop'] for x in cmp))}; "
              f"trích dẫn: {dict(Counter(x['that'] for x in cmp))}")
    else:
        out = args.out or "doi_chieu_" + Path(args.file_moi).stem[:40] + ".xlsx"
        rows, tong, cnt = doi_chieu(args.master, args.file_moi, out)
        print(f"Đã ghi {out}: {len(rows)} hạng mục; lệch: {dict(cnt)}")


if __name__ == "__main__":
    main()
