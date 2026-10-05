# -*- coding: utf-8 -*-
"""
trich_lo_hop_dong.py - Trích 36 tham số cho MỘT LOẠT hợp đồng, nhờ máy Local AI trong mạng LAN.

Luồng dữ liệu
    Máy này (lưu hợp đồng)                      Máy AI trong LAN (Ollama)
    1. đọc hợp đồng, trích bằng quy tắc
    2. cắt các Điều liên quan, ẩn danh   ───►   3. đọc đoạn văn, trả JSON từng nhóm tham số (8 nhóm)
    4. kiểm trích dẫn AI có thật trong HĐ  ◄───     (không lưu gì trên máy AI)
    5. gộp quy tắc + AI, ghi file Excel trên máy này

Lệnh
    python trich_lo_hop_dong.py kiem-tra                 kiểm tra kết nối máy AI, liệt kê model
    python trich_lo_hop_dong.py cau-hinh --may-ai 192.168.1.50 [--model qwen3:8b]
    (Dùng hằng ngày: cửa sổ giao_dien_trich_hd.py, mở bằng 5_trich_hop_dong_AI.bat)
    python trich_lo_hop_dong.py chay <thư mục hoặc file ...> [--out ket_qua_AI\\trich_lo] [--lam-lai] [--khong-ai]

Kết quả (mặc định ket_qua_AI\\trich_lo\\)
    master_<tên HĐ>.xlsx       master từng hợp đồng: giá trị đã gộp, trạng thái "Chưa xác nhận", nạp vào CSDL bằng 0_...bat
    so_sanh_AI_<tên HĐ>.xlsx   quy tắc vs AI từng tham số, trích dẫn AI có thật trong HĐ không
    TONG_HOP_<ngày giờ>.xlsx   mỗi hợp đồng một dòng × 36 tham số; danh sách cần xác nhận; nhật ký
    _ai_json\\                 câu trả lời gốc của AI (chạy lại không hỏi lại AI, trừ khi --lam-lai)

Không ghi đè, không sửa file hợp đồng. Không gửi file, chỉ gửi đoạn văn đã cắt (và đã ẩn danh nếu bật).
"""
import argparse
import datetime as dt
import hashlib
import json
import re
import sys
import time
import urllib.error
import urllib.request
from collections import Counter
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import master_hd as M

DUOI = {".doc", ".docx", ".pdf"}
GREY = PatternFill("solid", fgColor="EEEEEE")
HEAD = PatternFill("solid", fgColor="1F3864")
WHITE_B = Font(name="Arial", size=10, bold=True, color="FFFFFF")

# Tình trạng sau khi gộp quy tắc + AI -> (màu, cần người xem?)
TT_MAU = {
    "Khớp": (M.GREEN, False),
    "Chỉ quy tắc thấy": (M.GREEN, False),
    "Khác cách ghi": (M.YELLOW, True),
    "Chỉ AI thấy": (M.YELLOW, True),
    "Lệch": (M.RED, True),
    "AI bịa trích dẫn": (M.RED, True),
    "Không tìm thấy": (GREY, True),
    "Chưa hỏi AI": (M.YELLOW, False),
}
BAT_BUOC = {"TT-01", "TT-02", "TT-03", "TT-04", "TT-10", "TT-20", "TT-21"}


# ------------------------------------------------------------------ cấu hình
def ghi_cau_hinh(may_ai=None, model=None, an_danh=None):
    cfg = M.cau_hinh_ai()
    if may_ai:
        h = may_ai.strip().rstrip("/")
        if not h.startswith("http"):
            h = "http://" + h
        if not re.search(r":\d+$", h.split("//", 1)[1]):
            h += ":11434"
        cfg["may_ai"] = h
    if model:
        cfg["model"] = model.strip()
    if an_danh is not None:
        cfg["an_danh"] = an_danh
    M.CAU_HINH_AI.write_text(json.dumps(cfg, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Đã lưu {M.CAU_HINH_AI.name}: máy AI = {cfg['may_ai']}, model = {cfg['model']}, ẩn danh = {cfg['an_danh']}")
    return cfg


def kiem_tra(cfg=None, thu_hoi=True, bao=print):
    """Kiểm tra máy AI: kết nối, có model chưa, trả lời thử. Trả về True/False."""
    cfg = cfg or M.cau_hinh_ai()
    host, model = cfg["may_ai"], cfg["model"]
    bao(f"Máy AI: {host}   Model: {model}")
    try:
        with urllib.request.urlopen(host + "/api/tags", timeout=5) as r:
            tags = json.loads(r.read())
    except Exception as e:
        bao(f"  [LỖI] Không kết nối được ({e.__class__.__name__}: {e}).")
        bao("  Kiểm tra: máy AI đã bật? Ollama đang chạy? Đã đặt OLLAMA_HOST=0.0.0.0:11434 và mở tường lửa cổng 11434 "
              "cho máy này? IP máy AI có đổi không?")
        return False
    co = [m.get("name") for m in tags.get("models", [])]
    bao(f"  [OK] Kết nối được. Model trên máy AI: {', '.join(co) or '(chưa có model nào)'}")
    if model not in co and f"{model}:latest" not in co:
        bao(f"  [LỖI] Máy AI chưa có model '{model}'. Chọn một model trong danh sách trên "
              f"(tab Máy AI) hoặc chạy trên máy AI: ollama pull {model}")
        return False
    if thu_hoi:
        llm = M.LocalLLM(model, host, cfg["api"], num_ctx=2048, timeout=300)
        try:
            t0 = time.time()
            raw = llm.chat("Bạn chỉ trả lời JSON.", 'Trả về đúng JSON {"tra_loi": "OK"}')
            bao(f"  [OK] Model trả lời sau {time.time() - t0:.0f} giây: {raw.strip()[:60]}")
        except Exception as e:
            bao(f"  [LỖI] Model không trả lời được: {e}")
            return False
    return True


def tu_tim_lai(cfg, bao=print):
    """Máy AI không trả lời ở địa chỉ cũ (đổi mạng / đổi IP): dò trong mạng LAN.
    Thấy đúng một máy có model đang dùng thì tự cập nhật cấu hình và chạy tiếp."""
    import may_ai as A
    bao("\nĐịa chỉ cũ không trả lời. Đang dò máy AI trong mạng LAN …")
    try:
        thay = A.tim_may_ai()
    except Exception as e:
        bao(f"  Không dò được ({e}).")
        return None
    hop = [(ip, ms) for ip, ms in thay if cfg["model"] in ms or f"{cfg['model']}:latest" in ms]
    for ip, ms in thay:
        bao(f"  Thấy máy AI {ip}: {', '.join(ms)}")
    if len(hop) != 1:
        bao("  " + ("Không thấy máy AI nào." if not thay else
                      f"Có {len(hop)} máy có model {cfg['model']} – chọn trong tab Máy AI."))
        return None
    moi = dict(cfg, may_ai=A.chuan_host(hop[0][0], A.tach_host(cfg["may_ai"])[1]))
    A.luu_cau_hinh(moi)
    bao(f"  Máy AI đã đổi địa chỉ: {cfg['may_ai']} → {moi['may_ai']}. Đã cập nhật cấu hình, chạy tiếp.\n")
    return moi if kiem_tra(moi, thu_hoi=False, bao=bao) else None


# ------------------------------------------------------------------ gộp quy tắc + AI
def gia_tri_ai(v, don_vi):
    """Đổi giá trị AI (chuỗi) về kiểu như quy tắc: 10% -> 0.1, 1.234.567 -> 1234567."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return v
    t = M.nfc(str(v)).strip()
    if not t or t.upper() == "KHÔNG TÌM THẤY":
        return None
    if re.fullmatch(r"-?\d+([.,]\d+)?\s*%", t):
        return round(M.so(t.rstrip("% ").strip()) / 100, 6)
    if don_vi == "đồng" and re.fullmatch(r"\d{1,3}([.,]\d{3})+|\d+", t):
        return int(re.sub(r"[.,]", "", t))
    if don_vi in ("ngày", "tháng", "bộ"):
        m = re.fullmatch(r"(\d+)(\s*(ngày|tháng|bộ))?", t)
        if m:
            return int(m.group(1))
    return t


def gop(rows, cmp):
    """Gộp kết quả quy tắc và AI. Quy tắc có giá trị thì giữ (AI để đối chứng); quy tắc không thấy thì lấy AI
    nếu trích dẫn AI tìm lại được trong hợp đồng. Trả về rows đã cập nhật + danh sách tình trạng từng mã."""
    by = {x["ma"]: x for x in cmp}
    tinh_trang = {}
    for r in rows:
        x = by.get(r["ma"])
        if x is None:
            tinh_trang[r["ma"]] = "Chưa hỏi AI" if r["gia_tri"] is not None else "Không tìm thấy"
            continue
        that = x["that"] or ""
        bia = that.startswith("KHÔNG")
        k = x["khop"]
        if k == "Khác cách ghi" and r["gia_tri"] is not None and re.search(r"\[(BEN_A|BEN_B|CA_NHAN|SO)\]", str(x["ai"])):
            k = "Khớp"  # AI chỉ trả lại chuỗi đã ẩn danh, quy tắc có tên thật
        if k == "Chỉ AI thấy":
            if bia:
                tt = "AI bịa trích dẫn"
                r["ghi_chu"] = _noi(r["ghi_chu"], f"AI ghi '{x['ai']}' nhưng trích dẫn không có trong HĐ – không dùng")
            else:
                tt = "Chỉ AI thấy"
                r.update(gia_tri=gia_tri_ai(x["ai"], r["don_vi"]), cach="Local AI (quy tắc không thấy)",
                         dieu=x["dieu_ai"] or "", trich=str(x["trich_ai"])[:400])
                if r["ma"] in ("TT-60", "TT-61") and isinstance(r["gia_tri"], str):
                    r["items"] = [y.strip(" +-;.") for y in re.split(r";|\n", r["gia_tri"]) if y.strip(" +-;.")]
                r["ghi_chu"] = _noi(r["ghi_chu"], "Giá trị do AI trích – đọc trích dẫn trước khi xác nhận"
                                    + (f"; AI ghi chú: {x['ghi_chu_ai']}" if x["ghi_chu_ai"] else ""))
        elif k == "Lệch":
            tt = "Lệch"
            r["ghi_chu"] = _noi(r["ghi_chu"], f"LỆCH: AI đọc là '{x['ai']}'"
                                + (f" ({x['dieu_ai']})" if x["dieu_ai"] else "") + " – đọc lại điều khoản")
        elif k == "Khác cách ghi":
            tt = "Khác cách ghi"
            r["ghi_chu"] = _noi(r["ghi_chu"], f"AI ghi: '{x['ai']}'")
        elif k == "Cả hai không thấy":
            tt = "Không tìm thấy"
        else:
            tt = k  # Khớp / Chỉ quy tắc thấy
            if k == "Khớp":
                r["cach"] = r["cach"] + " + AI khớp"
                if bia:
                    r["ghi_chu"] = _noi(r["ghi_chu"], "AI cùng giá trị nhưng trích dẫn không đúng nguyên văn")
        if x["ghi_chu_ai"] and k != "Chỉ AI thấy" and re.search(r"mâu thuẫn|khác nhau|sai|bất thường", x["ghi_chu_ai"], re.I):
            r["ghi_chu"] = _noi(r["ghi_chu"], f"AI lưu ý: {x['ghi_chu_ai']}")
        tinh_trang[r["ma"]] = tt
    return rows, tinh_trang


def _noi(a, b):
    return f"{a}; {b}" if a else b


def _ten_an_toan(s, n=60):
    return re.sub(r'[\\/:*?"<>|]+', "_", s)[:n].strip(" ._")


def bam(path, model, an, ca):
    h = hashlib.sha1(Path(path).read_bytes()).hexdigest()[:16]
    return f"{h}_{_ten_an_toan(model, 30)}_{'ad' if an else 'goc'}_{'ca' if ca else 'rg'}"


# ------------------------------------------------------------------ xử lý một hợp đồng
class Dung(Exception):
    """Người dùng bấm Dừng."""


class LoiChay(Exception):
    """Lỗi làm dừng cả lô (không kết nối được máy AI, không có file...)."""


def _uoc_token(text):
    return int(len(text) / 3)  # tiếng Việt với tokenizer Qwen/Gemma: khoảng 3–3,5 ký tự / token


def hoi_ai_theo_nhom(llm, text, rows, bao=print, dung=None, so_lan=2):
    """Hỏi AI từng nhóm tham số (8 nhóm) thay vì 36 tham số một lần: mỗi lần trả lời ngắn, nhóm lỗi thì hỏi lại
    riêng nhóm đó. Hợp đồng đặt ở ĐẦU câu hỏi, giống hệt nhau giữa các nhóm, để Ollama dùng lại phần đã đọc
    (chỉ đọc hợp đồng một lần). Trả về (kết quả {mã: {...}}, lỗi {nhóm: lý do}, thống kê, cảnh báo)."""
    nhom = {}
    for r in rows:
        nhom.setdefault(r["nhom"], []).append(r)
    can = _uoc_token(text) + 3000
    llm.num_ctx = 32768 if can <= 32768 else (49152 if can <= 49152 else 65536)
    canh_bao = ""
    if can > 65536:
        canh_bao = (f"Hợp đồng quá dài (~{can:,} token): AI chỉ đọc được khoảng {65536 * 100 // can}% đầu. "
                    "Các tham số nằm ở phần sau có thể bị thiếu.")
        bao("    CẢNH BÁO: " + canh_bao)
    ket, loi = {}, {}
    tk = dict(token_vao=0, token_ra=0, giay=0.0)
    for i, (ten, ds) in enumerate(nhom.items(), 1):
        if dung is not None and dung.is_set():
            raise Dung()
        ma = {r["ma"] for r in ds}
        dsx = "\n".join(f"{r['ma']} | {r['ten']}" + (f" ({r['don_vi']})" if r["don_vi"] else "") for r in ds)
        user = (f"HỢP ĐỒNG\n{text}\n\n---\nĐọc hợp đồng trên và trích các tham số trong DANH SÁCH, giữ đúng mã."
                f"\n\nDANH SÁCH THAM SỐ (mã | tham số)\n{dsx}\n\nĐỊNH DẠNG ĐẦU RA\n{M.AI_FORMAT}")
        for lan in range(1, so_lan + 1):
            bao(f"    nhóm {i}/{len(nhom)} – {ten} ({len(ds)} tham số)" + (f" – thử lại lần {lan}" if lan > 1 else "") + " …")
            try:
                raw = llm.chat(M.AI_SYSTEM, user)
                m = re.search(r"\{.*\}", raw, re.S)
                got = json.loads(m.group(0) if m else raw)
                ket.update({k: v for k, v in got.items() if k in ma})
                for k in tk:
                    tk[k] += llm.stats.get(k) or 0
                if i == 1 and llm.stats.get("token_vao") and llm.stats["token_vao"] < _uoc_token(text) * 0.6:
                    canh_bao = canh_bao or ("Máy AI có thể đã cắt bớt hợp đồng (đọc ít token hơn dự kiến) – "
                                            "kiểm tra kỹ các tham số Không tìm thấy.")
                bao(f"      xong sau {llm.stats.get('giay')} giây")
                break
            except Exception as e:
                if lan < so_lan:
                    bao(f"      lỗi: {e} – chờ 5 giây rồi hỏi lại nhóm này")
                    time.sleep(5)
                else:
                    loi[ten] = str(e)
                    bao(f"      LỖI nhóm {ten}: {e} – bỏ qua nhóm này, các nhóm khác vẫn chạy")
    tk["giay"] = round(tk["giay"], 1)
    return ket, loi, tk, canh_bao


def chuan_ai(ai, rows):
    """AI hay ghi tỷ lệ thiếu dấu % ("90" thay vì "90%"): tham số đơn vị % mà AI ghi số trần > 1 thì thêm %."""
    dv = {r["ma"]: r["don_vi"] for r in rows}
    for ma, v in ai.items():
        if isinstance(v, dict) and str(dv.get(ma, "")).startswith("%"):
            g = str(v.get("gia_tri", "")).strip()
            if re.fullmatch(r"\d+([.,]\d+)?", g) and M.so(g) > 1:
                v["gia_tri"] = g + "%"
    return ai


def mot_hop_dong(path, cfg, out_dir, llm, lam_lai=False, bao=print, dung=None):
    t0 = time.time()
    paras = M.doc_hop_dong(path)
    if len(paras) < 20:
        raise RuntimeError(f"Đọc được quá ít chữ ({len(paras)} đoạn) – có thể là bản scan, cần OCR trước.")
    rows = M.trich_xuat(paras)
    log = dict(file=Path(path).name, so_doan=len(paras), ky_tu=0, token_vao=None, token_ra=None, giay_ai=None,
               nguon_ai="", ket_qua="")
    cmp = None
    if llm is not None:
        an, ca = cfg["an_danh"], cfg["gui_ca_hop_dong"]
        cache = out_dir / "_ai_json" / (bam(path, cfg["model"], an, ca) + ".json")
        text = "\n".join(t for _, t in (paras if ca else M.rut_gon(paras)))
        if an:
            text = M.an_danh(text, rows)
        log["ky_tu"] = len(text)
        if cache.exists() and not lam_lai:
            ai = json.loads(cache.read_text(encoding="utf-8"))
            log["nguon_ai"] = "Dùng lại kết quả AI lần trước"
            bao("    dùng lại kết quả AI lần trước (không hỏi lại)")
        else:
            ai, loi, tk, canh_bao = hoi_ai_theo_nhom(llm, text, rows, bao, dung)
            if not ai:
                raise RuntimeError("Máy AI không trả lời được nhóm nào: " + "; ".join(f"{k}: {v}" for k, v in loi.items()))
            if not loi:  # chỉ lưu khi đủ mọi nhóm, để lần sau hỏi lại nhóm lỗi
                cache.parent.mkdir(parents=True, exist_ok=True)
                cache.write_text(json.dumps(ai, ensure_ascii=False), encoding="utf-8")
            log.update(tk, giay_ai=tk["giay"], nguon_ai=f"Hỏi {cfg['model']} tại {cfg['may_ai']}")
            ghi_chu = ([f"Nhóm AI lỗi: " + "; ".join(f"{k} ({v[:80]})" for k, v in loi.items())] if loi else []) + \
                      ([canh_bao] if canh_bao else [])
            if ghi_chu:
                log["canh_bao"] = " | ".join(ghi_chu)
        goc = "\n".join(t for _, t in paras)
        chuan_ai(ai, rows)
        cmp = M.so_sanh_ai(rows, ai, goc, M.an_danh(goc, rows) if an else "")
        ten = _ten_an_toan(Path(path).stem)
        M.ghi_so_sanh(out_dir / f"so_sanh_AI_{ten}.xlsx", path, cmp,
                      dict(token_vao=log.get("token_vao"), token_ra=log.get("token_ra"), giay=log.get("giay_ai")),
                      cfg["model"], len(text), an)
    rows, tt = gop(rows, cmp or [])
    ten = _ten_an_toan(Path(path).stem)
    M.ghi_master(out_dir / f"master_{ten}.xlsx", path, rows)
    log["giay"] = round(time.time() - t0, 1)
    log["ket_qua"] = "OK" + (" – " + log["canh_bao"] if log.get("canh_bao") else "")
    log["master"] = f"master_{ten}.xlsx"
    return rows, tt, cmp, log


# ------------------------------------------------------------------ file tổng hợp
def _put(ws, r, c, v, fill=None, bold=False, wrap=False, fmt=None):
    cell = ws.cell(r, c, v)
    cell.font = M.BOLD if bold else M.BASE
    cell.alignment = Alignment(vertical="top", wrap_text=wrap)
    if fill:
        cell.fill = fill
    if fmt:
        cell.number_format = fmt
    return cell


def ghi_tong_hop(out, ket_qua, logs, cfg, dung_ai):
    wb = openpyxl.Workbook()
    # --- Tổng hợp: mỗi HĐ một dòng
    ws = wb.active
    ws.title = "Tong hop"
    _put(ws, 1, 1, "TỔNG HỢP THAM SỐ HỢP ĐỒNG – TRÍCH HÀNG LOẠT", bold=True)
    _put(ws, 2, 1, f"Lập lúc {dt.datetime.now():%d/%m/%Y %H:%M} | "
                   + (f"Máy AI {cfg['may_ai']} – model {cfg['model']} – ẩn danh: {'có' if cfg['an_danh'] else 'không'}"
                      if dung_ai else "Chỉ dùng quy tắc (không hỏi AI)")
                   + " | Mọi giá trị đều 'Chưa xác nhận' – xác nhận trong file master từng hợp đồng.")
    _put(ws, 3, 1, "Màu: xanh = quy tắc và AI khớp; vàng = chỉ AI thấy / khác cách ghi (đọc trích dẫn); "
                   "đỏ = quy tắc và AI lệch, hoặc AI bịa trích dẫn; xám = không tìm thấy. Rê chuột vào ô để xem ghi chú.")
    if not ket_qua:
        _put(ws, 5, 1, "Không có hợp đồng nào xử lý được – xem sheet Nhat ky.")
    else:
        mau = ket_qua[0][1]
        cols = ["File hợp đồng", "Cần xem", "Thiếu bắt buộc"] + [f"{r['ma']} {r['ten']}" for r in mau]
        for j, h in enumerate(cols, 1):
            c = _put(ws, 5, j, h, wrap=True)
            c.font, c.fill = WHITE_B, HEAD
        ws.row_dimensions[5].height = 60
        for i, (path, rows, tt, _) in enumerate(ket_qua, 6):
            can = sum(1 for m, t in tt.items() if TT_MAU.get(t, (None, False))[1])
            thieu = [r["ma"] for r in rows if r["ma"] in BAT_BUOC and r["gia_tri"] is None]
            if "TT-20" in thieu and "TT-21" not in thieu:
                thieu.remove("TT-20")
            if "TT-21" in thieu and "TT-20" not in thieu:
                thieu.remove("TT-21")
            ws.row_dimensions[i].height = 30
            _put(ws, i, 1, Path(path).name, wrap=True)
            _put(ws, i, 2, can, fill=M.RED if can else M.GREEN)
            _put(ws, i, 3, ", ".join(thieu) or "Đủ", fill=M.RED if thieu else M.GREEN, wrap=True)
            for j, r in enumerate(rows, 4):
                v = r["gia_tri"]
                fmt = "0.00%" if r["don_vi"] in ("%", "%/ngày") and isinstance(v, float) else (
                    "#,##0" if r["don_vi"] == "đồng" and isinstance(v, (int, float)) else None)
                c = _put(ws, i, j, v if not isinstance(v, str) else (v[:117] + "…" if len(v) > 120 else v), fmt=fmt,
                         fill=TT_MAU.get(tt.get(r["ma"]), (None,))[0])
                note = f"Tình trạng: {tt.get(r['ma'], '')}\nCách lấy: {r['cach']}\n{r['dieu']}"
                if r["ghi_chu"]:
                    note += f"\nGhi chú: {r['ghi_chu']}"
                if r["trich"]:
                    note += f"\nTrích dẫn: {r['trich'][:300]}"
                c.comment = Comment(note, "trich_lo")
                c.comment.width, c.comment.height = 420, 220
        ws.freeze_panes = "B6"
        ws.column_dimensions["A"].width = 42
        ws.column_dimensions["B"].width = 9
        ws.column_dimensions["C"].width = 14
        for j in range(4, len(cols) + 1):
            ws.column_dimensions[get_column_letter(j)].width = 18
        ws.auto_filter.ref = f"A5:{get_column_letter(len(cols))}{5 + len(ket_qua)}"

    # --- Cần xác nhận
    ws = wb.create_sheet("Can xac nhan")
    hdr = ["File hợp đồng", "Mã", "Tham số", "Tình trạng", "Giá trị đang lấy", "Theo AI", "Ghi chú",
           "Điều khoản", "Trích dẫn"]
    for j, h in enumerate(hdr, 1):
        c = _put(ws, 1, j, h)
        c.font, c.fill = WHITE_B, HEAD
    r0 = 2
    thu_tu = ["Lệch", "AI bịa trích dẫn", "Chỉ AI thấy", "Khác cách ghi", "Không tìm thấy"]
    dong = []
    for path, rows, tt, cmp in ket_qua:
        ai = {x["ma"]: x for x in (cmp or [])}
        for r in rows:
            t = tt.get(r["ma"])
            if t in thu_tu and not (t == "Không tìm thấy" and r["ma"] not in BAT_BUOC):
                dong.append((thu_tu.index(t), Path(path).name, r, t, ai.get(r["ma"], {}).get("ai")))
    for k, (_, f, r, t, a) in enumerate(sorted(dong, key=lambda x: (x[0], x[1], x[2]["ma"])), r0):
        vals = [f, r["ma"], r["ten"], t, r["gia_tri"], a, r["ghi_chu"], r["dieu"], r["trich"]]
        for j, v in enumerate(vals, 1):
            _put(ws, k, j, v if not isinstance(v, str) else v[:500], wrap=j in (1, 5, 6, 7, 9),
                 fill=TT_MAU[t][0] if j == 4 else None)
    if not dong:
        _put(ws, 2, 1, "Không có tham số nào cần xem thêm (vẫn phải xác nhận trong master trước khi dùng).")
    for j, w in enumerate([36, 8, 30, 16, 26, 26, 50, 26, 70], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:I{max(2, len(dong) + 1)}"

    # --- Nhật ký
    ws = wb.create_sheet("Nhat ky")
    hdr = ["File hợp đồng", "Kết quả", "Số đoạn", "Ký tự gửi AI", "Token vào", "Token ra", "Giây (AI)", "Giây (tổng)",
           "Nguồn AI", "File master"]
    for j, h in enumerate(hdr, 1):
        c = _put(ws, 1, j, h)
        c.font, c.fill = WHITE_B, HEAD
    for i, lg in enumerate(logs, 2):
        vals = [lg["file"], lg["ket_qua"], lg.get("so_doan"), lg.get("ky_tu"), lg.get("token_vao"), lg.get("token_ra"),
                lg.get("giay_ai"), lg.get("giay"), lg.get("nguon_ai"), lg.get("master", "")]
        for j, v in enumerate(vals, 1):
            _put(ws, i, j, v, wrap=j in (1, 2, 9), fill=(M.RED if j == 2 and not lg["ket_qua"].startswith("OK") else
                                                (M.YELLOW if j == 2 and lg["ket_qua"] != "OK" else None)))
    for j, w in enumerate([42, 40, 9, 12, 10, 10, 10, 10, 40, 40], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    wb.save(out)


# ------------------------------------------------------------------ chạy cả lô
def tim_file(nguon, bao=print):
    files = []
    for s in nguon:
        p = Path(s)
        if p.is_dir():
            files += [f for f in sorted(p.rglob("*")) if f.suffix.lower() in DUOI and not f.name.startswith("~$")]
        elif p.suffix.lower() in DUOI and p.exists():
            files.append(p)
        else:
            bao(f"  Bỏ qua (không phải hợp đồng .doc/.docx/.pdf): {s}")
    seen, out = set(), []
    for f in files:
        k = str(f.resolve()).lower()
        if k not in seen:
            seen.add(k)
            out.append(f)
    return out


def chay(nguon, out_dir, lam_lai=False, khong_ai=False, bao=print, dung=None, tien_do=None):
    """Chạy cả lô. bao(msg): in tiến độ; dung: threading.Event để dừng; tien_do(i, trang_thai): cập nhật từng file.
    Trả về đường dẫn file TONG_HOP. Lỗi làm dừng cả lô -> raise LoiChay."""
    cfg = M.cau_hinh_ai()
    files = tim_file(nguon, bao)
    if not files:
        raise LoiChay("Không thấy file hợp đồng (.doc, .docx, .pdf) nào trong danh sách.")
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    llm = None
    if not khong_ai:
        if not kiem_tra(cfg, thu_hoi=False, bao=bao):
            cfg = tu_tim_lai(cfg, bao)
            if cfg is None:
                raise LoiChay("Chưa kết nối được máy AI. Mở tab Máy AI để nhập hoặc tự tìm IP máy AI, "
                              "hoặc chọn 'Chỉ dùng quy tắc'.")
        llm = M.LocalLLM(cfg["model"], cfg["may_ai"], cfg["api"], timeout=int(cfg["timeout_giay"]))
    bao(f"\n{len(files)} hợp đồng → {out_dir}\n")
    ket_qua, logs = [], []
    for i, f in enumerate(files, 1):
        if dung is not None and dung.is_set():
            bao("Đã dừng theo yêu cầu. Lần chạy sau sẽ dùng lại kết quả AI đã có.")
            break
        bao(f"[{i}/{len(files)}] {f.name}")
        if tien_do:
            tien_do(i - 1, "Đang chạy")
        try:
            rows, tt, cmp, log = mot_hop_dong(f, cfg, out_dir, llm, lam_lai, bao, dung)
            c = Counter(tt.values())
            tom = ", ".join(f"{k}: {n}" for k, n in c.most_common())
            bao(f"    xong {log['giay']} giây – {log['nguon_ai'] or 'chỉ quy tắc'} – {tom}")
            ket_qua.append((f, rows, tt, cmp))
            logs.append(log)
            if tien_do:
                can = sum(1 for t in tt.values() if TT_MAU.get(t, (None, False))[1])
                tien_do(i - 1, ("Xong, có cảnh báo" if log.get("canh_bao") else "Xong") + f" – {can} ô cần xem")
        except (Dung, KeyboardInterrupt):
            bao("    Đã dừng theo yêu cầu. Hợp đồng này chưa xong; lần chạy sau sẽ làm tiếp.")
            logs.append(dict(file=f.name, ket_qua="Dừng giữa chừng"))
            if tien_do:
                tien_do(i - 1, "Đã dừng")
            break
        except Exception as e:
            msg = f"LỖI: {e}" if isinstance(e, RuntimeError) else f"LỖI: {e.__class__.__name__}: {e}"
            if llm is not None and isinstance(e, (urllib.error.URLError, TimeoutError, ConnectionError)):
                msg += " (máy AI không trả lời – kiểm tra máy AI rồi chạy lại; hợp đồng đã xong sẽ không hỏi lại)"
            bao("    " + msg)
            logs.append(dict(file=f.name, ket_qua=msg))
            if tien_do:
                tien_do(i - 1, "Lỗi")
    out = out_dir / f"TONG_HOP_{dt.datetime.now():%Y%m%d_%H%M%S}.xlsx"
    ghi_tong_hop(out, ket_qua, logs, cfg, llm is not None)
    loi = sum(1 for lg in logs if not lg["ket_qua"].startswith("OK"))
    bao(f"\nXong {len(ket_qua)}/{len(files)} hợp đồng" + (f", {loi} lỗi (xem sheet Nhat ky)" if loi else "")
        + f".\nTổng hợp: {out}")
    return out


def main():
    ap = argparse.ArgumentParser(description="Trích tham số hàng loạt hợp đồng qua máy Local AI trong LAN")
    sub = ap.add_subparsers(dest="lenh", required=True)
    sub.add_parser("kiem-tra", help="Kiểm tra kết nối máy AI")
    c = sub.add_parser("cau-hinh", help="Lưu địa chỉ máy AI / model vào cau_hinh_may_ai.json")
    c.add_argument("--may-ai", help="IP hoặc địa chỉ máy AI, vd 192.168.1.50")
    c.add_argument("--model")
    c.add_argument("--an-danh", choices=["co", "khong"])
    r = sub.add_parser("chay", help="Trích cả lô")
    r.add_argument("nguon", nargs="+", help="Thư mục chứa hợp đồng hoặc từng file")
    r.add_argument("--out", default=str(Path(__file__).with_name("ket_qua_AI") / "trich_lo"))
    r.add_argument("--lam-lai", action="store_true", help="Hỏi lại AI kể cả hợp đồng đã có kết quả")
    r.add_argument("--khong-ai", action="store_true", help="Chỉ dùng quy tắc, không hỏi AI")
    a = ap.parse_args()
    if a.lenh == "kiem-tra":
        sys.exit(0 if kiem_tra() else 1)
    elif a.lenh == "cau-hinh":
        cfg = ghi_cau_hinh(a.may_ai, a.model, None if a.an_danh is None else a.an_danh == "co")
        sys.exit(0 if kiem_tra(cfg) else 1)
    else:
        try:
            chay(a.nguon, a.out, a.lam_lai, a.khong_ai)
        except LoiChay as e:
            sys.exit(f"Dừng: {e}")


if __name__ == "__main__":
    main()
